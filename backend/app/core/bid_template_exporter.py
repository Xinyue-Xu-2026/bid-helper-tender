"""模板式商务标导出：以 data/商务标模板.docx 为底稿，替换其中
「三、人员配备表」与「四、相关业绩一览表」两节的表格内容，其余部分原样保留。

表定位策略：按 body 顺序遍历段落/表格，用 Heading 1/2/3 文本（strip 后前缀匹配）
维护当前所在章节状态机：
- H1 "三、人员配备表" 内：首个表格→表A（人员安排）；
  H2 "2、项目负责人" 后的 H3 "（2）项目负责人业绩" 内第 1/2 张表→表B/表C（负责人业绩）；
  H2 "3、项目组人员" 后的表→表D。
- H1 "四、相关业绩一览表" 内：H2 前的首表→表E；H2 "1、…" 后的表→表F；H2 "2、…" 后的表→表G。

行替换：保留表头行，deepcopy 首数据行作为样式 donor 逐条填充后插入，
删除全部原数据行（无记录时即只剩表头）。
"""
from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.table import Table
from docx.text.paragraph import Paragraph

_HEADING_STYLES = {"Heading 1": 1, "Heading 2": 2, "Heading 3": 3,
                   "标题 1": 1, "标题 2": 2, "标题 3": 3}


def _iter_block_items(doc):
    """按 body 顺序产出 Paragraph / Table。"""
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def _heading_level(para: Paragraph) -> int:
    style = para.style
    return _HEADING_STYLES.get(style.name if style is not None else "", 0)


def _locate_tables(doc) -> dict:
    """定位七张目标表，返回 {role: Table}，role ∈ a/b/c/d/e/f/g。"""
    roles: dict = {}
    section = 0          # 3=人员配备表 / 4=相关业绩一览表 / 0=其他
    sub3 = ""            # 三、内小节：""（表A区）/ lead / lead_perf / members
    sub4 = 0             # 四、内小节：0（表E区）/ 1 / 2
    lead_perf_tables = 0
    for block in _iter_block_items(doc):
        if isinstance(block, Paragraph):
            level = _heading_level(block)
            if not level:
                continue
            text = block.text.strip()
            if level == 1:
                if text.startswith("三、人员配备表"):
                    section, sub3, sub4 = 3, "", 0
                elif text.startswith("四、相关业绩一览表"):
                    section, sub3, sub4 = 4, "", 0
                else:
                    section = 0
            elif section == 3 and level == 2:
                if text.startswith("2、"):
                    sub3 = "lead"
                elif text.startswith("3、项目组人员"):
                    sub3 = "members"
            elif section == 3 and level == 3 and text.startswith("（2）项目负责人业绩"):
                sub3 = "lead_perf"
                lead_perf_tables = 0
            elif section == 4 and level == 2:
                if text.startswith("1、"):
                    sub4 = 1
                elif text.startswith("2、"):
                    sub4 = 2
        else:  # Table
            if section == 3:
                if sub3 == "":
                    roles.setdefault("a", block)
                elif sub3 == "lead_perf":
                    lead_perf_tables += 1
                    roles.setdefault("b" if lead_perf_tables == 1 else "c", block)
                elif sub3 == "members":
                    roles.setdefault("d", block)
            elif section == 4:
                if sub4 == 0:
                    roles.setdefault("e", block)
                elif sub4 == 1:
                    roles.setdefault("f", block)
                elif sub4 == 2:
                    roles.setdefault("g", block)
    return roles


def _set_cell_text(cell, text: str) -> None:
    """写单元格文本（\n 转软换行），尽量保留首段/首 run 的格式。"""
    lines = str(text if text is not None else "").split("\n")
    para = cell.paragraphs[0]
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    if para.runs:
        run = para.runs[0]
        for r in para.runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        run = para.add_run()
    run.text = lines[0]
    for line in lines[1:]:
        run.add_break()
        run.add_text(line)


def _fill_table(table: Table, rows_data: list) -> None:
    """数据行整体替换：保留表头；无数据时仅留表头。"""
    rows = table.rows
    donor_tr = deepcopy(rows[1]._tr if len(rows) >= 2 else rows[0]._tr)
    for r in list(table.rows[1:]):
        r._tr.getparent().remove(r._tr)
    for data in rows_data or []:
        table._tbl.append(deepcopy(donor_tr))
        new_row = table.rows[-1]
        for idx, cell in enumerate(new_row.cells):
            _set_cell_text(cell, data[idx] if idx < len(data) else "")


# 注册证书类占位标签关键词（职业资格证书/注册证书都算注册证书类）
_CERT_IMAGE_KEYWORDS = ("注册证书", "职业资格证书", "资格证书")


def _clear_table_images(table: Table) -> None:
    """清除占位表数据行里已存在的图片（模板是上一次投标填过的文档，会残留旧扫描件）。"""
    for row in table.rows[1:]:
        cell = row.cells[0]
        for dr in list(cell._tc.iter(qn("w:drawing"))):
            dr.getparent().remove(dr)


def _insert_image_into_table(table: Table, image_path: str) -> None:
    """往图片占位表（1列，首行是标签文字）插入图片。无文件/路径无效则跳过。"""
    if not image_path or not Path(image_path).exists():
        return
    row = table.rows[1] if len(table.rows) > 1 else table.rows[0]
    cell = row.cells[0]
    para = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    para.add_run().add_picture(image_path, width=Cm(8))


def _insert_person_images(doc, persons) -> None:
    """按人员顺序把图片插入「三、人员配备表」的占位表（1列，首行含标签文字）。
    persons 已是负责人优先排序：[{name, fields, is_lead}]。"""
    if not persons:
        return
    person_images = []
    for p in persons:
        f = p.get("fields") or {}
        certs = [c for c in (f.get("证书") or []) if isinstance(c, dict)]
        person_images.append({
            "职称证书": f.get("职称证书扫描件") or "",
            "社保": f.get("社保缴纳证明扫描件") or "",
            "certs": [c.get("扫描件") or "" for c in certs],
        })

    in_section3 = False
    sub = ""          # "" / "lead" / "members"
    cur_person = None  # person_images 下标
    member_idx = 0
    cert_slot = 0

    for block in _iter_block_items(doc):
        if isinstance(block, Paragraph):
            level = _heading_level(block)
            if not level:
                continue
            text = block.text.strip()
            if level == 1:
                if text.startswith("三、人员配备表"):
                    in_section3, sub, cur_person = True, "", None
                elif text.startswith("四、"):
                    in_section3 = False
                else:
                    in_section3 = False
            elif in_section3 and level == 2:
                if text.startswith("2、项目负责人"):
                    sub, cur_person, cert_slot = "lead", 0, 0
                elif text.startswith("3、项目组人员"):
                    sub, cur_person, member_idx, cert_slot = "members", None, 0, 0
                else:
                    sub = ""
            elif in_section3 and level == 3:
                if sub == "members" and "项目组其他人员" in text:
                    cur_person = 1 + member_idx
                    member_idx += 1
                    cert_slot = 0
        else:  # Table
            if not in_section3 or cur_person is None:
                continue
            if len(block.columns) != 1:      # 只处理 1 列图片占位表
                continue
            label = (block.rows[0].cells[0].text or "").strip()
            # 先清掉占位表里残留的旧扫描件（模板为上次投标填过的文档）
            _clear_table_images(block)
            if cur_person >= len(person_images):
                continue
            pi = person_images[cur_person]
            if "职称证书" in label:
                _insert_image_into_table(block, pi["职称证书"])
            elif "社保" in label:
                _insert_image_into_table(block, pi["社保"])
            elif any(k in label for k in _CERT_IMAGE_KEYWORDS):
                if cert_slot < len(pi["certs"]):
                    _insert_image_into_table(block, pi["certs"][cert_slot])
                    cert_slot += 1


# ---------- 残留文本替换（项目编号/项目名称/日期）与人员子标题改名 ----------
#
# 模板残留格式（以 data/商务标模板.docx 实测为准）：
# - 项目编号：封面「编       号：YZGKWD-2026-009号」、商务条款响应表
#   「项目编号：YZGKWD-2026-009号」、开标一览表单元格裸值「YZGKWD-2026-009号」，
#   共同子串为编号本体（含尾字「号」）。
# - 项目名称：封面「项 目 名 称：×××」（字间有空格）、正文「项目名称：×××」、
#   人员承诺书正文裸名「本单位参与本次 ×××项目投标活动」、开标一览表单元格裸值。
#   运行时先从带「项目名称」标签的段落中提取旧名，再全局替换该子串。
# - 日期：封面「日      期 ：2026年8月11日」、各表页脚「日 期：\t2026年8月11日」、
#   承诺函「日期：2026年8月11日」「日期:2026年8月11日」等变体，统一用
#   「日\s*期\s*[：:]」标签定位，仅替换其后的 Y年M月D日。
# - 目录（toc 样式段落）整段跳过，上述替换一律不触碰。

_PROJECT_NO_RE = re.compile(r"[A-Za-z0-9]+(?:-[A-Za-z0-9]+)+号")
_PROJECT_NAME_LABEL_RE = re.compile(r"项\s*目\s*名\s*称\s*[：:]\s*(\S+)")
_DOC_DATE_RE = re.compile(
    r"(日\s*期\s*[：:]\s*)\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日")
_MEMBER_HEADING_RE = re.compile(r"^（\d+）项目组其他人员-")
# 标签式空白占位（P3）：仅当标签后是空白/下划线占位（或行尾）时填充，
# 已含真实值的行（如"招标人：某某中心"）不动
_TENDERER_LABEL_RE = re.compile(r"(招标人[：:])(?=[\s_　＿]|$)[\s_　＿]*")
_BIDDER_LABEL_RE = re.compile(r"(投标人名称[：:])(?=[\s_　＿]|$)[\s_　＿]*")


# ---------- 通用占位符规则核心（V1.2） ----------
# 单一事实源：填充端（_replace_stale_text/compute_text_subs）与校验端
# （bid_verify）共用 build_placeholder_rules 产出的规则列表。

DEFAULT_PLACEHOLDER_SYNONYMS = {
    "工程项目名称": "项目名称", "工程名称": "项目名称",
    "采购人": "招标人", "供应商名称": "投标人名称",
}
PLACEHOLDER_LABEL_KEYS = {
    "项目名称": "project_name", "招标人": "tenderer",
    "投标人名称": "bidder_name", "项目编号": "project_no",
    "招标编号": "project_no", "标段名称": "section_name",
    "标段编号": "section_no", "日期": "doc_date",
}
BRACKET_LABEL_KEYS = {
    "项目名称": "project_name", "工程名称": "project_name",
    "工程项目名称": "project_name", "招标人名称": "tenderer",
    "招标人": "tenderer", "采购人": "tenderer",
    "投标人名称": "bidder_name", "供应商名称": "bidder_name",
    "标段名称": "section_name", "标段编号": "section_no",
    "标段": "section_name", "日期": "doc_date",
}
_SECTION_COMBO_RE = re.compile(
    r"([（(])\s*(?:项目名称|工程名称|工程项目名称)\s*[）)](\s*)"
    r"([（(])\s*标段(?:名称)?\s*[）)]")
_UNRECOGNIZED_UNDERSCORE_RE = re.compile(r"_{4,}|＿{4,}")
_DOC_DATE_BLANK_RE = re.compile(
    r"(日\s*期\s*[：:]?\s*)[_＿\s]*年[_＿\s]*月[_＿\s]*日")
_PLACEHOLDER_ONLY_RE = re.compile(r"^[_＿\s]+$")


def _effective_label_keys(synonyms):
    keys = dict(PLACEHOLDER_LABEL_KEYS)
    for alias, canonical in (synonyms or DEFAULT_PLACEHOLDER_SYNONYMS).items():
        key = PLACEHOLDER_LABEL_KEYS.get(canonical)
        if key and alias not in keys:
            keys[alias] = key
    return keys


def _date_text(doc_date):
    m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", (doc_date or "").strip())
    return f"{int(m.group(1))}年{int(m.group(2))}月{int(m.group(3))}日" if m else ""


def _iter_cell_paragraphs(cell):
    for para in cell.paragraphs:
        yield para
    for table in cell.tables:
        for row in table.rows:
            for sub in row.cells:
                yield from _iter_cell_paragraphs(sub)


_MC_FALLBACK_TAG = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


def _iter_textbox_paragraphs(doc):
    """文本框（含嵌套）内段落，不重复产出。
    Word 写出的文本框为 mc:AlternateContent：mc:Choice（DrawingML）+
    mc:Fallback（VML）各含一份相同 w:txbxContent，跳过 Fallback 内的，
    避免同一段落产出两次。"""
    for txbx in doc.element.body.iter(qn("w:txbxContent")):
        if any(a.tag == _MC_FALLBACK_TAG for a in txbx.iterancestors()):
            continue
        for child in txbx.iterchildren():
            if child.tag == qn("w:p"):
                yield Paragraph(child, doc)
            elif child.tag == qn("w:tbl"):
                for row in Table(child, doc).rows:
                    for cell in row.cells:
                        yield from _iter_cell_paragraphs(cell)


def _iter_all_paragraphs(doc):
    """正文段落 + 所有表格（含嵌套）单元格内段落 + 文本框内段落。"""
    for para in doc.paragraphs:
        yield para
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_cell_paragraphs(cell)
    yield from _iter_textbox_paragraphs(doc)


def _is_toc_paragraph(para: Paragraph) -> bool:
    style = para.style
    name = (style.name if style is not None else "") or ""
    return name.lower().startswith("toc")


def _para_text(para: Paragraph) -> str:
    return "".join(r.text for r in para.runs)


def _rewrite_paragraph_text(para: Paragraph, text: str) -> None:
    """整段改写：保留首个 run 的格式，删除其余 run。"""
    runs = para.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        para.add_run(text)


def _replace_in_paragraph(para: Paragraph, pattern, repl) -> bool:
    """对段落整体文本做正则替换；命中时以首 run 格式重写整段。"""
    full = _para_text(para)
    if not full:
        return False
    new, count = pattern.subn(repl, full)
    if not count:
        return False
    _rewrite_paragraph_text(para, new)
    return True


def _discover_stale_values(doc) -> tuple:
    """从模板中提取残留的项目编号与项目名称（带标签的首个命中）。"""
    stale_no = ""
    stale_name = ""
    for para in _iter_all_paragraphs(doc):
        text = _para_text(para).strip()
        if not text:
            continue
        if not stale_no:
            m = _PROJECT_NO_RE.search(text)
            if m:
                stale_no = m.group(0)
        if not stale_name:
            m = _PROJECT_NAME_LABEL_RE.search(text)
            if m:
                stale_name = m.group(1)
        if stale_no and stale_name:
            break
    return stale_no, stale_name


def build_placeholder_rules(doc, *, project_no: str = "", project_name: str = "",
                            doc_date: str = "", tenderer: str = "",
                            bidder_name: str = "", section_name: str = "",
                            section_no: str = "", synonyms=None) -> list:
    """生成占位符替换规则列表，每项 {"kind","label","key","value","pattern","repl"}。
    顺序：stale → 组合 → 标签（长度降序）→ 括号（长度降序）→ 日期。
    填充端与校验端共用本函数，保证"导出改了什么"与"校验豁免什么"一致。"""
    values = {
        "project_no": project_no, "project_name": project_name,
        "doc_date": doc_date, "tenderer": tenderer,
        "bidder_name": bidder_name, "section_name": section_name,
        "section_no": section_no,
    }
    rules = []
    # stale：模板残留的上次投标编号/名称（空白占位不算 stale）
    stale_no, stale_name = _discover_stale_values(doc)
    if project_no and stale_no:
        new_no = project_no if project_no.endswith("号") else f"{project_no}号"
        rules.append({"kind": "stale_no", "label": stale_no, "key": "project_no",
                      "value": new_no, "pattern": re.compile(re.escape(stale_no)),
                      "repl": new_no})
    if (project_name and stale_name and project_name != stale_name
            and not _PLACEHOLDER_ONLY_RE.match(stale_name)):
        rules.append({"kind": "stale_name", "label": stale_name,
                      "key": "project_name", "value": project_name,
                      "pattern": re.compile(re.escape(stale_name)),
                      "repl": project_name})
    # 组合：(项目名称)(标段名称) → 项目名 (标段名)
    if project_name and section_name:
        hit = any(_SECTION_COMBO_RE.search(_para_text(p))
                  for p in _iter_all_paragraphs(doc)
                  if not _is_toc_paragraph(p))
        if hit:
            # 保留底稿匹配到的括号风格与两组间空白
            # （半角输入→半角，全角输入→全角）
            def _combo_repl(m, pn=project_name, sn=section_name):
                close = "）" if m.group(3) == "（" else ")"
                return f"{pn}{m.group(2)}{m.group(3)}{sn}{close}"
            combo = f"{project_name} ({section_name})"
            rules.append({"kind": "section_combo",
                          "label": "(项目名称)(标段名称)",
                          "key": "project_name+section_name", "value": combo,
                          "pattern": _SECTION_COMBO_RE, "repl": _combo_repl})
    # 标签式空白占位：标签后仅空白/下划线（或行尾）时填充
    label_keys = _effective_label_keys(synonyms)
    for label in sorted(label_keys, key=len, reverse=True):
        key = label_keys[label]
        if key == "doc_date":
            continue  # 日期由专用日期规则处理（含 ____年__月__日 空白格式）
        value = str(values.get(key) or "").strip()
        if not value:
            continue
        pattern = re.compile(
            rf"({re.escape(label)}\s*[：:])(?=[\s_　＿]|$)[\s_　＿]*")
        rules.append({"kind": "label", "label": label, "key": key,
                      "value": value, "pattern": pattern,
                      "repl": (lambda m, l=label, v=value: m.group(1) + v)})
    # 括号式占位：(项目名称)/（招标人名称）等
    for label in sorted(BRACKET_LABEL_KEYS, key=len, reverse=True):
        key = BRACKET_LABEL_KEYS[label]
        value = str(values.get(key) or "").strip()
        if not value:
            continue
        if key == "doc_date":
            value = _date_text(doc_date)
            if not value:
                continue
        pattern = re.compile(
            rf"[（(【\[]\s*{re.escape(label)}\s*[）)】\]]")
        rules.append({"kind": "bracket", "label": label, "key": key,
                      "value": value, "pattern": pattern, "repl": value})
    # 日期：带数字的旧日期与空白日期占位
    date_text = _date_text(doc_date)
    if date_text:
        rules.append({"kind": "date", "label": "日期", "key": "doc_date",
                      "value": date_text, "pattern": _DOC_DATE_RE,
                      "repl": (lambda m, v=date_text: m.group(1) + v)})
        rules.append({"kind": "date_blank", "label": "日期", "key": "doc_date",
                      "value": date_text, "pattern": _DOC_DATE_BLANK_RE,
                      "repl": (lambda m, v=date_text: m.group(1) + v)})
    return rules


def compute_text_subs(doc, project_no: str = "", project_name: str = "",
                      doc_date: str = "", tenderer: str = "",
                      bidder_name: str = "", section_name: str = "",
                      section_no: str = "", synonyms=None) -> list:
    """计算残留文本替换规则 [(pattern, repl)]（= build_placeholder_rules 的
    (pattern, repl) 投影）。导出与防篡改校验（bid_verify）共用同一套规则，
    保证"导出改了什么"与"校验豁免什么"一致。"""
    return [(r["pattern"], r["repl"]) for r in build_placeholder_rules(
        doc, project_no=project_no, project_name=project_name,
        doc_date=doc_date, tenderer=tenderer, bidder_name=bidder_name,
        section_name=section_name, section_no=section_no, synonyms=synonyms)]


def scan_placeholders(doc, params: dict, synonyms=None) -> dict:
    """扫描底稿占位符命中情况：
    返回 {"matched": [{"label","value","count","kind"}],
          "suspicious": [{"text"}]}（未被任何规则命中的成串下划线段落）。"""
    params = params or {}
    if synonyms is None:
        synonyms = params.get("synonyms")
    rules = build_placeholder_rules(
        doc,
        project_no=str(params.get("project_no") or ""),
        project_name=str(params.get("project_name") or ""),
        doc_date=str(params.get("doc_date") or ""),
        tenderer=str(params.get("tenderer") or ""),
        bidder_name=str(params.get("bidder_name") or ""),
        section_name=str(params.get("section_name") or ""),
        section_no=str(params.get("section_no") or ""),
        synonyms=synonyms)
    texts = [_para_text(p) for p in _iter_all_paragraphs(doc)
             if not _is_toc_paragraph(p)]
    matched = []
    for r in rules:
        count = sum(len(r["pattern"].findall(t)) for t in texts)
        if count:
            matched.append({"label": r["label"], "value": r["value"],
                            "count": count, "kind": r["kind"]})
    suspicious = []
    for t in texts:
        if not t.strip():
            continue
        if (_UNRECOGNIZED_UNDERSCORE_RE.search(t)
                and not any(r["pattern"].search(t) for r in rules)):
            suspicious.append({"text": t})
    return {"matched": matched, "suspicious": suspicious}


def _replace_stale_text(doc, project_no: str = "", project_name: str = "",
                        doc_date: str = "", tenderer: str = "",
                        bidder_name: str = "", section_name: str = "",
                        section_no: str = "", synonyms=None) -> None:
    """替换模板中残留的上次投标文本：规则由 compute_text_subs 统一计算，
    此处仅逐段（含表格单元格与文本框，跳过 toc 段）应用。project_no/doc_date
    为空则跳过对应替换，project_name 为空则不替换；tenderer/bidder_name/
    section_name/section_no 为空则不填充对应标签/括号空白。"""
    subs = compute_text_subs(doc, project_no, project_name, doc_date,
                             tenderer=tenderer, bidder_name=bidder_name,
                             section_name=section_name, section_no=section_no,
                             synonyms=synonyms)
    if not subs:
        return
    for para in _iter_all_paragraphs(doc):
        if _is_toc_paragraph(para):
            continue
        for pattern, repl in subs:
            _replace_in_paragraph(para, pattern, repl)


def _rename_member_headings(doc, member_names: list) -> None:
    """H2「3、项目组人员」下 H3「（N）项目组其他人员-×××」的名字后缀
    按所选非负责人员顺序替换；多余标题/人员互不干扰。"""
    if not member_names:
        return
    in_members = False
    idx = 0
    for block in _iter_block_items(doc):
        if not isinstance(block, Paragraph):
            continue
        level = _heading_level(block)
        if not level:
            continue
        text = block.text.strip()
        if level == 1:
            in_members = False
        elif level == 2:
            in_members = text.startswith("3、项目组人员")
        elif level == 3 and in_members and _MEMBER_HEADING_RE.match(text):
            if idx >= len(member_names):
                break
            prefix = _MEMBER_HEADING_RE.match(text).group(0)
            _rewrite_paragraph_text(block, prefix + member_names[idx])
            idx += 1


def build_bid_docx_from_template(template_path: str, dest_path: str, data: dict) -> Path:
    """以模板为底稿填充七张表。data: {"a".."g": 各行列表}（assemble_bid_template_data 产出）。

    额外支持的 data 键：
    - persons: [{name, fields, is_lead}]（负责人优先排序），用于插图与人员子标题改名；
    - project_no / project_name / doc_date: 替换模板中残留的上次投标文本，
      project_no/doc_date 为空则跳过对应替换，project_name 为空则不替换。
    """
    doc = Document(template_path)
    tables = _locate_tables(doc)
    for role, rows in (data or {}).items():
        table = tables.get(role)
        if table is not None:
            _fill_table(table, rows or [])
    persons = (data or {}).get("persons") or []
    member_names = [str(p.get("name") or "") for p in persons if not p.get("is_lead")]
    _rename_member_headings(doc, member_names)
    _replace_stale_text(
        doc,
        project_no=str((data or {}).get("project_no") or "").strip(),
        project_name=str((data or {}).get("project_name") or "").strip(),
        doc_date=str((data or {}).get("doc_date") or "").strip(),
    )
    _insert_person_images(doc, persons)
    dest = Path(dest_path)
    doc.save(str(dest))
    return dest
