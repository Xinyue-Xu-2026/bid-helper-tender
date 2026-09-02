# -*- coding: utf-8 -*-
"""程序化导出 docx 段落锚点（index/style/text）与表格，供仿写定位与 LLM 规划。"""
from docx import Document


def dump_docx(path: str) -> dict:
    """返回 {"paragraphs": [{index, style, text}], "tables": [{index, rows}]}。

    段落 index 为 `enumerate(doc.paragraphs)` 的原始序号，是后续 insert_after
    的锚点；表格 cell 文本用 `.text.strip().replace('\\n', ' / ')` 归一化。
    注意 `doc.paragraphs` 只含正文段落，不含表格单元格内段落。
    """
    doc = Document(path)
    paragraphs = []
    for i, p in enumerate(doc.paragraphs):
        if not p.text.strip():
            continue
        paragraphs.append({
            "index": i,
            "style": p.style.name if p.style else "",
            "text": p.text,
        })
    tables = []
    for ti, t in enumerate(doc.tables):
        rows = [[c.text.strip().replace("\n", " / ") for c in row.cells]
                for row in t.rows]
        tables.append({"index": ti, "rows": rows})
    return {"paragraphs": paragraphs, "tables": tables}
