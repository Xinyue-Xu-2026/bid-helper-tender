"""Word 导出：模板底稿（保留封面/页眉页脚）+ 样式继承 + Markdown 表格/列表渲染。

- 打开模板后仅删除「首个 Heading 样式段落或 TOC 域」之后的正文，
  之前的封面/扉页/说明页（含表格、图片、分节符外的元素）原样保留；
  sectPr 与页眉页脚天然不动。
- 封面之后插入 Word 原生 TOC 域，再按章节树写内容。
- 标题/正文不设任何显式格式，全部走模板 styles.xml（Heading 1-4 / Normal）。
- 章节正文解析 Markdown 表格块 → 真 Word 表格；连续 "- " 行 → List Bullet
  （模板有该样式才用）；其余按行拆段落，连续空行折叠为一个空段保留节奏。
- 模板缺失/样式缺失时逐级回退，保证导出不炸。
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

_HEADING_STYLE_RE = re.compile(r"^(heading|标题)\s*[1-4]$", re.IGNORECASE)
# Markdown 表格分隔行：| --- | :--: | --- |
_MD_SEP_ROW_RE = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_BULLET_RE = re.compile(r"^[-*]\s+\S")
_BULLET_PREFIX_RE = re.compile(r"^[-*]\s+")
_TABLE_STYLE_PREFERRED = ("Table Grid", "网格型")
_TABLE_STYLE_SKIP = ("Normal Table", "普通表格")


def _para_style_name(para) -> str:
    try:
        style = getattr(para, "style", None)
        return ((style.name if style else "") or "").strip()
    except Exception:
        return ""


def _is_cut_point(el, doc) -> bool:
    """body 元素是否为正文起点（Heading 样式段落或含 TOC 域的段落）。"""
    if el.tag != qn("w:p"):
        return False
    para = Paragraph(el, doc)
    if _HEADING_STYLE_RE.match(_para_style_name(para)):
        return True
    try:
        xml = el.xml
        return "instrText" in xml and "TOC" in xml
    except Exception:
        return False


def _clear_body_keep_cover(doc):
    """删除正文，保留首个 Heading/TOC 域之前的封面元素与末尾 sectPr。"""
    body = doc.element.body
    children = list(body)
    cut = None
    for idx, el in enumerate(children):
        if _is_cut_point(el, doc):
            cut = idx
            break
    start = cut if cut is not None else 0  # 无标题/TOC 的模板：全清（旧行为）
    for el in children[start:]:
        if el.tag == qn("w:sectPr"):
            continue
        body.remove(el)


def insert_toc_field_at(para) -> None:
    """在既有段落 para 内插入 Word 原生 TOC 域（清空原段落内容）。"""
    para.clear()
    run = para.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-4" \\h \\z \\u'
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.append(begin); run._r.append(instr); run._r.append(separate); run._r.append(end)


def _insert_toc_field(doc):
    """插入 Word 原生 TOC 域（用户在 Word 中「更新域」生成目录，含 1-4 级标题）。"""
    insert_toc_field_at(doc.add_paragraph())


def _style_names(doc, style_type=None) -> set:
    try:
        return {s.name for s in doc.styles
                if style_type is None or s.type == style_type}
    except Exception:
        return set()


def _pick_table_style(doc) -> str:
    """优先用模板已有表格样式；取不到回退 'Table Grid'（默认模板自带）。"""
    names = _style_names(doc, WD_STYLE_TYPE.TABLE)
    for preferred in _TABLE_STYLE_PREFERRED:
        if preferred in names:
            return preferred
    for name in names - set(_TABLE_STYLE_SKIP):
        return name
    return "Table Grid"


def _add_heading(doc, title: str, level: int):
    """标题只挂样式不显式设格式；模板缺对应 Heading 样式时退化为正文段落。"""
    level = max(1, min(4, int(level or 1)))
    try:
        doc.add_heading(title, level=level)
    except KeyError:
        doc.add_paragraph(title)


def _split_md_row(line: str) -> list:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _add_md_table(doc, rows: list, style_name: str):
    cols = max(len(r) for r in rows)
    rows = [r + [""] * (cols - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=cols)
    if style_name:
        try:
            table.style = doc.styles[style_name]
        except KeyError:
            pass
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.cell(ri, ci).text = val


def _render_content(doc, content: str, table_style: str, bullet_ok: bool):
    """按行渲染章节正文：Markdown 表格块 → 真表格；'- ' 列表 → List Bullet。"""
    lines = (content or "").split("\n")
    i, n = 0, len(lines)
    blank_pending = False
    while i < n:
        stripped = lines[i].strip()
        if not stripped:
            blank_pending = True   # 连续空行折叠为一个空段，保留节奏
            i += 1
            continue
        # Markdown 表格块：当前行 + 次行分隔行
        if stripped.startswith("|") and i + 1 < n \
                and _MD_SEP_ROW_RE.match(lines[i + 1].strip()):
            rows = [_split_md_row(stripped)]
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append(_split_md_row(lines[i]))
                i += 1
            _add_md_table(doc, rows, table_style)
            blank_pending = False
            continue
        if blank_pending:
            doc.add_paragraph("")
            blank_pending = False
        # 连续 "- " 列表行
        if bullet_ok and _BULLET_RE.match(stripped):
            while i < n and _BULLET_RE.match(lines[i].strip()):
                doc.add_paragraph(_BULLET_PREFIX_RE.sub("", lines[i].strip(), count=1),
                                  style="List Bullet")
                i += 1
            continue
        doc.add_paragraph(stripped)
        i += 1


def export_word(project: dict, template: dict, sections: list, dest_path: str) -> str:
    """导出标书 Word。sections 为空抛 ValueError；模板缺失回退默认样式。返回写出路径。"""
    if not sections:
        raise ValueError("章节树为空，无法导出")
    template_path = (template or {}).get("file_path") or ""
    if template_path and Path(template_path).exists():
        doc = Document(template_path)   # 模板底稿：styles/页眉页脚/封面全继承
        _clear_body_keep_cover(doc)     # 只清正文，保留封面与 sectPr
    else:
        doc = Document()                # 回退 python-docx 默认样式
    table_style = _pick_table_style(doc)
    bullet_ok = "List Bullet" in _style_names(doc)
    _insert_toc_field(doc)
    for s in sections:
        _add_heading(doc, s["title"], s.get("level"))
        content = (s.get("content") or "").strip()
        if content:
            _render_content(doc, content, table_style, bullet_ok)
    doc.save(dest_path)
    return dest_path
