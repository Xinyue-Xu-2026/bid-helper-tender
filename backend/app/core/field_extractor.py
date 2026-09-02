"""资料字段抽取：Excel 映射解析 + 正则兜底 + OCR + 可插拔 LLM 钩子。

数据源优先级：Excel 汇总表 > 文字型 PDF/Word > OCR（扫描件/图片）> LLM（可选增强）。
OCR 依赖 rapidocr_onnxruntime，缺失时优雅降级（ocr_available() 为 False），不抛异常。
"""
import json
import re
from datetime import date, datetime
from typing import Callable, List, Optional, Tuple

from openpyxl import load_workbook

from app.core.extractor import extract_text
from app.services.asset_service import normalize_date

# ---------- OCR（可选依赖，导入失败降级） ----------
try:
    from rapidocr_onnxruntime import RapidOCR as _RapidOCR
    OCR_AVAILABLE = True
    OCR_UNAVAILABLE_REASON = ""
except Exception as exc:  # ImportError 或 onnxruntime 环境异常
    _RapidOCR = None
    OCR_AVAILABLE = False
    OCR_UNAVAILABLE_REASON = f"OCR 不可用（rapidocr_onnxruntime 未安装或加载失败：{exc}）"

_ocr_engine = None


def ocr_available() -> bool:
    return OCR_AVAILABLE


def _get_ocr():
    global _ocr_engine
    if _ocr_engine is None:
        _ocr_engine = _RapidOCR()
    return _ocr_engine


def _ocr_lines(result) -> List[str]:
    return [str(line[1]).strip() for line in (result or []) if len(line) >= 2 and str(line[1]).strip()]


def ocr_image(path: str) -> str:
    """图片文件 OCR，返回拼接文本。"""
    if not OCR_AVAILABLE:
        raise RuntimeError(OCR_UNAVAILABLE_REASON)
    return "\n".join(_ocr_lines(_get_ocr()(str(path))[0]))


def ocr_pdf_pages(path: str, dpi: int = 200) -> str:
    """扫描件 PDF：PyMuPDF 渲染成图后走 RapidOCR。"""
    if not OCR_AVAILABLE:
        raise RuntimeError(OCR_UNAVAILABLE_REASON)
    import fitz
    engine = _get_ocr()
    texts: List[str] = []
    with fitz.open(path) as doc:
        for page in doc:
            pix = page.get_pixmap(dpi=dpi)
            texts.extend(_ocr_lines(engine(pix.tobytes("png"))[0]))
    return "\n".join(texts)


# ---------- Excel 映射解析 ----------

def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return normalize_date(value)
    return str(value).strip()


def _match_columns(headers: List[str], mapping: dict) -> dict:
    """{列号: 目标字段}，按映射表的源列名精确匹配表头。"""
    return {idx: mapping[h] for idx, h in enumerate(headers) if h in mapping}


def _sheet_kind(headers: List[str], person_mapping: dict, contract_mapping: dict) -> Optional[str]:
    """判定 sheet 类型：匹配到映射字段多者胜；平手时特征字段（姓名/项目名称）优先。"""
    p_targets = {_match_columns(headers, person_mapping)[i] for i in _match_columns(headers, person_mapping)}
    c_targets = {_match_columns(headers, contract_mapping)[i] for i in _match_columns(headers, contract_mapping)}
    if not p_targets and not c_targets:
        return None
    if len(p_targets) > len(c_targets):
        return "person"
    if len(c_targets) > len(p_targets):
        return "contract"
    if "姓名" in p_targets and "项目名称" not in c_targets:
        return "person"
    if "项目名称" in c_targets and "姓名" not in p_targets:
        return "contract"
    return "person" if "姓名" in p_targets else "contract"


def parse_excel(file_path: str, person_mapping: dict,
                contract_mapping: dict) -> Tuple[List[dict], List[dict], List[str]]:
    """逐 sheet 按映射解析。返回 (人员行列表, 业绩行列表, 警告列表)。"""
    if file_path.lower().endswith(".xls"):
        raise ValueError("暂不支持 .xls 旧格式，请另存为 .xlsx 后重试")
    persons: List[dict] = []
    contracts: List[dict] = []
    warnings: List[str] = []
    wb = load_workbook(file_path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            try:
                header_row = next(rows)
            except StopIteration:
                continue
            headers = [str(h).strip() if h is not None else "" for h in header_row]
            kind = _sheet_kind(headers, person_mapping, contract_mapping)
            if kind is None:
                if any(headers):
                    warnings.append(f"工作表「{ws.title}」表头未匹配到映射列，已跳过")
                continue
            mapping = person_mapping if kind == "person" else contract_mapping
            col_index = _match_columns(headers, mapping)
            for row in rows:
                if not any(v is not None and str(v).strip() for v in row):
                    continue
                fields = {}
                for idx, target in col_index.items():
                    value = row[idx] if idx < len(row) else None
                    text = _cell_str(value)
                    if text:
                        fields[target] = text
                if kind == "person":
                    if fields.get("姓名"):
                        persons.append(fields)
                else:
                    if fields.get("项目名称") or fields.get("项目经理"):
                        contracts.append(fields)
    finally:
        wb.close()
    return persons, contracts, warnings


# ---------- 文本正则兜底抽取 ----------

CERT_KEYWORDS = (
    "一级注册建造师", "二级注册建造师", "注册建造师",
    "一级造价工程师", "二级造价工程师", "造价工程师",
    "注册监理工程师", "监理工程师", "咨询工程师", "注册安全工程师",
    "安全生产考核合格证", "安全生产考核", "职称证书", "执业资格证书",
    "岗位证书", "特种作业操作证",
)

CREDIT_KEYWORDS = ("企业资质", "建筑业企业资质", "营业执照", "安全生产许可证",
                   "质量管理体系认证", "环境管理体系认证", "职业健康安全管理体系认证")

_DATE_RE = r"(\d{4}\s*[年./\-]\s*\d{1,2}\s*[月./\-]\s*\d{1,2}\s*日?|\d{4}\s*年\s*\d{1,2}\s*月)"


def _search(pattern: str, text: str) -> str:
    m = re.search(pattern, text)
    return m.group(1).strip() if m else ""


def extract_fields_from_text(text: str) -> Tuple[dict, List[str]]:
    """从证书/资料文本中正则抽取关键字段；返回 (fields, warnings)。"""
    fields: dict = {}
    warnings: List[str] = []

    name = _search(r"姓\s*名[：:\s]*([一-龥]{2,4})", text)
    if name:
        fields["姓名"] = name
    id_no = _search(r"(?<!\d)(\d{17}[\dXx])(?!\d)", text)
    if id_no:
        fields["身份证号"] = id_no
    title = _search(r"职\s*称[：:\s]*([一-龥（）()a-zA-Z]{2,15})", text)
    if title:
        fields["职称"] = title
    phone = _search(r"(?<!\d)(1[3-9]\d{9})(?!\d)", text)
    if phone:
        fields["联系方式"] = phone

    cert = ""
    for kw in CERT_KEYWORDS:
        if kw in text:
            cert = kw
            break
    if not cert:
        cert = _search(r"证书名称[：:\s]*([^\n，。；]{2,30})", text)
    if cert:
        fields["证书名称"] = cert

    expiry_raw = _search(r"有效期(?:至|截止|止)?[：:\s]*" + _DATE_RE, text)
    if expiry_raw:
        normalized = normalize_date(expiry_raw.replace("年", "-").replace("月", "-").replace("日", ""))
        if normalized:
            fields["证书有效期至"] = normalized

    issuer = _search(r"发证机关[：:\s]*([^\n，。；]{2,40})", text)
    if issuer:
        fields["发证机关"] = issuer
    issue_date = _search(r"发证日期[：:\s]*" + _DATE_RE, text)
    if issue_date:
        normalized = normalize_date(issue_date.replace("年", "-").replace("月", "-").replace("日", ""))
        if normalized:
            fields["发证日期"] = normalized

    if not fields.get("姓名") and not fields.get("证书名称"):
        warnings.append("未能从文本中识别出姓名或证书名称，请人工核对")
    if not fields.get("证书有效期至"):
        warnings.append("未识别到证书有效期")
    return fields, warnings


def classify_text(text: str, fields: dict) -> str:
    """判定资料类型：person（人员证书）或 credit（企业资质）。"""
    if fields.get("姓名") or fields.get("身份证号"):
        return "person"
    if any(kw in text for kw in CREDIT_KEYWORDS):
        return "credit"
    if re.search(r"企业名称|统一社会信用代码", text):
        return "credit"
    return "person"


# ---------- 可插拔 LLM 结构化抽取（默认关闭，settings 开启） ----------

LLM_EXTRACT_PROMPT = """你是投标资料结构化抽取专家。从用户提供的证书/资料文本中抽取信息，只输出 JSON 对象：
{"asset_type": "person 或 credit", "fields": {"姓名": "", "证书名称": "", "职称": "", "身份证号": "", "联系方式": "", "证书有效期至": "YYYY-MM-DD", "发证机关": "", "发证日期": "YYYY-MM-DD"}}
person 表示人员证书，credit 表示企业资质证书。没有的字段留空字符串，不要输出任何其他文字。"""


def extract_with_llm(text: str, api_key: str = "", model: str = "") -> Optional[dict]:
    """Kimi LLM 结构化抽取钩子；任何失败返回 None（调用方回退正则结果）。

    调用方式参考 core/llm_parser.py（编程订阅 Key 走编程网关）。
    """
    from app import settings_store
    from app.core import llm_parser

    api_key = api_key or settings_store.get_api_key()
    if not api_key:
        return None
    model = model or settings_store.get_model()
    try:
        client = llm_parser._make_client(api_key)
        completion = client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system", "content": LLM_EXTRACT_PROMPT},
                {"role": "user", "content": text[:8000]},
            ],
            response_format={"type": "json_object"},
            max_tokens=2048,
        )
        if not completion.choices:
            return None
        data = json.loads(completion.choices[0].message.content)
    except Exception:
        return None
    if not isinstance(data, dict) or not isinstance(data.get("fields"), dict):
        return None
    fields = {str(k): str(v).strip() for k, v in data["fields"].items()
              if v is not None and str(v).strip()}
    asset_type = data.get("asset_type") if data.get("asset_type") in ("person", "credit") else None
    return {"asset_type": asset_type, "fields": fields}


# ---------- 文本文件入口 ----------

def extract_text_auto(path: str, suffix: str) -> Tuple[str, str]:
    """返回 (text, source)。PDF 先提文字，提不到再 OCR；图片直接 OCR。

    OCR 不可用且必须 OCR 时抛 RuntimeError（由上层转成 errors 条目）。
    """
    if suffix == ".pdf":
        text = extract_text(path)
        if len(text.strip()) >= 20:
            return text, "pdf"
        return ocr_pdf_pages(path), "ocr"
    if suffix in (".docx", ".doc"):
        return extract_text(path), "pdf"
    if suffix in (".png", ".jpg", ".jpeg"):
        return ocr_image(path), "ocr"
    raise ValueError(f"不支持的文件类型：{suffix}")
