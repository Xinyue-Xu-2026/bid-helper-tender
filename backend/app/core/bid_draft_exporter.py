"""底稿驱动商务标导出：按用户确认的 bindings 填充数据表/简历表/扫描件，
替换残留项目编号/名称/日期，可选目录换 TOC 域，输出填充报告 + 防篡改校验。

bindings JSON schema 权威定义见 bid_table_classifier 模块 docstring；
data 结构见 bid_service.assemble_bid_draft_data
（{"persons":[{name,fields,is_lead,role,sem}],"contracts":[{name,fields,
section,sem}],"lead_perfs_text":str}）。
"""
import os
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.table import Table

from app.core.bid_template_exporter import (
    _clear_table_images, _insert_image_into_table, _is_toc_paragraph,
    _para_text, _replace_stale_text, _rewrite_paragraph_text, _set_cell_text,
    insert_image_adaptive, page_text_width_cm,
)
from app.core.bid_page_setup import apply_page_setup
from app.core.bid_verify import _is_section_break_para, verify_draft_fill
from app.core.word_exporter import insert_toc_field_at


def _clear_tr_text(tr) -> None:
    """清空一行（deepcopy 出的 tr）全部单元格文本，保留单元格/段落属性。"""
    for tc in tr.iter(qn("w:tc")):
        for para in tc.findall(qn("w:p")):
            for child in list(para):
                if child.tag != qn("w:pPr"):
                    para.remove(child)


def _strip_tr_vmerge(tr) -> None:
    """剔除克隆行内全部 vMerge 元素：donor 克隆行作为数据行，
    不应启动/延续纵向合并区（表头行的 vMerge 不随克隆带入）。"""
    for vm in list(tr.iter(qn("w:vMerge"))):
        vm.getparent().remove(vm)


def _fill_table_semantic(table, columns: dict, sem_rows: list,
                         header_rows: int = 1) -> int:
    """语义列填充数据表：保留前 header_rows 个表头行，donor 取
    rows[header_rows]（首个真数据行——空白数据行正好携带正确
    vAlign/缩进/行高体例）deepcopy 逐行写入；无数据行（总行数 <=
    header_rows）时 donor = 清空文本的最后一个表头行克隆（保留单元格
    属性）。donor 克隆行剔除 vMerge。列下标 j → 反向查 columns 得语义键，
    有则写 sem_row.get(key,"")，无则清 ""；横向合并按 _tc 引用去重只写
    一次。sem_rows 空 → 仅留表头返回 0。返回写入行数。"""
    col_to_sem = {col: sem for sem, col in (columns or {}).items()}
    rows = table.rows
    if not rows:
        return 0
    header_rows = max(1, min(int(header_rows or 1), len(rows)))
    if len(rows) > header_rows:
        donor_tr = deepcopy(rows[header_rows]._tr)
    else:
        donor_tr = deepcopy(rows[header_rows - 1]._tr)
        _clear_tr_text(donor_tr)
    _strip_tr_vmerge(donor_tr)
    for r in list(table.rows[header_rows:]):
        r._tr.getparent().remove(r._tr)
    written = 0
    for sem_row in sem_rows or []:
        table._tbl.append(deepcopy(donor_tr))
        new_row = table.rows[-1]
        seen_tcs = []
        for j, cell in enumerate(new_row.cells):
            tc = cell._tc
            if any(tc is t for t in seen_tcs):
                continue
            seen_tcs.append(tc)
            sem = col_to_sem.get(j)
            _set_cell_text(cell, sem_row.get(sem, "") if sem else "")
        written += 1
    return written


def _clone_table_after(table):
    """XML 级 deepcopy 整表并插到该表之后（后随空段落分隔），完整保留
    行列数/合并单元格/行高/样式（一人一表克隆用，绝不重建结构）。"""
    tbl = table._tbl
    new = deepcopy(tbl)
    sep = tbl.makeelement(qn("w:p"), {})
    tbl.addnext(sep)
    tbl.addnext(new)
    return Table(new, table._parent)


def _fill_resume_table(table, label_to_sem: dict, sem: dict,
                       perfs_text: str) -> int:
    """简历键值表填充：逐行逐格，规范化格文本（去首尾空白 + 去尾部
    "："或":"）命中 label_to_sem → 写其右侧首个不同 _tc 的格子（无右格跳过）；
    语义键为 lead_perfs 时写 perfs_text（随人传入）。返回写入格数。"""
    written = 0
    for row in table.rows:
        cells = row.cells
        seen_tcs = []
        for j, cell in enumerate(cells):
            tc = cell._tc
            if any(tc is t for t in seen_tcs):
                continue
            seen_tcs.append(tc)
            label = (cell.text or "").strip().rstrip("：:").strip()
            sem_key = (label_to_sem or {}).get(label)
            if not sem_key:
                continue
            target = None
            for nxt in cells[j + 1:]:
                if nxt._tc is not tc:
                    target = nxt
                    break
            if target is None:
                continue
            value = (perfs_text if sem_key == "lead_perfs"
                     else (sem or {}).get(sem_key, ""))
            _set_cell_text(target, value or "")
            written += 1
    return written


def _fill_image_slots(doc, slot_bindings, persons, report) -> None:
    """逐图片占位表插图。slot_bindings: [(binding, table)]。
    person "lead" → 负责人；"member:N" → 第 N 个非负责人（0 起，
    越界 → report["skipped"] 记 reason="人员不足"）；缺省按负责人。
    图片路径按 label_kind 取 fields：职称证书 → 职称证书扫描件；
    社保 → 社保缴纳证明扫描件；身份证 → 身份证扫描件；
    注册证书 → 按该 person 出现顺序逐次取 fields["证书"][k]["扫描件"]
    （每人独立计数器）。先 _clear_table_images 再 _insert_image_into_table
    （文件不存在内部跳过）。成功/失败记入 report["images"]。"""
    persons = persons or []
    lead = next((p for p in persons if p.get("is_lead")), None)
    members = [p for p in persons if not p.get("is_lead")]
    cert_counters = {}  # id(person) → 已取注册证书数（每人独立）

    for binding, table in slot_bindings:
        who = binding.get("person") or "lead"
        if who == "lead":
            person = lead
        elif who.startswith("member:"):
            try:
                n = int(who.split(":", 1)[1])
            except ValueError:
                n = -1
            person = members[n] if 0 <= n < len(members) else None
        else:
            person = lead
        if person is None:
            report["skipped"].append({
                "table_index": binding.get("table_index"),
                "role": "image_slot", "reason": "人员不足"})
            continue
        fields = person.get("fields") or {}
        kind = binding.get("label_kind") or ""
        if kind == "职称证书":
            image_path = str(fields.get("职称证书扫描件") or "")
        elif kind == "社保":
            image_path = str(fields.get("社保缴纳证明扫描件") or "")
        elif kind == "身份证":
            image_path = str(fields.get("身份证扫描件") or "")
        elif kind == "注册证书":
            certs = [c for c in (fields.get("证书") or [])
                     if isinstance(c, dict)]
            key = id(person)
            k = cert_counters.get(key, 0)
            cert_counters[key] = k + 1
            image_path = (str(certs[k].get("扫描件") or "")
                          if k < len(certs) else "")
        else:
            image_path = ""
        _clear_table_images(table)
        # 等比自适应插入；失败（文件缺失/格式不支持/插入异常）显式记入报告
        _insert_image_into_table(
            table, image_path, report=report,
            table_index=binding.get("table_index"),
            person=str(person.get("name") or ""), label_kind=kind)


def _swap_toc(doc) -> bool:
    """首个连续 toc 样式段落组 → 保留首段位置、删除组内其余段，
    在首段 insert_toc_field_at；无 toc 段 → False。"""
    group = []
    for para in doc.paragraphs:
        if _is_toc_paragraph(para):
            group.append(para)
        elif group:
            break
    if not group:
        return False
    first = group[0]
    for para in group[1:]:
        para._element.getparent().remove(para._element)
    insert_toc_field_at(first)
    return True


def _scope_persons(persons, scope: str) -> list:
    if scope == "lead":
        return [p for p in persons if p.get("is_lead")]
    if scope == "members":
        return [p for p in persons if not p.get("is_lead")]
    return list(persons)


# ---------- 授权页正文填充（法定代表人/委托代理人） ----------
_ID_CARD_WIDTH = Cm(7)


def _date_zh(doc_date: str) -> str:
    """'YYYY-MM-DD' → 'YYYY年M月D日'；空/非法返回 ''。"""
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", (doc_date or "").strip())
    if not m:
        return ""
    return f"{int(m.group(1))}年{int(m.group(2))}月{int(m.group(3))}日"


def _auth_signature_block(doc, bound, report, auth) -> None:
    """授权页正文填充：处理「法定代表人（单位负责人）身份证明」与
    「授权委托书」两页的正文占位与签名区。只命中文本特征明确的段落：
    - （姓名）/（投标人名称）括号占位（姓名按序=法人、代理人）；
    - 「身份证号：」行（紧跟在 法人/代理人签名行之后）→ 填对应证件号；
    - 盖公章行后的「年月日」行 → 填投标文件日期；
    - 「附：…身份证…电子扫描件」段末追加身份证正反面图片。
    其余段落（承诺函等）一律不动。改动段记入 bound（verify 同坐标系
    段下标：toc/分节符/空段过滤后的「非空段」序列），图片追加不计文本
    故不需 bound。"""
    legal = auth.get("legal_rep") or {}
    agent = auth.get("agent") or {}
    legal_name = str(legal.get("name") or "").strip()
    agent_name = str(agent.get("name") or "").strip()
    legal_id = str(legal.get("身份证号") or "").strip()
    agent_id = str(agent.get("身份证号") or "").strip()
    bidder = str(auth.get("bidder_name") or "").strip()
    date_text = _date_zh(auth.get("doc_date") or "")

    def _img_paths(person, dual):
        if dual:
            # 委托双方：法人+代理人各正反面
            out = []
            for p in (legal, agent):
                if p:
                    out += [str(p.get("身份证正面扫描件") or ""),
                            str(p.get("身份证反面扫描件") or "")]
            return out
        if not person:
            return []
        return [str(person.get("身份证正面扫描件") or ""),
                str(person.get("身份证反面扫描件") or "")]

    def _append_images(para, person, dual) -> int:
        inserted = 0
        for pth in _img_paths(person, dual):
            if not pth:
                continue
            # 身份证扫描件：自适应（页宽为可用宽，_ID_CARD_WIDTH 为退化宽度）；
            # 失败显式记入 report["images"]，不再静默吞异常
            result = insert_image_adaptive(
                para, pth, page_text_width_cm(doc),
                fallback_w_cm=_ID_CARD_WIDTH.cm)
            report["images"].append({
                "table_index": None,
                "person": str((person or {}).get("name") or ""),
                "label_kind": "身份证", "ok": result["ok"],
                "reason": result["reason"], "too_long": result["too_long"]})
            if result["ok"]:
                inserted += 1
        return inserted

    # 与 verify 段落比对同坐标系：toc/分节符/空段过滤后的「非空段」序列
    # （空段不携带内容；克隆表分隔空段等结构插入不影响本坐标系）
    paras = [p for p in doc.paragraphs
             if not _is_toc_paragraph(p) and not _is_section_break_para(p)
             and _para_text(p)]

    def _filter_index(p) -> int:
        return paras.index(p)

    def _rewrite(para, text):
        _rewrite_paragraph_text(para, text)
        bound.add(_filter_index(para))

    sig_role = None        # 'legal'/'agent'：最近签名行（其后一行身份证号归属）
    prev_is_sig_ctx = False  # 上一段是盖公章/签名上下文（其后年月日归属）

    for para in paras:
        t = (para.text or "").strip()
        if not t:
            prev_is_sig_ctx = False
            continue

        # 1) 括号占位：本人（姓名）→法人名；现委托（姓名）→代理人名；
        #    （投标人名称）→投标人名称（两页通用）
        nt = t
        if "（姓名）" in nt:
            if legal_name:
                nt = nt.replace("本人（姓名）", f"本人{legal_name}", 1)
            if agent_name:
                nt = nt.replace("现委托（姓名）", f"现委托{agent_name}", 1)
        if "（投标人名称）" in nt and bidder:
            nt = nt.replace("（投标人名称）", bidder)
        if nt != t:
            _rewrite(para, nt)
            t = nt

        # 2) 签名角色跟踪
        if "（签名）" in t:
            if "法定代表人" in t or "单位负责人" in t:
                sig_role = "legal"
            elif "委托代理人" in t:
                sig_role = "agent"
            prev_is_sig_ctx = True
        elif "盖单位公章" in t:
            prev_is_sig_ctx = True
            sig_role = None
        # 3) 身份证号行 → 填当前签名角色证件号
        elif sig_role and re.match(r"^身份证号[：:]\s*$", t):
            person = legal if sig_role == "legal" else agent
            pid = legal_id if sig_role == "legal" else agent_id
            if pid:
                _rewrite(para, f"身份证号：{pid}")
                prev_is_sig_ctx = True
            sig_role = None
        # 4) 年月日行（盖公章/签名上下文之后）→ 日期
        elif prev_is_sig_ctx and date_text and re.fullmatch(r"年\s*月\s*日", t):
            _rewrite(para, date_text)
            prev_is_sig_ctx = False
        else:
            prev_is_sig_ctx = False

        # 5) 附：…身份证…电子扫描件 → 追图
        if "附：" in t and "身份证" in t and "扫描件" in t:
            dual = "双方" in t or "委托双方" in t
            person = legal if not dual else None
            n = _append_images(para, person, dual)
            if n:
                report["authority_images"] += n
        # 状态续到下一行判断
        # prev_is_sig_ctx 由上面各分支设置，正常流转
        if sig_role == "legal" or sig_role == "agent":
            prev_is_sig_ctx = True


def _scope_contracts(contracts, persons, scope: str) -> list:
    if scope == "lead":
        lead = next((p for p in persons if p.get("is_lead")), None)
        lead_name = str(lead.get("name") or "").strip() if lead else ""
        if not lead_name:
            return []
        return [c for c in contracts
                if str((c.get("fields") or {}).get("项目负责人")
                       or "").strip() == lead_name]
    if scope == "section1":
        return [c for c in contracts if c.get("section") == 1]
    if scope == "section2":
        return [c for c in contracts if c.get("section") == 2]
    return list(contracts)


def fill_draft(draft_path: str, dest_path: str, bindings: dict, data: dict,
               project_no: str = "", project_name: str = "",
               doc_date: str = "", tenderer: str = "",
               bidder_name: str = "", auth: dict = None,
               section_name: str = "", section_no: str = "") -> dict:
    """按确认的 bindings 填充底稿并另存 dest_path，返回 fill_report。
    dest 与 draft 不得同路径。tenderer/bidder_name 为空则跳过对应
    标签空白（招标人：____ / 投标人名称：____）填充；section_name/
    section_no 为空则跳过标段标签/括号占位（标段名称：____ /
    （标段名称）等）填充。
    auth：可选，{"legal_rep": {姓名/身份证号/正反面扫描件路径...} | {},
    "agent": {...}, "doc_date": str}——授权页（法定代表人身份证明/授权
    委托书）签名区与身份证附图填充；未提供或字段为空则跳过。"""
    if os.path.abspath(str(draft_path)) == os.path.abspath(str(dest_path)):
        raise ValueError("dest 与 draft 不得同路径")
    doc = Document(draft_path)
    # doc.tables：顶层表格、文档顺序（python-docx 按 body 中 w:tbl 顺序收集，
    # 已验证与块遍历一致），与 T5 table_index 同坐标系。
    tables = doc.tables
    data = data or {}
    persons = data.get("persons") or []
    contracts = data.get("contracts") or []
    report = {"filled": [], "images": [], "skipped": [], "verify": None,
              "page_setup": None, "authority_images": 0, "auth_bound": []}

    # 残留文本替换须在填充之前执行：替换规则（compute_text_subs）基于
    # 填充前的底稿计算，与 verify_draft_fill 的底稿侧同源——否则 stale 值
    # 仅存在于绑定表数据单元格时，填充覆写后规则发现不到旧值，
    # 封面/正文残留不被替换且 verify 误报。填充整体覆写数据单元格，
    # 提前替换不改变产物内容。
    _replace_stale_text(doc, project_no=project_no,
                        project_name=project_name, doc_date=doc_date,
                        tenderer=tenderer, bidder_name=bidder_name,
                        section_name=section_name, section_no=section_no)

    # 授权页正文填充：改动的是普通正文段，须把被改段加入 verify 的
    # bound（paragraph 下标，toc/分节符过滤坐标系）避免误报
    auth_bound = set()
    if auth:
        auth = {**auth, "bidder_name": bidder_name}
        _auth_signature_block(doc, auth_bound, report, auth)
        report["auth_bound"] = sorted(auth_bound)

    image_slots = []
    bound_indices = set()
    table_insertions = {}   # {底稿表下标: 插入克隆数}（resume_each 一人一表，
                            # 填充端与校验端共用同一登记）
    for binding in (bindings or {}).get("tables") or []:
        if not binding.get("confirmed"):
            continue
        role = binding.get("role") or "ignore"
        if role == "ignore":
            continue
        idx = binding.get("table_index")
        if not isinstance(idx, int) or not 0 <= idx < len(tables):
            report["skipped"].append({
                "table_index": idx, "role": role,
                "reason": "表不存在，可能底稿已重新生成"})
            continue
        # 仅 fill 实际会改动的表豁免防篡改校验；quote（手工填写）与未知
        # 角色不改动 → 不豁免，未绑定表同等受校验保护
        if role in ("person_roster", "perf_list", "lead_resume",
                    "resume_each", "image_slot"):
            bound_indices.add(idx)
        table = tables[idx]
        if role == "person_roster":
            picked = _scope_persons(persons, binding.get("person_scope") or "all")
            rows = _fill_table_semantic(table, binding.get("columns") or {},
                                        [p.get("sem") or {} for p in picked],
                                        header_rows=binding.get("header_rows") or 1)
            report["filled"].append(
                {"table_index": idx, "role": role, "rows": rows})
        elif role == "perf_list":
            picked = _scope_contracts(contracts, persons,
                                      binding.get("perf_scope") or "all")
            rows = _fill_table_semantic(table, binding.get("columns") or {},
                                        [c.get("sem") or {} for c in picked],
                                        header_rows=binding.get("header_rows") or 1)
            report["filled"].append(
                {"table_index": idx, "role": role, "rows": rows})
        elif role == "lead_resume":
            lead = next((p for p in persons if p.get("is_lead")), None)
            written = 0
            if lead is not None:
                written = _fill_resume_table(
                    table, binding.get("columns") or {},
                    lead.get("sem") or {},
                    str(lead.get("perfs_text")
                        or data.get("lead_perfs_text") or ""))
            report["filled"].append(
                {"table_index": idx, "role": role, "rows": written})
        elif role == "resume_each":
            if binding.get("mode") == "per_person":
                # 一人一表（V1.2 7.4）：第 1 份填原样表，其余先克隆再填；
                # 克隆全部基于填充前的样表，链式 addnext 保持人员顺序
                picked = _scope_persons(persons,
                                        binding.get("person_scope") or "all")
                targets = [table]
                for _ in range(max(0, len(picked) - 1)):
                    targets.append(_clone_table_after(targets[-1]))
                total = 0
                for t, p in zip(targets, picked):
                    perfs = str(p.get("perfs_text") or (
                        data.get("lead_perfs_text")
                        if p.get("is_lead") else "") or "")
                    total += _fill_resume_table(
                        t, binding.get("columns") or {},
                        p.get("sem") or {}, perfs)
                if len(picked) > 1:
                    table_insertions[idx] = len(picked) - 1
                report["filled"].append(
                    {"table_index": idx, "role": role, "rows": total,
                     "cloned": len(picked)})
            else:
                # 未开 per_person：等同 lead_resume（单份 lead 填充，旧行为）
                lead = next((p for p in persons if p.get("is_lead")), None)
                written = 0
                if lead is not None:
                    written = _fill_resume_table(
                        table, binding.get("columns") or {},
                        lead.get("sem") or {},
                        str(lead.get("perfs_text")
                            or data.get("lead_perfs_text") or ""))
                report["filled"].append(
                    {"table_index": idx, "role": role, "rows": written})
        elif role == "quote":
            report["skipped"].append(
                {"table_index": idx, "role": role, "reason": "报价表需手工填写"})
        elif role == "image_slot":
            image_slots.append((binding, table))

    _fill_image_slots(doc, image_slots, persons, report)
    swap_toc = bool((bindings or {}).get("swap_toc"))
    if swap_toc:
        _swap_toc(doc)
    doc.save(str(dest_path))
    # 页眉移植 + 分节页码（参考缺失自动降级，见 bid_page_setup）
    report["page_setup"] = apply_page_setup(str(dest_path))
    report["verify"] = verify_draft_fill(
        str(draft_path), str(dest_path),
        bound_table_indices=bound_indices,
        replace_params={"project_no": project_no, "project_name": project_name,
                        "doc_date": doc_date, "tenderer": tenderer,
                        "bidder_name": bidder_name,
                        "section_name": section_name,
                        "section_no": section_no},
        swapped_toc=swap_toc,
        bound_paragraph_indices=auth_bound,
        table_insertions=table_insertions)
    return report
