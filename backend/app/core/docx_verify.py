# -*- coding: utf-8 -*-
"""校验输出 docx 是否保持模板格式（移植 skill verify_output 逻辑，返回结构化结果）。

校验规则：
1. 未改动段落（text 在模板中出现过）格式必须与模板逐项一致，否则记 drift。
2. 新增段落（text 不在模板中）的 style 必须存在于模板同样式格式集合，且格式
   匹配其一（因为新增段落应是模板段落的克隆），否则记 invalid_new。
"""
from collections import defaultdict

from docx import Document

FMT_KEYS = ('style', 'font', 'size', 'bold', 'line_spacing', 'first_indent')


def _fmt(p):
    r = p.runs[0] if p.runs else None
    pf = p.paragraph_format
    return (
        p.style.name if p.style else None,
        r.font.name if r else None,
        str(r.font.size) if r else None,
        r.font.bold if r else None,
        str(pf.line_spacing),
        str(pf.first_line_indent),
    )


def verify_output(tpl_path: str, out_path: str) -> dict:
    tpl = Document(tpl_path)
    out = Document(out_path)

    tpl_fmt_by_text = {}
    tpl_fmts_by_style = defaultdict(set)
    for p in tpl.paragraphs:
        if not (p.style and p.text.strip()):
            continue
        f = _fmt(p)
        tpl_fmt_by_text.setdefault(p.text, f)
        tpl_fmts_by_style[p.style.name].add(f)

    drift = 0
    invalid_new = 0
    new_count = 0
    issues = []
    for i, p in enumerate(out.paragraphs):
        if not (p.style and p.text.strip()):
            continue
        f = _fmt(p)
        if p.text in tpl_fmt_by_text:
            tf = tpl_fmt_by_text[p.text]
            if f != tf:
                diff = {k: (a, b) for k, a, b in zip(FMT_KEYS, tf, f) if a != b}
                issues.append(f"段落[{i}]“{p.text[:30]}”格式漂移：{diff}")
                drift += 1
        else:
            new_count += 1
            if p.style.name not in tpl_fmts_by_style:
                issues.append(f"段落[{i}]“{p.text[:30]}”使用了模板中不存在的样式 <{p.style.name}>")
                invalid_new += 1
            elif f not in tpl_fmts_by_style[p.style.name]:
                issues.append(f"段落[{i}]“{p.text[:30]}”格式与模板同样式不符：{dict(zip(FMT_KEYS, f))}")
                invalid_new += 1

    outline = []
    for p in out.paragraphs:
        name = p.style.name if p.style and p.style.name else ''
        if name.startswith('Heading') and p.text.strip():
            outline.append(p.text)

    return {
        "drift": drift,
        "invalid_new": invalid_new,
        "new_count": new_count,
        "tpl_paragraphs": len(tpl.paragraphs),
        "out_paragraphs": len(out.paragraphs),
        "outline": outline,
        "issues": issues,
    }
