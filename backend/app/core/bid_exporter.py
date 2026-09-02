"""商务标导出：Excel / Word 两种格式内容一致（设计文档 §6）。

- 人员配备表：序号 / 姓名 / 拟派岗位 / 职称 / 证书明细 / 有效期状态；
- 企业业绩表：序号 / 项目名称 + 字段配置（settings_store field_config.contract）驱动列；
  未配置（空列表）时默认列：类型/项目经理/合同金额/年份。
- Word 从零新建 Document()，不打开任何模板，规避格式变形问题。
"""
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app import settings_store
from app.services.bid_service import _cert_expiry

PERSON_HEADERS = ["序号", "姓名", "拟派岗位", "职称", "证书明细", "有效期状态"]
PERSON_WIDTHS = [6, 12, 16, 14, 42, 30]
CONTRACT_BASE_HEADERS = ["序号", "项目名称"]
CONTRACT_BASE_WIDTHS = [6, 30]
DEFAULT_CONTRACT_COLUMNS = ["类型", "项目经理", "合同金额", "年份"]

_HEADER_FONT = Font(bold=True)
_WRAP_TOP = Alignment(wrap_text=True, vertical="top")


def _contract_columns() -> list:
    """企业业绩表的配置驱动列；剔除"项目名称"（已固定为第二列），空配置用默认列。"""
    cfg = settings_store.get_field_config().get("contract") or []
    keys = [f["key"] for f in cfg
            if isinstance(f, dict) and f.get("key") and f["key"] != "项目名称"]
    return keys or DEFAULT_CONTRACT_COLUMNS


def _cert_detail(fields: dict) -> str:
    """证书明细：每本 "类型·专业（有效期至YYYY-MM-DD）"（无专业时 "类型（有效期至…）"），
    多本换行拼接。"""
    parts = []
    for cert in (fields or {}).get("证书") or []:
        if not isinstance(cert, dict):
            continue
        label = cert.get("类型") or "证书"
        major = cert.get("专业")
        if major:
            label = f"{label}·{major}"
        parts.append(f"{label}（有效期至{_cert_expiry(cert)}）")
    return "\n".join(parts)


def _person_rows(persons: list) -> list:
    rows = []
    for i, p in enumerate(persons, start=1):
        fields = p.get("fields") or {}
        warnings = p.get("cert_warnings") or []
        status = "\n".join(w["message"] for w in warnings) if warnings else "正常"
        rows.append([i, p.get("name", ""), p.get("role", ""),
                     str(fields.get("职称", "") or ""),
                     _cert_detail(fields), status])
    return rows


def _contract_rows(contracts: list, columns: list) -> list:
    rows = []
    for i, c in enumerate(contracts, start=1):
        fields = c.get("fields") or {}
        rows.append([i, c.get("name", "")] +
                    [str(fields.get(k, "") or "") for k in columns])
    return rows


def _write_sheet(ws, headers: list, rows: list, widths: list, wrap_cols: tuple):
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        ws.cell(row=1, column=col_idx).font = _HEADER_FONT
    for r_idx, row in enumerate(rows, start=2):
        ws.append(row)
        for col_idx in wrap_cols:
            ws.cell(row=r_idx, column=col_idx).alignment = _WRAP_TOP
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"


def build_bid_xlsx(project_name: str, persons: list, contracts: list, dest_path: str) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "人员配备"
    _write_sheet(ws, PERSON_HEADERS, _person_rows(persons), PERSON_WIDTHS,
                 wrap_cols=(5, 6))
    columns = _contract_columns()
    ws2 = wb.create_sheet("企业业绩")
    _write_sheet(ws2, CONTRACT_BASE_HEADERS + columns,
                 _contract_rows(contracts, columns),
                 CONTRACT_BASE_WIDTHS + [16] * len(columns),
                 wrap_cols=())
    dest = Path(dest_path)
    wb.save(dest)
    return dest


def _add_table(doc, headers: list, rows: list):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for ci, header in enumerate(headers):
        cell = table.cell(0, ci)
        cell.text = str(header)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            table.cell(ri, ci).text = str(val)


def build_bid_docx(project_name: str, persons: list, contracts: list, dest_path: str) -> Path:
    doc = Document()  # 从零新建，不打开任何模板
    doc.add_heading(f"{project_name} 商务标", level=1)
    doc.add_heading("一、人员配备表", level=2)
    _add_table(doc, PERSON_HEADERS, _person_rows(persons))
    doc.add_heading("二、企业业绩表", level=2)
    columns = _contract_columns()
    _add_table(doc, CONTRACT_BASE_HEADERS + columns,
               _contract_rows(contracts, columns))
    dest = Path(dest_path)
    doc.save(dest)
    return dest
