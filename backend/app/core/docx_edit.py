# -*- coding: utf-8 -*-
"""Format-preserving docx edit helpers.

Core idea: never build paragraphs from scratch. Always deepcopy an existing
template paragraph's XML so style, font, spacing, and indent come along for
free; then only replace the text of the first run (which keeps run-level
formatting). Never touch styles.xml, sections, headers/footers, or images.
"""
from copy import deepcopy

from docx.table import _Row
from docx.text.paragraph import Paragraph

W_NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def set_text(p, text):
    """Replace paragraph text, keeping the first run's formatting."""
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        p.add_run(text)


def clean_clone(p):
    """Deep-copy paragraph XML, stripping bookmarks to avoid duplicate ids."""
    el = deepcopy(p._p)
    for tag in ('bookmarkStart', 'bookmarkEnd'):
        for b in el.findall(f'.//{W_NS}{tag}'):
            b.getparent().remove(b)
    return el


def insert_after(anchor_p, template_p, text):
    """Insert a cloned paragraph (styled like template_p) after anchor_p.

    Returns the new Paragraph so successive insertions can chain:
        p = insert_after(anchor, tpl, 'first')
        p = insert_after(p, tpl, 'second')

    IMPORTANT: when inserting at several anchor positions located by index,
    insert bottom-up (largest paragraph index first) so earlier indices stay
    valid.
    """
    el = clean_clone(template_p)
    anchor_p._p.addnext(el)
    np = Paragraph(el, anchor_p._parent)
    set_text(np, text)
    return np


def add_row_after(table, match_text, new_cell_texts, col=0):
    """Clone the table row whose cell[col] contains match_text, insert the
    clone right after it, and set each cell's text from new_cell_texts.
    Cell formatting is inherited from the cloned row."""
    for row in table.rows:
        if match_text in row.cells[col].text:
            new_tr = deepcopy(row._tr)
            row._tr.addnext(new_tr)
            new_row = _Row(new_tr, table)
            for i, t in enumerate(new_cell_texts):
                set_text(new_row.cells[i].paragraphs[0], t)
            return new_row
    raise ValueError(f'row containing {match_text!r} not found')


def first_with_style(doc, style_name):
    """First paragraph using the given style (e.g. 'Heading 3', 'Normal').
    Use these as clone templates so inserted paragraphs match the document."""
    for p in doc.paragraphs:
        if p.style and p.style.name == style_name and p.text.strip():
            return p
    raise ValueError(f'no paragraph with style {style_name!r}')


def para_index(doc, startswith):
    """Index of the first paragraph whose text starts with the given string.
    Always assert anchor text before editing — indices shift after inserts."""
    for i, p in enumerate(doc.paragraphs):
        if p.text.startswith(startswith):
            return i
    raise ValueError(f'no paragraph starting with {startswith!r}')
