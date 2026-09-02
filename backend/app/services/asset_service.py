"""资产库 Excel 批量导入 + 人员证书结构工具。"""
from datetime import date, datetime

from openpyxl import load_workbook

from app import settings_store
from app.db import Database

COLUMN_MAPS = {
    "credit": {
        "名称": "name", "证书名称": "name",
        "发证机关": "fields.发证机关", "发证日期": "fields.发证日期",
        "有效期至": "expiry_date", "有效期": "expiry_date",
    },
}

# 人员 Excel 导入的旧列名兼容（字段配置之外的证书列 → 证书数组条目）
_PERSON_CERT_COLUMN_MAP = {
    "类型": "类型", "证书.类型": "类型",
    "证书有效期至": "有效期至", "有效期至": "有效期至", "有效期": "有效期至",
    "证书.有效期至": "有效期至",
    "编号": "编号", "证书编号": "编号", "证书.编号": "编号",
    "专业": "专业", "证书.专业": "专业",
    "执业时间": "执业时间", "证书.执业时间": "执业时间",
}


def _person_col_map() -> dict:
    """人员导入列映射：field-config 的 key → fields.<key>，另加姓名/证书列。"""
    col_map = {"姓名": "name"}
    for f in settings_store.get_field_config()["person"]:
        if f["key"] != "姓名":
            col_map.setdefault(f["key"], f"fields.{f['key']}")
    for header, cert_key in _PERSON_CERT_COLUMN_MAP.items():
        col_map.setdefault(header, f"cert.{cert_key}")
    return col_map


def _contract_col_map() -> dict:
    """合同导入列映射：项目名称 → name，field-config 其余 key（含项目经理）→ fields.<key>。"""
    col_map = {"项目名称": "name"}
    for f in settings_store.get_field_config()["contract"]:
        if f["key"] != "项目名称":
            col_map.setdefault(f["key"], f"fields.{f['key']}")
    return col_map


# ---------- 人员证书（一人多证） ----------

CERT_ITEM_KEYS = ("类型", "编号", "专业", "执业时间", "有效期至")
# 旧单值字段 → 证书条目键（懒迁移用）
LEGACY_CERT_FIELD_MAP = {"证书有效期至": "有效期至",
                         "证书编号": "编号", "编号": "编号",
                         "类型": "类型", "专业": "专业", "执业时间": "执业时间"}


def person_certs(fields: dict) -> list:
    """取人员证书数组；无"证书"数组但有旧单值字段时懒合成一条（不改原 dict）。"""
    fields = fields or {}
    certs = fields.get("证书")
    if isinstance(certs, list):
        return [dict(c) for c in certs if isinstance(c, dict)]
    legacy = {}
    for old_key, new_key in LEGACY_CERT_FIELD_MAP.items():
        value = fields.get(old_key)
        if value not in (None, ""):
            legacy[new_key] = value
    if legacy:
        return [legacy]
    return []


def earliest_expiry(certs: list) -> str:
    """所有证书中最早的有效期（YYYY-MM-DD），无有效日期返回空串。"""
    dates = []
    for c in certs or []:
        d = normalize_date((c or {}).get("有效期至") or "")
        if d:
            dates.append(d)
    return min(dates) if dates else ""


def solidify_person_fields(fields: dict) -> dict:
    """固化新结构：保证"证书"键为数组，移除已迁移的旧单值证书字段。"""
    fields = dict(fields or {})
    certs = person_certs(fields)
    for old_key in LEGACY_CERT_FIELD_MAP:
        fields.pop(old_key, None)
    fields["证书"] = certs
    return fields


def normalize_person_asset(asset: dict) -> dict:
    """GET 输出用：person 资产 fields 保证含"证书"数组（懒迁移），
    expiry_date 取列值与证书最早有效期中的较早者。"""
    a = dict(asset)
    a["fields"] = solidify_person_fields(a.get("fields") or {})
    computed = earliest_expiry(a["fields"]["证书"])
    candidates = [d for d in (a.get("expiry_date") or "", computed) if d]
    if candidates:
        a["expiry_date"] = min(candidates)
    return a


def normalize_date(value) -> str:
    """统一为 YYYY-MM-DD；无法解析返回空串。"""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    s = str(value).strip().replace("/", "-").replace(".", "-")
    if not s:
        return ""
    try:
        parts = [int(p) for p in s.split("-")]
        if len(parts) == 3:
            return date(parts[0], parts[1], parts[2]).isoformat()
    except (ValueError, IndexError):
        pass
    return ""


_CERT_DATE_KEYS = ("有效期至", "执业时间")


def _cert_cell_str(cert_key: str, value) -> str:
    """证书列值规范化：日期统一 YYYY-MM-DD（含 2013.7.23 形态）；
    数字（如编号 32092664 / 32092664.0）转字符串并去掉浮点小数点。"""
    if isinstance(value, (datetime, date)):
        return normalize_date(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    s = str(value).strip()
    if cert_key in _CERT_DATE_KEYS:
        return normalize_date(s) or s
    return s


def import_assets_excel(db: Database, asset_type: str, file_path: str) -> dict:
    if asset_type == "credit":
        col_map = COLUMN_MAPS["credit"]
    elif asset_type == "contract":
        col_map = _contract_col_map()
    else:
        col_map = _person_col_map()
    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        wb.close()
        return {"imported": 0, "errors": [{"row": 1, "reason": "文件中没有工作表"}]}
    rows = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h else "" for h in next(rows)]
    col_index = {}
    for idx, header in enumerate(headers):
        if header in col_map:
            col_index[idx] = col_map[header]
    if "name" not in col_index.values():
        wb.close()
        return {"imported": 0, "errors": [{"row": 1, "reason": "未找到名称列（首行需包含名称/证书名称/姓名/项目名称）"}]}

    imported = 0
    errors = []
    person_ids: dict = {}  # 姓名 → asset_id（人员同一人多行证书合并）
    for row_num, row in enumerate(rows, start=2):
        record = {"name": "", "expiry_date": "", "fields": {}, "cert": {}}
        for idx, target in col_index.items():
            value = row[idx] if idx < len(row) else None
            if target == "name":
                record["name"] = str(value).strip() if value else ""
            elif target == "expiry_date":
                record["expiry_date"] = normalize_date(value)
            elif target.startswith("cert."):
                cert_key = target.split(".", 1)[1]
                if value is not None and str(value).strip():
                    record["cert"][cert_key] = _cert_cell_str(cert_key, value)
            else:
                field_name = target.split(".", 1)[1]
                if value is not None and str(value).strip():
                    if isinstance(value, (datetime, date)):
                        record["fields"][field_name] = normalize_date(value)
                    else:
                        record["fields"][field_name] = str(value).strip()
        if not record["name"]:
            if any(v is not None and str(v).strip() for v in row):
                errors.append({"row": row_num, "reason": "名称为空"})
            continue
        if asset_type == "person":
            imported += _import_person_row(db, person_ids, record)
        else:
            db.create_asset(asset_type, record["name"], record["fields"],
                            expiry_date=record["expiry_date"])
            imported += 1
    wb.close()
    return {"imported": imported, "errors": errors}


def _import_person_row(db: Database, person_ids: dict, record: dict) -> int:
    """人员行入库：同姓名多行往"证书"数组追加（同人不同证不覆盖）。返回计数增量。"""
    fields = dict(record["fields"])
    cert = record["cert"]
    if cert:
        cert.setdefault("类型", fields.get("类型", ""))
        cert.setdefault("编号", "")
        cert["有效期至"] = normalize_date(cert.get("有效期至") or record["expiry_date"])
        fields["证书"] = [cert]
    else:
        fields["证书"] = []
    expiry = earliest_expiry(fields["证书"]) or record["expiry_date"]

    existing_id = person_ids.get(record["name"])
    if existing_id is None:
        new_id = db.create_asset("person", record["name"], fields, expiry_date=expiry)
        person_ids[record["name"]] = new_id
        return 1

    existing = db.get_asset(existing_id) or {}
    merged = solidify_person_fields(existing.get("fields") or {})
    for k, v in fields.items():
        if k != "证书" and v not in (None, ""):
            merged[k] = v
    if cert:
        merged["证书"] = merge_certs(merged["证书"], fields["证书"])
    new_expiry = earliest_expiry(merged["证书"]) or expiry
    db.update_asset(existing_id, fields=merged, expiry_date=new_expiry)
    return 1


def merge_certs(existing_certs: list, new_certs: list) -> list:
    """合并证书数组：同（类型, 编号, 专业）更新条目，否则追加（不覆盖其他证书）。
    去重键空值归一为 ""。"""
    merged = [dict(c) for c in existing_certs or [] if isinstance(c, dict)]

    def _key(c: dict):
        return (str(c.get("类型") or ""), str(c.get("编号") or ""),
                str(c.get("专业") or ""))

    index = {}
    for i, c in enumerate(merged):
        index.setdefault(_key(c), i)
    for cert in new_certs or []:
        if not isinstance(cert, dict):
            continue
        key = _key(cert)
        if key in index:
            merged[index[key]].update(
                {k: v for k, v in cert.items() if v not in (None, "")})
        else:
            index[key] = len(merged)
            merged.append(dict(cert))
    return merged


# ---------- 人员关联业绩（只读组装，不落库） ----------

def person_perfs(db: Database, person_name: str) -> list:
    """按 contract 资产中 项目经理==person_name 实时组装"业绩"数组。"""
    result = []
    for c in db.get_assets(type="contract"):
        fields = c.get("fields") or {}
        if fields.get("项目经理") != person_name:
            continue
        entry = {"项目名称": c["name"]}
        entry.update({k: v for k, v in fields.items() if k != "项目经理"})
        result.append(entry)
    return result


# ---------- 导入模板 ----------

TEMPLATE_TYPES = ("person", "contract")


def import_template_headers(asset_type: str) -> list:
    """按当前 field-config 生成模板表头。"""
    cfg = settings_store.get_field_config()
    if asset_type == "person":
        return (["姓名"] + [f["key"] for f in cfg["person"]]
                + ["类型", "证书.编号", "专业", "执业时间", "证书.有效期至"])
    if asset_type == "contract":
        return [f["key"] for f in cfg["contract"]]
    raise ValueError(f"不支持的模板类型：{asset_type}")


def build_import_template(asset_type: str) -> bytes:
    """生成导入模板 xlsx（仅表头行），返回字节流。"""
    from io import BytesIO

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    if ws is not None:
        ws.append(import_template_headers(asset_type))
    buf = BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()
