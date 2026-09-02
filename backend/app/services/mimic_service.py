# -*- coding: utf-8 -*-
"""保格式仿写业务编排：锚点转储 → LLM 规划编辑指令 → 应用并校验。"""
import shutil
import time
from pathlib import Path

from docx import Document

from app import config
from app.core.docx_dump import dump_docx
from app.core.docx_edit import add_row_after, first_with_style, insert_after, set_text
from app.core.docx_verify import verify_output
from app.core.llm_parser import LLMParseError
from app.core.mimic_planner import plan_edits
from app.db import Database
from app.settings_store import get_api_key, get_model


class MimicService:
    def __init__(self, db_path=None):
        self.db = Database(db_path)
        self.db.init_schema()

    def _template_path(self, template_id: int) -> str:
        if not self.db.get_template(template_id):
            raise ValueError("模板不存在")
        path = self.db.get_template(template_id).get("file_path") or ""
        if not path or not Path(path).exists():
            raise ValueError("模板文件不存在")
        return path

    @staticmethod
    def _project_info_text(project: dict) -> str:
        lines = []
        if project.get("name"):
            lines.append(f"项目名称：{project['name']}")
        if project.get("client"):
            lines.append(f"招标人/客户：{project['client']}")
        if project.get("project_type"):
            lines.append(f"项目类型：{project['project_type']}")
        if project.get("bid_date"):
            lines.append(f"投标截止/开标日期：{project['bid_date']}")
        if project.get("notes"):
            lines.append(f"项目备注：{project['notes']}")
        return ("项目信息：\n" + "\n".join(lines)) if lines else ""

    @staticmethod
    def _anchor_dump_text(dump: dict) -> str:
        lines = [f"[{p['index']}] <{p['style']}> {p['text']}" for p in dump["paragraphs"]]
        for t in dump["tables"]:
            lines.append(f"===== 表格{t['index']} =====")
            for row in t["rows"]:
                lines.append(" | ".join(row))
        return "\n".join(lines)

    @staticmethod
    def _op_text(o: dict) -> str:
        parts = [o.get("text") or ""]
        for it in (o.get("items") or []):
            parts.append(it.get("text") or "")
        parts.extend(o.get("cells") or [])
        return "".join(parts)

    def build_plan(self, project_id: int, template_id: int) -> dict:
        if not self.db.get_project(project_id):
            raise ValueError("项目不存在")
        tpl_path = self._template_path(template_id)
        api_key = get_api_key()
        if not api_key:
            raise LLMParseError("未配置 API Key")
        requirements = self.db.get_requirements(project_id)
        requirements_text = "\n".join(r["content"] for r in requirements)
        project_info = self._project_info_text(self.db.get_project(project_id) or {})
        anchor_text = self._anchor_dump_text(dump_docx(tpl_path))
        ops = plan_edits(anchor_text, requirements_text, project_info,
                         api_key, get_model())

        n_insert = sum(1 for o in ops if o["op"] == "insert")
        n_set = sum(1 for o in ops if o["op"] == "set_text")
        n_row = sum(1 for o in ops if o["op"] == "add_row")
        summary = f"共 {len(ops)} 处修改：新增 {n_insert} 处、改写 {n_set} 处、表格追加 {n_row} 行"
        placeholder_count = sum(1 for o in ops if "【需人工替换" in self._op_text(o))
        return {"ops": ops, "summary": summary, "placeholder_count": placeholder_count}

    def apply_plan(self, project_id: int, template_id: int, ops: list) -> dict:
        if not self.db.get_project(project_id):
            raise ValueError("项目不存在")
        tpl_path = self._template_path(template_id)
        if Path(tpl_path).suffix.lower() != ".docx":
            raise ValueError("仅支持 .docx 模板")

        out_dir = config.FILES_DIR / "mimic_outputs"
        out_dir.mkdir(parents=True, exist_ok=True)
        out_name = f"mimic_{project_id}_{int(time.time())}.docx"
        out_path = out_dir / out_name
        shutil.copy2(tpl_path, out_path)

        doc = Document(str(out_path))
        paras = doc.paragraphs

        def check_anchor(index, anchor, k):
            if index < 0 or index >= len(paras):
                raise ValueError(f"第{k}条操作段落序号 {index} 越界")
            if not paras[index].text.startswith(anchor):
                raise ValueError(f"第{k}条操作锚点不匹配：第{index}段应以“{anchor[:20]}”开头")

        # 1) set_text / add_row（顺序无关，按 ops 原始顺序执行）
        for k, o in enumerate(ops, 1):
            if o["op"] == "set_text":
                check_anchor(o["index"], o["anchor"], k)
                set_text(paras[o["index"]], o["text"])
            elif o["op"] == "add_row":
                if o["table"] >= len(doc.tables):
                    raise ValueError(f"第{k}条操作表格序号 {o['table']} 越界")
                add_row_after(doc.tables[o["table"]], o["match"], o["cells"])

        # 2) insert（自底向上，避免前面的插入使后面段落序号整体偏移）
        inserts = [o for o in ops if o["op"] == "insert"]
        for o in sorted(inserts, key=lambda x: x["index"], reverse=True):
            check_anchor(o["index"], o["anchor"], "insert")
            anchor_p = paras[o["index"]]
            first = o["items"][0]
            p = insert_after(anchor_p, first_with_style(doc, first["style"]), first["text"])
            for it in o["items"][1:]:
                p = insert_after(p, first_with_style(doc, it["style"]), it["text"])

        doc.save(str(out_path))

        report = verify_output(tpl_path, str(out_path))
        return {
            "output_file": out_name,
            "download_url": f"/api/projects/{project_id}/mimic/download?file={out_name}",
            "report": report,
            "warnings": [
                "目录为 Word 域，请在 Word 中右键目录→更新域→更新整个目录",
                "模板中的图片（示意图/截图）请人工确认是否需替换为本项目",
                "所有【需人工替换：…】标注位置请人工填写",
            ],
        }
