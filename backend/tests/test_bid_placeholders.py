import re
from docx import Document
from docx.shared import Cm

from app.core.bid_template_exporter import (
    DEFAULT_PLACEHOLDER_SYNONYMS, _iter_all_paragraphs,
    _iter_textbox_paragraphs, build_placeholder_rules, scan_placeholders,
)


def _doc_with(paras):
    doc = Document()
    for p in paras:
        doc.add_paragraph(p)
    return doc


def test_label_blank_fills_expanded_labels():
    doc = _doc_with(["项目名称：____", "采购人：____", "标段名称：____",
                     "标段编号：", "日期：____年__月__日"])
    rules = build_placeholder_rules(
        doc, project_name="里下河项目", tenderer="某中心",
        section_name="一标段", section_no="A1", doc_date="2026-09-10")
    full = "\n".join(p.text for p in _iter_all_paragraphs(doc))
    for r in rules:
        full = r["pattern"].sub(r["repl"], full)
    assert "项目名称：里下河项目" in full
    assert "采购人：某中心" in full          # 同义词 采购人→招标人
    assert "标段名称：一标段" in full
    assert "标段编号：A1" in full
    assert "日期：2026年9月10日" in full


def test_bracket_half_and_full_width_filled():
    doc = _doc_with(["(项目名称) (标段名称)", "（招标人名称）（投标人名称）"])
    rules = build_placeholder_rules(
        doc, project_name="P", tenderer="T", bidder_name="B",
        section_name="S")
    full = "\n".join(p.text for p in _iter_all_paragraphs(doc))
    for r in rules:
        full = r["pattern"].sub(r["repl"], full)
    assert "P (S)" in full
    assert "T" in full and "B" in full
    assert "（" not in full and "(项目名称)" not in full


def test_section_combo_keeps_full_width_brackets():
    """全角组合占位（项目名称）（标段名称）→ 保留全角括号风格。"""
    doc = _doc_with(["（项目名称）（标段名称）"])
    rules = build_placeholder_rules(doc, project_name="P", section_name="S")
    full = "\n".join(p.text for p in _iter_all_paragraphs(doc))
    for r in rules:
        full = r["pattern"].sub(r["repl"], full)
    assert "P（S）" in full
    assert "（项目名称）" not in full


def test_textbox_paragraphs_are_traversed_and_no_dup():
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run()
    # 构造一个含 w:txbxContent 的极简 drawing
    xml = (
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:r><w:drawing><wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
        '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:graphicData '
        'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:wsp xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:txbx><w:txbxContent><w:p><w:r><w:t>招标人：____</w:t></w:r></w:p>'
        '</w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:inline>'
        '</w:drawing></w:r></w:p>')
    from lxml import etree
    p._p.addnext(etree.fromstring(xml))
    texts = [t.text for t in _iter_textbox_paragraphs(doc)]
    assert texts.count("招标人：____") == 1


def test_textbox_alternate_content_fallback_not_duplicated():
    """Word 真实写法的文本框：mc:AlternateContent 中 mc:Choice（DrawingML）
    与 mc:Fallback（VML）各含一份相同 w:txbxContent —— 只应产出一次。"""
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run()
    inner = ('<w:p><w:r><w:t>招标人：____</w:t></w:r></w:p>')
    xml = (
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'xmlns:v="urn:schemas-microsoft-com:vml">'
        '<w:r><mc:AlternateContent>'
        '<mc:Choice Requires="wps">'
        '<w:drawing><wp:inline><a:graphic><a:graphicData '
        'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:wsp><wps:txbx><w:txbxContent>' + inner +
        '</w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic>'
        '</wp:inline></w:drawing>'
        '</mc:Choice>'
        '<mc:Fallback>'
        '<w:pict><v:shape><v:textbox><w:txbxContent>' + inner +
        '</w:txbxContent></v:textbox></v:shape></w:pict>'
        '</mc:Fallback>'
        '</mc:AlternateContent></w:r></w:p>')
    from lxml import etree
    p._p.addnext(etree.fromstring(xml))
    texts = [t.text for t in _iter_textbox_paragraphs(doc)]
    assert texts.count("招标人：____") == 1


def test_scan_placeholders_reports_match_and_suspicious():
    doc = _doc_with(["招标人：____", "______________________________"])
    res = scan_placeholders(doc, {"tenderer": "某中心"})
    assert any(m["label"] == "招标人" and m["value"] == "某中心"
               for m in res["matched"])
    assert res["suspicious"], "孤立下划线应列入可疑清单"


def test_fill_draft_synonyms_fill_alias(tmp_path):
    """导出端接入同义词库（V1.2 4.5）：synonyms={"甲方":"招标人"} 时底稿
    "甲方：____"被填充且 verify 同源不误报；不传 synonyms 时不填。"""
    from app.core.bid_draft_exporter import fill_draft
    doc = Document()
    doc.add_paragraph("甲方：____")
    draft = tmp_path / "d.docx"
    doc.save(str(draft))
    out = tmp_path / "o.docx"
    report = fill_draft(str(draft), str(out),
                        {"tables": [], "swap_toc": False}, {},
                        tenderer="某中心", synonyms={"甲方": "招标人"})
    texts = [p.text for p in Document(str(out)).paragraphs]
    assert "甲方：某中心" in texts
    assert report["verify"]["ok"] is True

    out2 = tmp_path / "o2.docx"
    fill_draft(str(draft), str(out2), {"tables": [], "swap_toc": False}, {},
               tenderer="某中心")
    texts2 = [p.text for p in Document(str(out2)).paragraphs]
    assert "甲方：____" in texts2
