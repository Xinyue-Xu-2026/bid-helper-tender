from PIL import Image
from docx import Document
from app.core.bid_template_exporter import (
    fit_image_cm, insert_image_adaptive, cell_available_width_cm)

def _png(path, w, h):
    Image.new("RGB", (w, h), "white").save(path)

def test_fit_image_limits_by_height(tmp_path):
    p = tmp_path / "tall.png"; _png(p, 400, 4000)   # 10:1 竖长
    r = fit_image_cm(str(p), avail_w_cm=7.0, avail_h_cm=24.0)
    assert r["height_cm"] <= 24.0 + 1e-6
    assert r["width_cm"] < 7.0
    assert r["too_long"] is True

def test_fit_image_uses_cell_width_when_wide(tmp_path):
    p = tmp_path / "wide.png"; _png(p, 2000, 1000)
    r = fit_image_cm(str(p), avail_w_cm=7.0, avail_h_cm=24.0)
    assert abs(r["width_cm"] - 7.0) < 1e-6
    assert abs(r["height_cm"] - 3.5) < 0.01

def test_insert_reports_missing_file(tmp_path):
    doc = Document(); cell = doc.add_table(1, 1).rows[0].cells[0]
    r = insert_image_adaptive(cell, str(tmp_path / "nope.png"), 6.0)
    assert r["ok"] is False and "不存在" in r["reason"]

def test_insert_reports_unsupported_format(tmp_path):
    f = tmp_path / "x.emf"; f.write_bytes(b"not an image")
    doc = Document(); cell = doc.add_table(1, 1).rows[0].cells[0]
    r = insert_image_adaptive(cell, str(f), 6.0)
    assert r["ok"] is False and r["reason"]

def test_cell_width_from_tblgrid(tmp_path):
    from docx.shared import Cm
    doc = Document(); t = doc.add_table(1, 1); t.columns[0].width = Cm(6)
    assert 5.5 < cell_available_width_cm(t.rows[0].cells[0]) < 6.5
