"""招标文件输入归一化测试：docx 原样返回、不支持后缀报错、扫描版 PDF 判定、
文本型 PDF 转 docx、.doc 无 Word 时优雅报错、内容缓存命中不重复转换。"""
from pathlib import Path

import fitz
import pytest
from docx import Document

from app.core import input_convert
from app.core.input_convert import InputConvertError, ensure_docx


def _make_blank_pdf(path: Path, pages: int = 2):
    """程序化生成无文字的扫描版 PDF fixture。"""
    doc = fitz.open()
    for _ in range(pages):
        doc.new_page()
    doc.save(str(path))
    doc.close()


def _make_text_pdf(path: Path, lines):
    """程序化生成含文字的文本型 PDF fixture。"""
    doc = fitz.open()
    page = doc.new_page()
    y = 72
    for line in lines:
        page.insert_text((72, y), line)
        y += 14
    doc.save(str(path))
    doc.close()


def test_docx_passthrough(tmp_path):
    """docx 原样返回同一路径，不写缓存。"""
    src = tmp_path / "招标文件.docx"
    Document().save(str(src))
    dest_dir = tmp_path / "cache"
    result = ensure_docx(str(src), str(dest_dir))
    assert result == src
    assert not dest_dir.exists()  # 未创建缓存目录、未复制


def test_unsupported_suffix(tmp_path):
    """.txt 等不支持的后缀抛 InputConvertError。"""
    src = tmp_path / "notes.txt"
    src.write_text("随便写点内容", encoding="utf-8")
    with pytest.raises(InputConvertError, match=r"\.txt"):
        ensure_docx(str(src), str(tmp_path / "cache"))


def test_scanned_pdf_rejected(tmp_path):
    """2 页无文字 PDF 判定为扫描版，错误消息含"扫描版"。"""
    src = tmp_path / "scan.pdf"
    _make_blank_pdf(src, pages=2)
    with pytest.raises(InputConvertError, match="扫描版"):
        ensure_docx(str(src), str(tmp_path / "cache"))


def test_doc_without_word(tmp_path, monkeypatch):
    """DispatchEx 抛异常（模拟本机无 Word）→ InputConvertError 且消息含 Microsoft Word。"""
    src = tmp_path / "旧版招标文件.doc"
    src.write_bytes(b"\xd0\xcf\x11\xe0fake-doc")
    import win32com.client

    def _boom(*args, **kwargs):
        raise RuntimeError("Word 未安装")

    monkeypatch.setattr(win32com.client, "DispatchEx", _boom)
    with pytest.raises(InputConvertError, match="Microsoft Word"):
        ensure_docx(str(src), str(tmp_path / "cache"))


def test_text_pdf_converted(tmp_path):
    """文本型 PDF 经 pdf2docx 转换：产物存在、python-docx 可打开、段落文本含原文。"""
    pytest.importorskip("pdf2docx")
    lines = ["Chapter One General Provisions of bidding doc"] * 10
    src = tmp_path / "bid.pdf"
    _make_text_pdf(src, lines)
    dest_dir = tmp_path / "cache"
    result = ensure_docx(str(src), str(dest_dir))
    assert result.suffix == ".docx" and result.exists()
    assert result.parent == dest_dir
    doc = Document(str(result))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Chapter One" in all_text


def test_pdf_cache_hit_skips_conversion(tmp_path, monkeypatch):
    """同一 src 第二次调用命中内容缓存，不再触发转换器。"""
    src = tmp_path / "bid.pdf"
    _make_text_pdf(src, ["缓存测试文本行"] * 10)
    calls = {"n": 0}

    def _fake_convert(s, d):
        calls["n"] += 1
        Document().save(str(d))

    monkeypatch.setattr(input_convert, "_convert_pdf", _fake_convert)
    dest_dir = tmp_path / "cache"
    r1 = ensure_docx(str(src), str(dest_dir))
    r2 = ensure_docx(str(src), str(dest_dir))
    assert calls["n"] == 1
    assert r1 == r2 and r1.exists()


def test_failed_conversion_leaves_no_poisoned_cache(tmp_path, monkeypatch):
    """转换中途写了一半即失败：缓存目标不得残留半成品（否则同路径+大小+
    mtime 的下次调用会永久命中坏缓存）；失败后重试须可成功。"""
    src = tmp_path / "bid.pdf"
    _make_text_pdf(src, ["缓存原子性测试文本"] * 10)

    def _broken(s, d):
        Path(d).write_bytes(b"partial-broken")  # 写了一半就崩
        raise RuntimeError("转换中途失败")

    monkeypatch.setattr(input_convert, "_convert_pdf", _broken)
    dest_dir = tmp_path / "cache"
    with pytest.raises(RuntimeError, match="转换中途失败"):
        ensure_docx(str(src), str(dest_dir))
    # 最终缓存键不得存在；半成品临时文件也被清理
    assert not list(dest_dir.glob("*.docx"))

    def _good(s, d):
        Document().save(str(d))

    monkeypatch.setattr(input_convert, "_convert_pdf", _good)
    result = ensure_docx(str(src), str(dest_dir))
    assert result.exists()
    assert Document(str(result)) is not None  # 重试成功且产物可解析
