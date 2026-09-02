"""contract 资产类型、导入模板、旧业绩迁移、人员业绩关联注入的测试。"""
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import settings_store
from app.db import Database
from app.deps import get_db
from app.main import app
from app.services import import_service


@pytest.fixture()
def db(db_path):
    d = Database(db_path)
    d.init_schema()
    return d


@pytest.fixture()
def client(db):
    app.dependency_overrides[get_db] = lambda: db
    yield TestClient(app)
    app.dependency_overrides.clear()


def _xlsx_bytes(headers, rows):
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet_headers(content: bytes):
    wb = load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.active
    assert ws is not None
    headers = [c.value for c in next(ws.iter_rows(max_row=1))]
    wb.close()
    return headers


# ---------- 字段配置：已存配置含"姓名"时 GET 剔除 ----------

def test_field_config_strips_saved_name(client):
    settings_store.save_settings({"field_config": {
        "person": [{"key": "姓名", "type": "text", "options": []},
                   {"key": "职称", "type": "text", "options": []}],
        "contract": [],
    }})
    cfg = client.get("/api/assets/field-config").json()
    assert [f["key"] for f in cfg["person"]] == ["职称"]


# ---------- 导入模板 ----------

def test_import_template_person(client):
    r = client.get("/api/assets/import-template", params={"type": "person"})
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert "filename*=UTF-8''" in r.headers["content-disposition"]
    assert _sheet_headers(r.content) == [
        "姓名", "部门", "职称", "联系方式",
        "类型", "证书.编号", "专业", "执业时间", "证书.有效期至"]


def test_import_template_contract(client):
    r = client.get("/api/assets/import-template", params={"type": "contract"})
    assert r.status_code == 200
    assert _sheet_headers(r.content) == [
        "项目名称", "类型", "合同金额", "年份", "甲方", "项目经理"]


def test_import_template_bad_type(client):
    assert client.get("/api/assets/import-template", params={"type": "credit"}).status_code == 400


# ---------- contract Excel 导入 ----------

def test_contract_excel_import(client):
    content = _xlsx_bytes(
        ["项目名称", "类型", "合同金额", "年份", "甲方", "项目经理"],
        [["项目A", "编标", "120万", "2024", "甲公司", "张三"],
         ["项目B", "结算", "80万", "2023", "乙公司", "李四"]])
    r = client.post("/api/assets/import", params={"type": "contract"},
                    files={"file": ("contracts.xlsx", content,
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200
    assert r.json()["imported"] == 2
    assets = client.get("/api/assets", params={"type": "contract"}).json()
    by_name = {a["name"]: a for a in assets}
    assert by_name["项目A"]["fields"] == {
        "类型": "编标", "合同金额": "120万", "年份": "2024", "甲方": "甲公司", "项目经理": "张三"}
    assert by_name["项目B"]["fields"]["项目经理"] == "李四"


# ---------- 旧 person.业绩 迁移 ----------

def _seed_legacy_perfs(db_path):
    """不触发 init_schema 迁移，直接写旧结构数据。"""
    d = Database(db_path)
    d.init_schema()
    d.create_asset("person", "张三", {"职称": "高工", "业绩": [
        {"项目名称": "项目A", "年份": "2024", "合同金额": "120万"},
        {"项目名称": "项目B", "年份": "2023", "甲方": "乙公司"},
    ]})
    d.create_asset("person", "李四", {"业绩": [
        {"项目名称": "项目A", "年份": "2024", "甲方": "甲公司"},  # 与张三重复 → 合并
    ]})
    return d


def test_legacy_perfs_migration_idempotent(db_path):
    _seed_legacy_perfs(db_path)
    d = Database(db_path)  # 重新 init_schema 触发迁移
    d.init_schema()
    contracts = d.get_assets(type="contract")
    assert len(contracts) == 2  # 项目A 合并为一条
    by_name = {c["name"]: c for c in contracts}
    assert by_name["项目A"]["fields"]["项目经理"] == "张三"  # 已有非空值优先
    assert by_name["项目A"]["fields"]["甲方"] == "甲公司"   # 空值被李四的条目补齐
    assert by_name["项目B"]["fields"]["项目经理"] == "张三"
    # person.fields 不再含"业绩"
    for p in d.get_assets(type="person"):
        assert "业绩" not in p["fields"]
    # 幂等：再次 init_schema 不重复迁移
    d.init_schema()
    assert len(d.get_assets(type="contract")) == 2


# ---------- person GET 业绩关联注入（只读） ----------

def test_person_perfs_injection(client, db):
    aid = client.post("/api/assets", json={"type": "person", "name": "张三"}).json()["id"]
    client.post("/api/assets", json={"type": "contract", "name": "项目A", "fields": {
        "类型": "编标", "年份": "2024", "项目经理": "张三"}})
    client.post("/api/assets", json={"type": "contract", "name": "项目B", "fields": {
        "类型": "结算", "年份": "2023", "项目经理": "李四"}})
    got = client.get(f"/api/assets/{aid}").json()
    assert got["fields"]["业绩"] == [{"项目名称": "项目A", "类型": "编标", "年份": "2024"}]
    listed = client.get("/api/assets", params={"type": "person"}).json()
    assert len(listed[0]["fields"]["业绩"]) == 1
    # 只读：不落库到 person
    assert "业绩" not in (db.get_asset(aid)["fields"])


# ---------- import_service：业绩 confirm 落 contract ----------

def test_confirm_contract_creates_asset_and_placeholder_person(db):
    items = [{"asset_type": "contract", "action": "new", "fields": {
        "项目名称": "项目A", "类型": "编标", "年份": "2024",
        "甲方": "甲公司", "项目经理": "王五"}}]
    result = import_service.confirm_items(db, items)
    assert result["created"] == 1
    assert result["new_persons"] == ["王五"]  # 项目经理自动新建占位
    contracts = db.get_assets(type="contract")
    assert len(contracts) == 1
    assert contracts[0]["name"] == "项目A"
    assert contracts[0]["fields"]["项目经理"] == "王五"
    person = db.get_assets(type="person")[0]
    assert person["name"] == "王五" and person["fields"] == {"证书": []}


def test_confirm_contract_dedup_by_name_and_year(db):
    item = {"asset_type": "contract", "action": "new", "fields": {
        "项目名称": "项目A", "年份": "2024", "甲方": "甲公司"}}
    import_service.confirm_items(db, [item])
    # 同名同年份 → 合并更新；同名不同年份 → 新建
    import_service.confirm_items(db, [{"asset_type": "contract", "action": "update", "fields": {
        "项目名称": "项目A", "年份": "2024", "合同金额": "120万"}}])
    import_service.confirm_items(db, [{"asset_type": "contract", "action": "new", "fields": {
        "项目名称": "项目A", "年份": "2025", "甲方": "丙公司"}}])
    contracts = db.get_assets(type="contract")
    assert len(contracts) == 2
    c2024 = next(c for c in contracts if c["fields"]["年份"] == "2024")
    assert c2024["fields"]["甲方"] == "甲公司" and c2024["fields"]["合同金额"] == "120万"


def test_confirm_legacy_person_perfs_converted(db):
    """旧格式 person 条目携带"业绩"数组 → 转 contract 资产。"""
    import_service.confirm_items(db, [{"asset_type": "person", "action": "new", "fields": {
        "姓名": "赵六", "职称": "高工",
        "业绩": [{"项目名称": "项目C", "年份": "2022", "类型": "跟踪"}]}}])
    contracts = db.get_assets(type="contract")
    assert len(contracts) == 1
    assert contracts[0]["name"] == "项目C"
    assert contracts[0]["fields"]["项目经理"] == "赵六"
    person = db.get_assets(type="person")[0]
    assert "业绩" not in person["fields"]
    assert person["fields"]["职称"] == "高工"
