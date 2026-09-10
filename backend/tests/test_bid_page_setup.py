"""页眉移植与分节页码测试（P2）：参考包页眉（含图片）/页脚（PAGE 域）
包级移植进导出产物；封面+目录节无页眉无页码，正文节页码从 1 起；
参考缺失降级不炸；产物结构有效可重开；verify 不因分节符误报。"""
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from app import config
from app.core.bid_draft_exporter import fill_draft
from app.core.bid_page_setup import _attach_refs
from tests.test_bid_draft_export import (
    PNG_1X1, _bindings, _build_draft, _data,
)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _wq(tag):
    return f"{{{W}}}{tag}"


def _build_reference(path: Path):
    """程序化参考文件：图片页眉 + PAGE 域页脚（模拟宏信天德参考包结构）。"""
    logo = path.parent / "ref_logo.png"
    logo.write_bytes(PNG_1X1)
    doc = Document()
    section = doc.sections[0]
    section.header.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.add_run().add_picture(str(logo))
    hp.add_run("宏信天德")
    section.footer.is_linked_to_previous = False
    fp = section.footer.paragraphs[0]
    run = fp.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE \\* MERGEFORMAT"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.append(begin); run._r.append(instr); run._r.append(end)
    doc.save(str(path))
    return path


@pytest.fixture()
def reference(tmp_path, monkeypatch):
    ref = _build_reference(tmp_path / "ref.docx")
    monkeypatch.setattr(config, "BID_HEADER_SOURCE_PATH", ref)
    return ref


def _sectprs(zf):
    root = etree.fromstring(zf.read("word/document.xml"))
    body = root.find(_wq("body"))
    in_p = [p.find(f"{_wq('pPr')}/{_wq('sectPr')}")
            for p in body.findall(_wq("p"))]
    in_p = [sp for sp in in_p if sp is not None]
    body_sp = body.find(_wq("sectPr"))
    return in_p, body_sp


def _refs(sp, tag):
    return sp.findall(_wq(tag)) if sp is not None else []


# ---------- 1. 封面+目录节排除页眉页码；正文节页码从 1 起 ----------

def test_header_footer_transplant_with_toc_split(tmp_path, reference):
    draft = _build_draft(tmp_path / "draft.docx")  # 3 封面段 + 2 toc 段 + 表
    out = str(tmp_path / "out.docx")
    report = fill_draft(draft, out, _bindings(swap_toc=True), _data())
    assert report["page_setup"]["header_applied"] is True
    assert report["page_setup"]["page_numbers"] is True
    assert report["page_setup"]["degraded"] == ""
    # 分节符插入不引起防篡改误报
    assert report["verify"]["ok"] is True, report["verify"]["issues"]

    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert "word/header_hxtd.xml" in names
        assert "word/footer_hxtd.xml" in names
        assert any(n.startswith("word/media/hxtd_") for n in names)  # logo
        footer = z.read("word/footer_hxtd.xml").decode("utf-8")
        assert "PAGE" in footer
        header = z.read("word/header_hxtd.xml").decode("utf-8")
        assert "宏信天德" in header
        # document.xml.rels 挂接移植 part
        rels = z.read("word/_rels/document.xml.rels").decode("utf-8")
        assert "header_hxtd.xml" in rels and "footer_hxtd.xml" in rels
        in_p, body_sp = _sectprs(z)
    # 拆成两节：节1（封面+目录）无 header/footer 引用；节2 有且页码从 1 起
    assert len(in_p) == 1
    assert _refs(in_p[0], "headerReference") == []
    assert _refs(in_p[0], "footerReference") == []
    assert _refs(body_sp, "headerReference")
    assert _refs(body_sp, "footerReference")
    pg = body_sp.find(_wq("pgNumType"))
    assert pg is not None and pg.get(_wq("start")) == "1"


def test_header_applied_single_section_without_toc(tmp_path, reference):
    """底稿无封面/目录（无 toc 段）→ 整篇一节挂页眉+页码，不拆节。"""
    draft = _build_draft(tmp_path / "draft.docx", with_toc=False)
    out = str(tmp_path / "out.docx")
    report = fill_draft(draft, out, _bindings(), _data())
    assert report["page_setup"]["header_applied"] is True
    assert report["page_setup"]["page_numbers"] is True
    assert report["verify"]["ok"] is True, report["verify"]["issues"]
    with zipfile.ZipFile(out) as z:
        in_p, body_sp = _sectprs(z)
    assert in_p == []  # 未拆节
    assert _refs(body_sp, "headerReference")
    assert _refs(body_sp, "footerReference")
    assert body_sp.find(_wq("pgNumType")) is None  # 单节无需重起页码


# ---------- 2. 参考文件缺失 → 降级不炸 ----------

def test_missing_reference_degrades_gracefully(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "BID_HEADER_SOURCE_PATH",
                        tmp_path / "不存在.docx")
    draft = _build_draft(tmp_path / "draft.docx")
    out = str(tmp_path / "out.docx")
    report = fill_draft(draft, out, _bindings(), _data())
    assert report["page_setup"]["header_applied"] is False
    assert report["page_setup"]["page_numbers"] is False
    assert report["page_setup"]["degraded"]
    assert report["verify"]["ok"] is True
    assert Document(out) is not None  # 产物正常


# ---------- 3. 产物结构有效：python-docx 重开 + rels/ContentTypes 完整 ----------

def test_output_reopens_and_structure_valid(tmp_path, reference):
    draft = _build_draft(tmp_path / "draft.docx")
    out = str(tmp_path / "out.docx")
    fill_draft(draft, out, _bindings(swap_toc=True), _data())
    # python-docx 重开（结构有效）
    doc = Document(out)
    assert doc.tables
    with zipfile.ZipFile(out) as z:
        ct = etree.fromstring(z.read("[Content_Types].xml"))
        ns = {"ct": "http://schemas.openxmlformats.org/package/2006/content-types"}
        overrides = {el.get("PartName") for el in ct.findall("ct:Override", ns)}
        assert "/word/header_hxtd.xml" in overrides
        assert "/word/footer_hxtd.xml" in overrides
        defaults = {el.get("Extension") for el in ct.findall("ct:Default", ns)}
        assert "png" in defaults
        # 移植 part 的 rels 指向重命名后的媒体
        hrels = z.read("word/_rels/header_hxtd.xml.rels").decode("utf-8")
        assert "hxtd_" in hrels
        # 媒体文件确实存在
        media = [n for n in z.namelist() if n.startswith("word/media/hxtd_")]
        assert media and all(z.read(n) for n in media)


# ---------- 5. pgNumType 插入位置符合 CT_SectPr 序列 ----------

def test_pg_num_type_inserted_before_later_elements():
    """sectPr 含 docGrid/vAlign（序列排在 pgNumType 之后）且无 cols 时，
    pgNumType 仍须插到它们之前（严格校验的 Word 不接受乱序）。"""
    sp = etree.Element(_wq("sectPr"))
    etree.SubElement(sp, _wq("pgSz"))
    etree.SubElement(sp, _wq("pgMar"))
    etree.SubElement(sp, _wq("vAlign"))
    etree.SubElement(sp, _wq("docGrid"))
    _attach_refs(sp, "rId9001", "rId9002", restart_page=True)
    tags = [etree.QName(c).localname for c in sp]
    pg_pos = tags.index("pgNumType")
    # pgNumType 先于所有后置元素（vAlign/docGrid），且在引用/pgSz/pgMar 之后
    assert pg_pos < tags.index("vAlign")
    assert pg_pos < tags.index("docGrid")
    assert pg_pos > tags.index("pgMar")
    assert tags.index("headerReference") == 0
    assert tags.index("footerReference") == 1


# ---------- 6. 静态文本目录分节 + 页码居中断言 + 分节报告（V1.2 7.5） ----------

from app.core.bid_page_setup import (
    apply_page_setup, check_footer_page_centered, is_static_toc_group)


def _page_field_footer_doc(path, alignment=None):
    """构造页脚含 PAGE 域段落的 docx；alignment 非空时显式设置对齐。"""
    doc = Document()
    section = doc.sections[0]
    section.footer.is_linked_to_previous = False
    fp = section.footer.paragraphs[0]
    if alignment is not None:
        fp.alignment = alignment
    run = fp.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE \\* MERGEFORMAT"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.append(begin); run._r.append(instr); run._r.append(end)
    doc.save(str(path))
    return path


def test_is_static_toc_group_detected():
    body = etree.fromstring(
        f'<w:body xmlns:w="{W}">'
        '<w:p><w:r><w:t>投标文件封面</w:t></w:r></w:p>'
        '<w:p><w:r><w:t>目录</w:t></w:r></w:p>'
        '<w:p><w:r><w:t>第一章 概述 ........ 1</w:t></w:r></w:p>'
        '<w:p><w:r><w:t>第二章 方案 ........ 5</w:t></w:r></w:p>'
        '<w:p><w:r><w:t>正文开始</w:t></w:r></w:p></w:body>')
    children = list(body)
    idx = is_static_toc_group(children)
    assert idx is not None
    assert idx == 3  # 目录组末元素 = 第二个页码段


def test_is_static_toc_group_requires_both_signals():
    """双重特征缺一不可：仅「目录」标题无页码段组 → None；
    仅页码形态段无「目录」标题 → None（防正文误判）。"""
    only_title = etree.fromstring(
        f'<w:body xmlns:w="{W}">'
        '<w:p><w:r><w:t>目录</w:t></w:r></w:p>'
        '<w:p><w:r><w:t>正文段落</w:t></w:r></w:p></w:body>')
    assert is_static_toc_group(list(only_title)) is None
    only_entries = etree.fromstring(
        f'<w:body xmlns:w="{W}">'
        '<w:p><w:r><w:t>第一章 概述 ........ 1</w:t></w:r></w:p>'
        '<w:p><w:r><w:t>第二章 方案 ........ 5</w:t></w:r></w:p></w:body>')
    assert is_static_toc_group(list(only_entries)) is None


def test_footer_page_centered_warns_when_left(tmp_path):
    """页脚 PAGE 域段落显式非居中（jc=left，有意设置）→ 非空告警含 part 名。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p = _page_field_footer_doc(tmp_path / "x.docx",
                               alignment=WD_ALIGN_PARAGRAPH.LEFT)
    warnings = check_footer_page_centered(str(p))
    assert warnings and any("footer" in w for w in warnings)


def test_footer_page_centered_autofix_missing_jc(tmp_path):
    """页脚 PAGE 域段落无 jc（缺省）→ 安全补 center 后复检通过：无告警，
    且产物被修正（幂等：二次调用仍无告警）。"""
    p = _page_field_footer_doc(tmp_path / "x.docx")
    assert check_footer_page_centered(str(p)) == []
    assert check_footer_page_centered(str(p)) == []
    with zipfile.ZipFile(str(p)) as z:
        footers = [n for n in z.namelist()
                   if n.startswith("word/footer") and n.endswith(".xml")]
        assert footers
        assert 'w:val="center"' in z.read(footers[0]).decode("utf-8")


def test_footer_page_centered_ok_when_center(tmp_path):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p = _page_field_footer_doc(tmp_path / "x.docx",
                               alignment=WD_ALIGN_PARAGRAPH.CENTER)
    assert check_footer_page_centered(str(p)) == []


def test_apply_page_setup_reports_sections(tmp_path, reference):
    """静态文本目录（无 toc 样式段）也分节：封面+目录为节1（无引用无重起），
    正文节挂参考页眉页脚且页码从 1 起；result 报告 sections /
    section_page_numbers / warnings。"""
    doc = Document()
    doc.add_paragraph("投标文件封面")
    doc.add_paragraph("目录")
    doc.add_paragraph("第一章 概述 ........ 1")
    doc.add_paragraph("第二章 方案 ........ 5")
    doc.add_paragraph("正文开始")
    p = tmp_path / "out.docx"
    doc.save(str(p))
    result = apply_page_setup(str(p))
    assert result["header_applied"] is True
    assert result["sections"] >= 2
    assert result["section_page_numbers"][0] is False   # 封面+目录节
    assert result["section_page_numbers"][-1] is True   # 正文节
    assert result["warnings"] == []
    with zipfile.ZipFile(str(p)) as z:
        in_p, body_sp = _sectprs(z)
    assert len(in_p) == 1
    assert _refs(in_p[0], "headerReference") == []
    pg = body_sp.find(_wq("pgNumType"))
    assert pg is not None and pg.get(_wq("start")) == "1"


def test_apply_page_setup_degraded_still_reports_keys(tmp_path, monkeypatch):
    """参考缺失降级路径：新增报告键仍在（sections/section_page_numbers/
    warnings），行为不变不报错。"""
    monkeypatch.setattr(config, "BID_HEADER_SOURCE_PATH",
                        tmp_path / "不存在.docx")
    p = tmp_path / "out.docx"
    Document().save(str(p))
    result = apply_page_setup(str(p))
    assert result["degraded"]
    assert result["sections"] >= 1
    assert isinstance(result["section_page_numbers"], list)
    assert result["warnings"] == []
