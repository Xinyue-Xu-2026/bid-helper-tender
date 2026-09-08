"""模板章节解析与篇幅统计测试（python-docx 构造模板，零网络）。"""
import pytest
from docx import Document

from app.core.template_outline import (
    DEFAULT_TARGET_CHARS,
    TemplateOutlineError,
    build_tree,
    para_heading_level,
    parse_outline,
    section_target_chars,
)


def _make_heading_template(path):
    """Heading 1/2/3 模板：第一章正文 2000 字（含 1.1 的 1500 与 1.2 的 500）。"""
    doc = Document()
    doc.add_heading("第一章 项目概述", level=1)
    doc.add_heading("1.1 项目背景", level=2)
    doc.add_paragraph("背" * 1500)
    doc.add_heading("1.2 建设目标", level=2)
    doc.add_paragraph("标" * 500)
    doc.add_heading("第二章 技术方案", level=1)
    doc.add_heading("2.1 总体设计", level=2)
    doc.add_heading("2.1.1 架构设计", level=3)
    doc.add_paragraph("构" * 300)
    doc.save(str(path))


def _make_plain_template(path):
    """无标题样式的模板：正文段落按文本模式兜底识别。"""
    doc = Document()
    doc.add_paragraph("第一章 资格审查资料")
    doc.add_paragraph("一、营业执照")
    doc.add_paragraph("（一）副本复印件")
    doc.add_paragraph("这里是正文内容。" * 10)
    doc.add_paragraph("二、资质证书")
    doc.save(str(path))


def test_parse_outline_heading_styles(tmp_path):
    path = tmp_path / "tpl.docx"
    _make_heading_template(path)
    outline = parse_outline(str(path))
    assert outline == [
        (1, "第一章 项目概述"),
        (2, "1.1 项目背景"),
        (2, "1.2 建设目标"),
        (1, "第二章 技术方案"),
        (2, "2.1 总体设计"),
        (3, "2.1.1 架构设计"),
    ]


def test_parse_outline_text_pattern_fallback(tmp_path):
    path = tmp_path / "plain.docx"
    _make_plain_template(path)
    outline = parse_outline(str(path))
    assert outline == [
        (1, "第一章 资格审查资料"),
        (2, "一、营业执照"),
        (3, "（一）副本复印件"),
        (2, "二、资质证书"),
    ]


def test_parse_outline_no_headings_422(tmp_path):
    doc = Document()
    doc.add_paragraph("这段文字没有任何标题结构，只是普通段落。" * 5)
    path = tmp_path / "empty.docx"
    doc.save(str(path))
    with pytest.raises(TemplateOutlineError, match="模板中未识别到章节标题，请检查模板样式"):
        parse_outline(str(path))


def test_build_tree_nests_by_level():
    tree = build_tree([(1, "A"), (2, "A1"), (3, "A1a"), (2, "A2"), (1, "B")])
    assert [n["title"] for n in tree] == ["A", "B"]
    assert [n["title"] for n in tree[0]["children"]] == ["A1", "A2"]
    assert tree[0]["children"][0]["children"][0]["title"] == "A1a"
    assert tree[0]["level"] == 1
    assert tree[0]["children"][0]["level"] == 2
    assert tree[0]["children"][0]["children"][0]["level"] == 3


def test_section_target_chars_measured(tmp_path):
    path = tmp_path / "tpl.docx"
    _make_heading_template(path)
    # 叶子节：本标题到下一同级标题之间
    assert section_target_chars(str(path), "1.1 项目背景") == 1500
    # 章：含子节正文（1500 + 500）
    assert section_target_chars(str(path), "第一章 项目概述") == 2000
    assert section_target_chars(str(path), "2.1.1 架构设计", level=3) == 300


def test_section_target_chars_missing_fallback(tmp_path):
    path = tmp_path / "tpl.docx"
    _make_heading_template(path)
    # 找不到的二级节：同级（level 2）平均 = (1500+500+300)/3 = 766
    target = section_target_chars(str(path), "不存在的章节", level=2)
    assert target == (1500 + 500 + 300) // 3
    # 不给 level：全体标题跨度平均 = (2000+1500+500+300+300+300)/6 = 816
    assert section_target_chars(str(path), "不存在的章节") == (2000 + 1500 + 500 + 300 + 300 + 300) // 6


def test_section_target_chars_default(tmp_path):
    doc = Document()
    doc.add_heading("第一章 空章", level=1)
    path = tmp_path / "bare.docx"
    doc.save(str(path))
    assert section_target_chars(str(path), "不存在", level=5) == DEFAULT_TARGET_CHARS


# ---------- TOC（目录）污染 ----------

def _make_toc_polluted_template(path):
    """目录条目（toc 样式 / 制表符页码 / TOC 域代码）+ 正文真实标题混合。"""
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml.ns import qn
    doc = Document()
    # 目录区：toc 1/2 样式段落（带制表符页码）
    doc.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
    doc.styles.add_style("toc 2", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("第一章 项目概述\t1", style="toc 1")
    doc.add_paragraph("1.1 项目背景\t1", style="toc 2")
    doc.add_paragraph("第二章 技术方案\t5", style="toc 1")
    # 目录区：Normal 样式但文本以 制表符+页码 结尾（文本模式兜底也不能误收）
    doc.add_paragraph("一、审计质量目标\t1")
    doc.add_paragraph("（一）政府审计核减控制目标\t1")
    # 目录区：含 TOC 域代码的段落（文本无制表符页码特征）
    p = doc.add_paragraph("第五章 域代码目录条目  9")
    run = p._p.makeelement(qn("w:r"), {})
    instr = p._p.makeelement(qn("w:instrText"), {})
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    run.append(instr)
    p._p.append(run)
    # 正文：真实标题与正文
    doc.add_heading("第一章 项目概述", level=1)
    doc.add_paragraph("正" * 100)
    doc.add_heading("1.1 项目背景", level=2)
    doc.add_paragraph("背" * 50)
    doc.add_heading("第二章 技术方案", level=1)
    doc.add_paragraph("设" * 80)
    doc.save(str(path))


def test_parse_outline_skips_toc_paragraphs(tmp_path):
    path = tmp_path / "toc.docx"
    _make_toc_polluted_template(path)
    outline = parse_outline(str(path))
    assert outline == [
        (1, "第一章 项目概述"),
        (2, "1.1 项目背景"),
        (1, "第二章 技术方案"),
    ]
    # 无制表符页码尾巴
    assert all("\t" not in t for _, t in outline)


def test_toc_paragraphs_not_counted_as_body_chars(tmp_path):
    path = tmp_path / "toc.docx"
    _make_toc_polluted_template(path)
    # 目录条目字数不计入正文章节篇幅：第一章跨度 = 100 + 50
    assert section_target_chars(str(path), "第一章 项目概述") == 150
    assert section_target_chars(str(path), "第二章 技术方案") == 80


def test_strip_tab_pageno_defensive():
    from app.core.template_outline import _strip_tab_pageno
    assert _strip_tab_pageno("第三章 目标控制措施和手段\t1") == "第三章 目标控制措施和手段"
    assert _strip_tab_pageno("（一）政府审计核减控制目标\t12 ") == "（一）政府审计核减控制目标"
    assert _strip_tab_pageno("一、审计质量目标") == "一、审计质量目标"


# ---------- para_heading_level（公开化的标题判定） ----------

def test_para_heading_level_heading_style():
    doc = Document()
    doc.add_heading("第一章 项目概述", level=1)
    doc.add_heading("1.1 项目背景", level=2)
    doc.add_heading("1.1.1 架构设计", level=3)
    assert para_heading_level(doc.paragraphs[0]) == 1
    assert para_heading_level(doc.paragraphs[1]) == 2
    assert para_heading_level(doc.paragraphs[2]) == 3


def test_para_heading_level_text_fallback():
    doc = Document()
    doc.add_paragraph("第三章 投标文件格式")
    doc.add_paragraph("一、投标函")
    doc.add_paragraph("（一）投标函附录")
    doc.add_paragraph("1.2 建设目标")
    assert para_heading_level(doc.paragraphs[0]) == 1
    assert para_heading_level(doc.paragraphs[1]) == 2
    assert para_heading_level(doc.paragraphs[2]) == 3
    assert para_heading_level(doc.paragraphs[3]) == 2


def test_para_heading_level_long_body_returns_zero():
    doc = Document()
    doc.add_paragraph("这里是超过六十字的正文内容" + "长" * 60)
    doc.add_paragraph("普通正文")
    assert para_heading_level(doc.paragraphs[0]) == 0
    assert para_heading_level(doc.paragraphs[1]) == 0


def test_para_heading_level_toc_para_always_zero():
    from docx.enum.style import WD_STYLE_TYPE
    doc = Document()
    doc.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("第三章 投标文件格式\t1", style="toc 1")
    doc.add_paragraph("一、审计质量目标\t1")  # 制表符页码特征
    assert para_heading_level(doc.paragraphs[0]) == 0
    assert para_heading_level(doc.paragraphs[1]) == 0
