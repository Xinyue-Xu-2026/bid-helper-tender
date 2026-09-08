"""防篡改校验测试（T8）：产物 vs 底稿——正常产物 ok；篡改正文段 /
未绑定表格被逮；绑定表改动不报；编号/日期替换与 TOC 替换不误报。
fixture 构造法复用 test_bid_draft_export。"""
from docx import Document

from app.core.bid_draft_exporter import fill_draft
from app.core.bid_verify import verify_draft_fill
from tests.test_bid_draft_export import PROMISE_TEXT, _bindings, _build_draft, _data

BOUND = {0, 1, 2, 3}
NO_REPLACE = {"project_no": "", "project_name": "", "doc_date": ""}


def _produce(tmp_path, bindings=None, data=None, **replace):
    """用 fill_draft 产出 (draft_path, out_path)。"""
    draft = _build_draft(tmp_path / "draft.docx")
    out = str(tmp_path / "out.docx")
    fill_draft(draft, out, bindings or _bindings(),
               data if data is not None else _data(), **replace)
    return draft, out


# ---------- 1. 正常导出产物 → ok ----------

def test_clean_output_ok(tmp_path):
    draft, out = _produce(tmp_path)
    result = verify_draft_fill(draft, out, BOUND, NO_REPLACE)
    assert result["ok"] is True
    assert result["issues"] == []
    assert result["checked_tables"] == 1        # 仅承诺函表未绑定
    assert result["checked_paragraphs"] == 5    # 3 封面 + 2 toc


# ---------- 2. 篡改产物未绑定正文段 → 逮住 ----------

def test_tamper_paragraph_detected(tmp_path):
    draft, out = _produce(tmp_path)
    doc = Document(out)
    para = doc.paragraphs[2]  # 封面日期段（未绑定区域）
    para.runs[0].text = "日      期 ：2099年9月9日"
    doc.save(out)
    result = verify_draft_fill(draft, out, BOUND, NO_REPLACE)
    assert result["ok"] is False
    assert any("段落[2]" in issue for issue in result["issues"])


# ---------- 3. 篡改产物未绑定表格单元格 → 逮住 ----------

def test_tamper_unbound_table_cell_detected(tmp_path):
    draft, out = _produce(tmp_path)
    doc = Document(out)
    doc.tables[4].cell(0, 0).text = "被篡改的承诺"
    doc.save(out)
    result = verify_draft_fill(draft, out, BOUND, NO_REPLACE)
    assert result["ok"] is False
    assert any("表格[4]" in issue for issue in result["issues"])


# ---------- 4. 改动绑定表 → 不产生 issue ----------

def test_bound_table_change_no_issue(tmp_path):
    draft, out = _produce(tmp_path)
    doc = Document(out)
    doc.tables[0].cell(1, 1).text = "绑定表随意改动"
    doc.save(out)
    result = verify_draft_fill(draft, out, BOUND, NO_REPLACE)
    assert result["ok"] is True
    assert result["issues"] == []


# ---------- 5. 编号/名称/日期替换不误报 ----------

def test_stale_replacement_no_false_positive(tmp_path):
    params = {"project_no": "NEW-2026-123", "project_name": "新项目XYZ",
              "doc_date": "2026-09-08"}
    draft, out = _produce(tmp_path, **params)
    # 确认替换确实发生（旧值在产物中已无残留）
    out_texts = [p.text for p in Document(out).paragraphs]
    assert "编       号：NEW-2026-123号" in out_texts
    assert not any("OLD-2020-001" in t for t in out_texts)
    result = verify_draft_fill(draft, out, BOUND, params)
    assert result["ok"] is True
    assert result["issues"] == []


# ---------- 6. swap_toc=True 且产物 toc 段已替换 → 不误报 ----------

def test_swap_toc_no_false_positive(tmp_path):
    bindings = _bindings(swap_toc=True)
    draft, out = _produce(tmp_path, bindings=bindings)
    # 确认 toc 段已换成 TOC 域
    toc_paras = [p for p in Document(out).paragraphs
                 if (p.style.name or "").lower().startswith("toc")]
    assert len(toc_paras) == 1 and "instrText" in toc_paras[0]._element.xml
    result = verify_draft_fill(draft, out, BOUND, NO_REPLACE, swapped_toc=True)
    assert result["ok"] is True
    assert result["issues"] == []
    # 承诺函表仍未被触碰
    assert Document(out).tables[4].cell(0, 0).text == PROMISE_TEXT
