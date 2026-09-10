"""底稿表格角色分类器测试：程序化构造 docx（零网络），
覆盖人员表（标准/交通局式）、业绩表、报价表、图片占位表、
简历键值表、跨行表头、普通表 ignore、上下文 scope/person 建议、
横向合并表头去重。"""
from docx import Document

from app.core.bid_table_classifier import classify_tables


def _make_table(doc, rows_data):
    """按二维文本数组造表并填充单元格。"""
    table = doc.add_table(rows=len(rows_data), cols=len(rows_data[0]))
    for r, row in enumerate(rows_data):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    return table


def _save(doc, tmp_path, name="draft.docx"):
    path = tmp_path / name
    doc.save(str(path))
    return str(path)


# ---------- 1. 标准 5 列人员表 ----------

def test_person_roster_five_columns(tmp_path):
    doc = Document()
    _make_table(doc, [
        ["序号", "姓名", "职称", "专业工作年限", "执业资格"],
        ["1", "张三", "高级工程师", "10", "造价工程师"],
    ])
    result = classify_tables(_save(doc, tmp_path))
    assert len(result) == 1
    item = result[0]
    assert item["table_index"] == 0
    assert item["role"] == "person_roster"
    assert item["columns"] == {
        "seq": 0, "name": 1, "title": 2, "work_years": 3, "certs": 4}
    assert item["person_scope"] == "all"
    assert item["confidence"] == "高"
    assert item["confirmed"] is False
    assert item["header"] == ["序号", "姓名", "职称", "专业工作年限", "执业资格"]


# ---------- 2. 交通局式人员表（性别/年龄/学历均映射） ----------

def test_person_roster_jiaotongju_style(tmp_path):
    doc = Document()
    _make_table(doc, [
        ["序号", "姓名", "岗位", "性别", "年龄", "学历", "职称"],
        ["1", "李四", "技术员", "男", "35", "本科", "工程师"],
    ])
    result = classify_tables(_save(doc, tmp_path))
    item = result[0]
    assert item["role"] == "person_roster"
    assert item["columns"]["gender"] == 3
    assert item["columns"]["age"] == 4
    assert item["columns"]["education"] == 5
    assert item["columns"]["label"] == 2   # "岗位" → 职务栏语义键 label
    assert "role" not in item["columns"]  # 无"拟派/本岗位"关键词不进 role
    assert item["confidence"] == "高"


# ---------- 3. PERF7 业绩表 ----------

def test_perf_list_seven_columns(tmp_path):
    doc = Document()
    _make_table(doc, [
        ["序号", "项目名称", "委托单位", "合同签订时间", "合同金额",
         "服务类型", "证明材料页码"],
        ["1", "某项目", "某单位", "2024-01-01", "100万", "造价咨询", "P12"],
    ])
    result = classify_tables(_save(doc, tmp_path))
    item = result[0]
    assert item["role"] == "perf_list"
    assert item["columns"]["sign_date"] == 3
    assert item["columns"]["service_type"] == 5
    assert item["columns"]["proof_page"] == 6
    assert item["columns"]["project_name"] == 1
    assert item["columns"]["client"] == 2
    assert item["perf_scope"] == "all"
    assert item["confidence"] == "高"


def test_perf_list_client_xiangmu_danwei(tmp_path):
    """业绩表委托方列关键词"项目单位"（真实招标文件 GZ511 业绩表回归：
    表头 编号|项目名称|项目单位|服务总价 原为 ignore，属关键词库缺口）。"""
    doc = Document()
    _make_table(doc, [
        ["编号", "项目名称", "项目单位", "服务总价"],
        ["1", "某项目", "某单位", "100万"],
    ])
    result = classify_tables(_save(doc, tmp_path))
    item = result[0]
    assert item["role"] == "perf_list"
    assert item["columns"]["project_name"] == 1
    assert item["columns"]["client"] == 2


# ---------- header_rows（P1：两行表头检出） ----------

def test_header_rows_two_when_split_header(tmp_path):
    """跨行表头拼接兜底触发（row0 命中<2）→ 建议项 header_rows=2。"""
    doc = Document()
    _make_table(doc, [
        ["项目", "委托", "合同金额"],
        ["名称", "单位", ""],
        ["某项目", "某单位", "100万"],
    ])
    item = classify_tables(_save(doc, tmp_path))[0]
    assert item["role"] == "perf_list"
    assert item["header_rows"] == 2


def test_header_rows_default_one(tmp_path):
    """普通单行表头 → header_rows=1。"""
    doc = Document()
    _make_table(doc, [
        ["序号", "姓名", "职称", "专业工作年限", "执业资格"],
        ["1", "张三", "高级工程师", "10", "造价工程师"],
    ])
    item = classify_tables(_save(doc, tmp_path))[0]
    assert item["role"] == "person_roster"
    assert item["header_rows"] == 1


def test_header_rows_two_when_subheader_repeats_row0(tmp_path):
    """断裂 vMerge 副表头（row1 在同网格列重复 row0 标签 ≥2 格）→ header_rows=2。
    真实底稿 17×10 人员汇总表回归：row0 命中 4 个关键词不触发拼接兜底，
    但 row1（序号/本项目任职/姓名/职称/专业 + 证书名称/级别/证号）是副表头。"""
    doc = Document()
    _make_table(doc, [
        ["序号", "本项目任职", "姓名", "职称"],
        ["序号", "本项目任职", "姓名", "证号"],
        ["", "", "", ""],
    ])
    item = classify_tables(_save(doc, tmp_path))[0]
    assert item["role"] == "person_roster"
    assert item["header_rows"] == 2


# ---------- 4. 报价表（表头含"报价" / 前两行单元格含"费率"） ----------

def test_quote_by_header_and_by_cell(tmp_path):
    doc = Document()
    _make_table(doc, [
        ["磋商响应报价表", "", ""],
        ["项目", "下浮率", "备注"],
        ["a", "b", "c"],
    ])
    result = classify_tables(_save(doc, tmp_path, "q1.docx"))
    assert result[0]["role"] == "quote"

    doc2 = Document()
    _make_table(doc2, [
        ["项目", "标准", "备注"],
        ["咨询费", "费率（‰）", ""],
        ["a", "b", "c"],
    ])
    result2 = classify_tables(_save(doc2, tmp_path, "q2.docx"))
    assert result2[0]["role"] == "quote"


# ---------- 5. 1 列图片占位表 ----------

def test_image_slot_single_column(tmp_path):
    doc = Document()
    _make_table(doc, [["注册证书扫描件"], [""]])
    result = classify_tables(_save(doc, tmp_path))
    item = result[0]
    assert item["role"] == "image_slot"
    assert item["label_kind"] == "注册证书"
    assert item["confidence"] == "高"


# ---------- 6. 2 列简历键值表 ----------

def test_lead_resume_two_column(tmp_path):
    doc = Document()
    _make_table(doc, [
        ["姓名", ""], ["性别", ""], ["年龄", ""],
        ["学历", ""], ["职称", ""], ["执业资格", ""],
    ])
    result = classify_tables(_save(doc, tmp_path))
    item = result[0]
    assert item["role"] == "lead_resume"
    assert item["columns"] == {
        "姓名": "name", "性别": "gender", "年龄": "age",
        "学历": "education", "职称": "title", "执业资格": "certs"}
    assert item["confidence"] == "高"


# ---------- 7. 跨两行表头拼接兜底 ----------

def test_perf_list_split_header_rows(tmp_path):
    doc = Document()
    _make_table(doc, [
        ["项目", "委托", "合同金额"],
        ["名称", "单位", ""],
        ["某项目", "某单位", "100万"],
    ])
    result = classify_tables(_save(doc, tmp_path))
    item = result[0]
    assert item["role"] == "perf_list"
    assert item["columns"]["project_name"] == 0
    assert item["columns"]["client"] == 1
    assert item["columns"]["amount"] == 2


# ---------- 8. 无关键词普通表 → ignore ----------
def test_plain_table_ignore(tmp_path):
    doc = Document()
    _make_table(doc, [
        ["苹果", "香蕉", "橘子"],
        ["白菜", "萝卜", "土豆"],
        ["甲", "乙", "丙"],
    ])
    result = classify_tables(_save(doc, tmp_path))
    item = result[0]
    assert item["role"] == "ignore"
    assert item["confidence"] == "低"
    assert item["columns"] == {}


# ---------- 9. 上下文标题 → scope / person 建议 ----------

def test_context_heading_scopes(tmp_path):
    doc = Document()
    doc.add_heading("项目负责人", level=2)
    _make_table(doc, [
        ["序号", "姓名", "职称", "专业工作年限", "执业资格"],
        ["1", "张三", "高工", "10", "造价师"],
    ])
    doc.add_heading("（1）项目组其他人员-张三", level=3)
    _make_table(doc, [["职称证书扫描件"], [""]])
    result = classify_tables(_save(doc, tmp_path))
    assert len(result) == 2
    roster = result[0]
    assert roster["role"] == "person_roster"
    assert roster["person_scope"] == "lead"
    assert roster["context_heading"] == "项目负责人"
    slot = result[1]
    assert slot["role"] == "image_slot"
    assert slot["label_kind"] == "职称证书"
    assert slot["person"] == "member:0"
    assert slot["context_heading"] == "（1）项目组其他人员-张三"


# ---------- 10. 横向合并表头：不炸、下标不错位 ----------

def test_merged_header_cells(tmp_path):
    doc = Document()
    table = _make_table(doc, [
        ["序号", "姓名", "职称", "专业工作年限", "执业资格"],
        ["1", "张三", "高工", "10", "造价师"],
    ])
    table.cell(0, 0).merge(table.cell(0, 1))
    result = classify_tables(_save(doc, tmp_path))
    item = result[0]
    assert item["role"] == "person_roster"
    # 合并后 序号/姓名 同占网格列 0，其余列按网格下标不错位
    assert item["columns"]["title"] == 2
    assert item["columns"]["work_years"] == 3
    assert item["columns"]["certs"] == 4
    col0_sem = "name" if item["columns"].get("name") == 0 else "seq"
    assert item["columns"][col0_sem] == 0
    assert max(item["columns"].values()) < 5


# ---------- 拟派岗位列（V1.2 Task 4） ----------

def test_role_column_keywords_include_bendan():
    from app.core.bid_table_classifier import PERSON_COL_KEYWORDS
    assert "本项目任职" in PERSON_COL_KEYWORDS["role"]
    assert "本项目职务" in PERSON_COL_KEYWORDS["role"]


def test_person_roster_maps_role_column(tmp_path):
    doc = Document()
    _make_table(doc, [
        ["姓名", "学历", "本项目任职"],
        ["张三", "硕士", "项目经理"],
    ])
    item = next(s for s in classify_tables(_save(doc, tmp_path))
                if s["role"] == "person_roster")
    assert "role" in item["columns"]
    assert item["columns"]["role"] == 2


def test_role_beats_label_on_bendan_zhiwu(tmp_path):
    """「本项目职务」含 label 关键词「职务」：长关键词优先，role 命中抢列，
    label 不得截胡（V1.2 7.6 防回归）。"""
    doc = Document()
    _make_table(doc, [
        ["姓名", "学历", "本项目职务"],
        ["张三", "硕士", "项目经理"],
    ])
    item = classify_tables(_save(doc, tmp_path))[0]
    assert item["columns"].get("role") == 2
    assert item["columns"].get("label") != 2


def test_header_rows_two_when_row0_no_data_and_row1_header(tmp_path):
    """双行表头第三规则（7.1 补强）：row0 命中表头关键词但整行无数据语义
    （无数字/金额等），且 row1 也命中表头关键词签名（≥2）→ header_rows=2。"""
    doc = Document()
    _make_table(doc, [
        ["姓名", "职称", "执业资格"],        # row0：上层表头（无数字）
        ["职务", "专业工作年限", "资格证书"],  # row1：下层表头（同样命中关键词）
        ["张三", "工程师", "造价工程师"],     # row2 起为数据
    ])
    item = classify_tables(_save(doc, tmp_path))[0]
    assert item["role"] == "person_roster"
    assert item["header_rows"] == 2


def test_export_role_column_filled(tmp_path):
    """端到端：人员表 role 列 → fill_draft 写入 person_semantics 的 role 值。"""
    from app.core.bid_draft_exporter import fill_draft
    doc = Document()
    _make_table(doc, [
        ["姓名", "学历", "本项目任职"],
        ["旧名", "旧学历", "旧岗位"],
    ])
    draft = _save(doc, tmp_path)
    out = str(tmp_path / "out.docx")
    bindings = {"tables": [{
        "table_index": 0, "role": "person_roster",
        "columns": {"name": 0, "education": 1, "role": 2},
        "person_scope": "all", "confirmed": True}], "swap_toc": False}
    data = {"persons": [{"name": "张三", "is_lead": True,
                         "sem": {"name": "张三", "education": "硕士",
                                 "role": "项目经理"}}],
            "contracts": []}
    report = fill_draft(draft, out, bindings, data)
    rows = [[c.text for c in r.cells] for r in Document(out).tables[0].rows]
    assert rows[1][2] == "项目经理"
    assert report["verify"]["ok"] is True