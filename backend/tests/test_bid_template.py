"""模板式商务标导出测试：程序化构造最小模板 fixture（七张表 + Heading 定位锚点），
覆盖表A/D 人员填充、执业资格拼接、表E 主要内容生成、表F/G 按 section 分填、
负责人业绩匹配/无匹配清空、空选择 422、模板缺失 503、服务类型映射。"""
from io import BytesIO
from urllib.parse import quote

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import config
from app.db import Database
from app.deps import get_db
from app.main import app
from app.services import bid_service

PERSON_HEADERS = ["人员安排", "姓名", "职称", "专业工作年限", "执业资格"]
PERF5_HEADERS = ["序号", "项目名称", "业主单位", "业主联系方式", "主要内容"]
PERF7_HEADERS = ["序号", "项目名称", "业主单位", "合同签订时间",
                 "项目类型", "服务类型", "证明材料页码"]


def _add_table(doc, headers, data_rows=1):
    t = doc.add_table(rows=1 + data_rows, cols=len(headers))
    for j, h in enumerate(headers):
        t.cell(0, j).text = h
    for i in range(1, 1 + data_rows):
        for j in range(len(headers)):
            t.cell(i, j).text = f"样例{i}-{j}"
    return t


def _build_template(path):
    """构造含全部定位锚点与七张表的最小模板文档（含残留的上次投标文本）。"""
    from docx.enum.style import WD_STYLE_TYPE
    doc = Document()
    # 封面残留：旧项目名/旧编号/旧日期
    doc.add_paragraph("项 目 名 称：旧项目名AAA")
    doc.add_paragraph("编       号：OLD-2020-001号")
    doc.add_paragraph("日      期 ：2020年1月2日")
    # 目录条目（toc 样式）：含日期样式文本，替换必须跳过
    doc.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("一、开标一览表\t1", style="toc 1")
    doc.add_paragraph("日 期：2020年1月2日\t3", style="toc 1")
    doc.add_heading("商务标", level=1)
    doc.add_paragraph("一、投标函")
    doc.add_heading("三、人员配备表", level=1)
    doc.add_paragraph("项目人员基本情况表")
    _add_table(doc, PERSON_HEADERS)                       # 表A
    doc.add_paragraph("项目名称：旧项目名AAA")
    doc.add_paragraph("项目编号：OLD-2020-001号")
    doc.add_paragraph("日 期：\t2020年1月2日")
    doc.add_paragraph("本单位参与本次 旧项目名AAA项目投标活动")
    doc.add_heading("2、项目负责人", level=2)
    doc.add_heading("（1）项目负责人基本情况", level=3)
    doc.add_heading("（2）项目负责人业绩", level=3)
    doc.add_paragraph("相关业绩一览表")
    _add_table(doc, PERF5_HEADERS)                        # 表B
    doc.add_paragraph("业绩证明材料表")
    _add_table(doc, PERF7_HEADERS)                        # 表C
    doc.add_heading("3、项目组人员", level=2)
    doc.add_heading("（1）项目组其他人员-旧甲", level=3)
    doc.add_heading("（2）项目组其他人员-旧乙", level=3)
    doc.add_heading("（3）社保缴纳证明", level=3)
    _add_table(doc, PERSON_HEADERS)                       # 表D
    doc.add_heading("四、相关业绩一览表", level=1)
    _add_table(doc, PERF5_HEADERS)                        # 表E
    doc.add_heading("1、2024年业绩", level=2)
    _add_table(doc, PERF7_HEADERS)                        # 表F
    doc.add_heading("2、2023年业绩", level=2)
    _add_table(doc, PERF7_HEADERS)                        # 表G
    doc.save(path)


@pytest.fixture()
def client(db_path):
    def override():
        db = Database(db_path)
        db.init_schema()
        yield db
    app.dependency_overrides[get_db] = override
    yield TestClient(app)
    app.dependency_overrides.clear()


@pytest.fixture()
def template_path(tmp_path, monkeypatch):
    path = tmp_path / "商务标模板.docx"
    _build_template(path)
    monkeypatch.setattr(config, "BID_TEMPLATE_PATH", path)
    return path


def _make_project(client, name="模板项目"):
    return client.post("/api/projects", json={"name": name}).json()["id"]


def _make_person(client, name, fields=None):
    return client.post("/api/assets", json={
        "type": "person", "name": name, "fields": fields or {}}).json()["id"]


def _make_contract(client, name, fields=None):
    return client.post("/api/assets", json={
        "type": "contract", "name": name, "fields": fields or {}}).json()["id"]


def _seed(client):
    """两名人员（张三负责）+ 三份合同（A/B section1，C section2；A/C 项目负责人=张三）。"""
    zhang = _make_person(client, "张三", {
        "职称": "高级工程师",
        "证书": [{"类型": "一级造价师", "专业": "土建", "有效期至": "2027-01-01"},
                 {"类型": "监理工程师", "有效期至": "2028-01-01"}]})
    li = _make_person(client, "李四", {
        "职称": "工程师",
        "证书": [{"类型": "一级建造师", "专业": "建筑工程"}]})
    a = _make_contract(client, "项目A", {
        "类型": "跟踪", "委托单位": "甲公司", "签订日期": "2024-03-15",
        "项目负责人": "张三"})
    b = _make_contract(client, "项目B", {
        "类型": "结算", "委托单位": "乙公司", "签订日期": "2023-01-01",
        "项目负责人": "王五"})
    c = _make_contract(client, "项目C", {
        "类型": "编标", "委托单位": "丙公司", "签订日期": "2022-06-01",
        "项目负责人": "张三"})
    return zhang, li, a, b, c


def _post(client, pid, persons, contracts):
    return client.post(f"/api/projects/{pid}/bid-assets/export-template",
                       json={"persons": persons, "contracts": contracts})


def _tables(content):
    doc = Document(BytesIO(content))
    return doc, doc.tables  # 顺序即 A..G


def _data_rows(table):
    return [[cell.text for cell in row.cells] for row in table.rows[1:]]


# ---------- 主流程 ----------

def test_export_template_full(client, template_path):
    pid = _make_project(client)
    zhang, li, ca, cb, cc = _seed(client)
    r = _post(client, pid,
              persons=[{"asset_id": li, "is_lead": False},
                       {"asset_id": zhang, "is_lead": True}],
              contracts=[{"asset_id": ca, "section": 1},
                         {"asset_id": cb, "section": 1},
                         {"asset_id": cc, "section": 2}])
    assert r.status_code == 200
    disposition = r.headers["content-disposition"]
    assert disposition.startswith("attachment; filename*=utf-8''")
    assert disposition.endswith(".docx")
    assert "%E5%95%86%E5%8A%A1%E6%A0%87" in disposition  # "商务标" 的 quote 编码
    doc, tables = _tables(r.content)
    assert len(tables) == 7
    ta, tb, tc, td, te, tf, tg = tables

    # 表A：全体人员，负责人排最前；职称/执业资格拼接（含无专业仅类型）
    rows_a = _data_rows(ta)
    assert len(rows_a) == 2
    assert rows_a[0][0] == "1.项目负责人" and rows_a[0][1] == "张三"
    assert rows_a[0][2] == "高级工程师" and rows_a[0][3] == ""
    assert "土建专业一级造价师" in rows_a[0][4] and "监理工程师" in rows_a[0][4]
    assert rows_a[1][0] == "2.项目组其他人员" and rows_a[1][1] == "李四"
    assert "建筑工程专业一级建造师" in rows_a[1][4]

    # 表D：仅非负责人，编号从 1 起
    rows_d = _data_rows(td)
    assert len(rows_d) == 1
    assert rows_d[0][0] == "1.项目组其他人员" and rows_d[0][1] == "李四"

    # 表E：全部已选合同按传入顺序，主要内容自动生成
    rows_e = _data_rows(te)
    assert len(rows_e) == 3
    assert rows_e[0][:3] == ["1", "项目A", "甲公司"]
    assert rows_e[0][4] == "提供项目A的跟踪审计服务"
    assert rows_e[1][4] == "提供项目B的结算审核服务"
    assert rows_e[2][4] == "提供项目C的编标（清单及控制价编制）服务"

    # 表F/G：按 section 分填，序号各自从 1 起
    rows_f = _data_rows(tf)
    assert [r[1] for r in rows_f] == ["项目A", "项目B"]
    assert rows_f[0] == ["1", "项目A", "甲公司", "2024-03-15", "", "跟踪审计", ""]
    assert rows_f[1][5] == "结算审核"
    rows_g = _data_rows(tg)
    assert len(rows_g) == 1
    assert rows_g[0] == ["1", "项目C", "丙公司", "2022-06-01", "",
                         "编标（清单及控制价编制）", ""]

    # 表B/C：负责人（张三）业绩 = 项目负责人字段匹配的合同（A、C）
    rows_b = _data_rows(tb)
    assert [r[1] for r in rows_b] == ["项目A", "项目C"]
    assert rows_b[0][0] == "1" and rows_b[1][0] == "2"
    rows_c = _data_rows(tc)
    assert [r[1] for r in rows_c] == ["项目A", "项目C"]
    assert rows_c[1][3] == "2022-06-01"

    # 模板其余部分原样保留
    assert any(p.text == "一、投标函" for p in doc.paragraphs)


def test_export_template_no_lead_match_clears_lead_tables(client, template_path):
    pid = _make_project(client)
    zhang, _, ca, _, _ = _seed(client)
    r = _post(client, pid,
              persons=[{"asset_id": zhang, "is_lead": True}],
              contracts=[{"asset_id": ca, "section": 1}])
    assert r.status_code == 200
    _, tables = _tables(r.content)
    # 把项目负责人改成一个不在业绩里的人：直接用无匹配合同
    # 本用例改选无负责人匹配的合同（项目B 项目负责人=王五）
    _, _, _, cb, _ = _seed(client)
    r = _post(client, pid,
              persons=[{"asset_id": zhang, "is_lead": True}],
              contracts=[{"asset_id": cb, "section": 1}])
    assert r.status_code == 200
    _, tables = _tables(r.content)
    tb, tc = tables[1], tables[2]
    assert _data_rows(tb) == []  # 无匹配：只留表头
    assert _data_rows(tc) == []


def test_export_template_section_without_records_keeps_header(client, template_path):
    pid = _make_project(client)
    _, _, ca, _, _ = _seed(client)
    r = _post(client, pid, persons=[],
              contracts=[{"asset_id": ca, "section": 1}])
    assert r.status_code == 200
    _, tables = _tables(r.content)
    assert len(_data_rows(tables[5])) == 1   # 表F 有一条
    assert _data_rows(tables[6]) == []       # 表G 只留表头
    assert _data_rows(tables[0]) == []       # 表A 无人员只留表头
    assert _data_rows(tables[3]) == []       # 表D 同


# ---------- 错误分支 ----------

def test_export_template_empty_selection_422(client, template_path):
    pid = _make_project(client)
    r = _post(client, pid, persons=[], contracts=[])
    assert r.status_code == 422


def test_export_template_project_not_found(client, template_path):
    r = _post(client, 99999, persons=[{"asset_id": 1, "is_lead": True}],
              contracts=[])
    assert r.status_code == 404


def test_export_template_missing_503(client, monkeypatch, tmp_path):
    monkeypatch.setattr(config, "BID_TEMPLATE_PATH", tmp_path / "不存在.docx")
    pid = _make_project(client)
    zhang = _make_person(client, "张三")
    r = _post(client, pid, persons=[{"asset_id": zhang, "is_lead": True}],
              contracts=[])
    assert r.status_code == 503
    assert "模板" in r.json()["detail"]


# ---------- 残留文本替换（project_no / project_name / doc_date / 人员子标题） ----------

def _post_full(client, pid, persons, contracts, **extra):
    payload = {"persons": persons, "contracts": contracts}
    payload.update(extra)
    return client.post(f"/api/projects/{pid}/bid-assets/export-template",
                       json=payload)


def test_export_template_replaces_stale_text(client, template_path):
    """新字段全给：编号/名称/日期替换 + 人员子标题改名 + 目录日期不动。"""
    pid = _make_project(client, name="库项目名")
    zhang, li, ca, _, _ = _seed(client)
    r = _post_full(client, pid,
                   persons=[{"asset_id": li, "is_lead": False},
                            {"asset_id": zhang, "is_lead": True}],
                   contracts=[{"asset_id": ca, "section": 1}],
                   project_no="NEW-2026-123",
                   project_name="新项目XYZ",
                   doc_date="2026-09-08")
    assert r.status_code == 200
    # 下载文件名使用覆盖后的项目名
    assert quote("新项目XYZ") in r.headers["content-disposition"]

    doc = Document(BytesIO(r.content))
    texts = [p.text for p in doc.paragraphs]
    # 项目编号：封面与正文均替换，保留模板尾字「号」风格
    assert "编       号：NEW-2026-123号" in texts
    assert "项目编号：NEW-2026-123号" in texts
    # 项目名称：封面（字间空格）/正文标签行/承诺书裸名句子均替换
    assert "项 目 名 称：新项目XYZ" in texts
    assert "项目名称：新项目XYZ" in texts
    assert "本单位参与本次 新项目XYZ项目投标活动" in texts
    # 日期：封面与页脚均替换为 Y年M月D日（月日不补零），保留标签与原空白
    assert "日      期 ：2026年9月8日" in texts
    assert "日 期：\t2026年9月8日" in texts
    # 旧文本无残留
    assert not any("OLD-2020-001" in t for t in texts)
    assert not any("旧项目名AAA" in t for t in texts)
    # 人员子标题：非负责人按顺序改名，多余标题保持原样，社保标题不动
    assert "（1）项目组其他人员-李四" in texts
    assert "（2）项目组其他人员-旧乙" in texts
    assert "（3）社保缴纳证明" in texts
    # 目录（toc 样式）日期行不受影响
    toc_texts = [p.text for p in doc.paragraphs
                 if (p.style.name or "").lower().startswith("toc")]
    assert any("2020年1月2日" in t for t in toc_texts)
    assert not any("2026年9月8日" in t for t in toc_texts)


def test_export_template_optional_fields_empty_keep_stale(client, template_path):
    """新字段缺省：编号/日期不动，项目名称回落为数据库项目名。"""
    pid = _make_project(client, name="库项目名")
    zhang, _, ca, _, _ = _seed(client)
    r = _post_full(client, pid,
                   persons=[{"asset_id": zhang, "is_lead": True}],
                   contracts=[{"asset_id": ca, "section": 1}])
    assert r.status_code == 200
    # 文件名回落为数据库项目名
    assert quote("库项目名") in r.headers["content-disposition"]

    doc = Document(BytesIO(r.content))
    texts = [p.text for p in doc.paragraphs]
    # project_no / doc_date 为空：模板残留原样保留
    assert "编       号：OLD-2020-001号" in texts
    assert "日      期 ：2020年1月2日" in texts
    # project_name 为空：以数据库项目名替换残留旧名
    assert "项 目 名 称：库项目名" in texts
    assert "本单位参与本次 库项目名项目投标活动" in texts
    assert not any("旧项目名AAA" in t for t in texts)


def test_export_template_member_headings_more_members_than_headings(
        client, template_path):
    """非负责人多于子标题时：标题改完即止，不报错。"""
    pid = _make_project(client)
    zhang, li, _, _, _ = _seed(client)
    wang = _make_person(client, "王五", {"职称": "工程师"})
    zhao = _make_person(client, "赵六", {"职称": "工程师"})
    r = _post_full(client, pid,
                   persons=[{"asset_id": zhang, "is_lead": True},
                            {"asset_id": li}, {"asset_id": wang},
                            {"asset_id": zhao}],
                   contracts=[])
    assert r.status_code == 200
    doc = Document(BytesIO(r.content))
    texts = [p.text for p in doc.paragraphs]
    assert "（1）项目组其他人员-李四" in texts
    assert "（2）项目组其他人员-王五" in texts
    # 赵六没有对应标题，仅出现在表A/表D 数据行中
    rows_a = _data_rows(doc.tables[0])
    assert [row[1] for row in rows_a] == ["张三", "李四", "王五", "赵六"]


def test_replace_stale_text_in_table_cells(tmp_path):
    """表格单元格内的残留编号/名称/日期同样替换（直接调导出器）。"""
    from app.core.bid_template_exporter import build_bid_docx_from_template
    doc = Document()
    doc.add_paragraph("项目编号：OLD-2020-001号")
    doc.add_paragraph("项目名称：旧项目名AAA")  # 带标签的旧名，供运行时提取
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "项目编号"
    t.cell(0, 1).text = "OLD-2020-001号"
    t.cell(1, 0).text = "项目名称"
    t.cell(1, 1).text = "旧项目名AAA"
    src = tmp_path / "src.docx"
    doc.save(src)
    dest = tmp_path / "out.docx"
    build_bid_docx_from_template(str(src), str(dest), {
        "project_no": "NEW-2026-123",
        "project_name": "新项目XYZ",
        "doc_date": "2026-09-08",
    })
    out = Document(dest)
    assert out.tables[0].cell(0, 1).text == "NEW-2026-123号"
    assert out.tables[0].cell(1, 1).text == "新项目XYZ"
    assert out.paragraphs[0].text == "项目编号：NEW-2026-123号"


# ---------- 服务类型映射 ----------
def test_service_type_label():
    assert bid_service.service_type_label({"类型": "跟踪"}) == "跟踪审计"
    assert bid_service.service_type_label({"类型": "结算"}) == "结算审核"
    assert bid_service.service_type_label({"类型": "编标"}) == "编标（清单及控制价编制）"
    assert bid_service.service_type_label({"类型": "审标"}) == "审标"
    assert bid_service.service_type_label({"类型": "水利审计"}) == "水利审计"
    assert bid_service.service_type_label({"类型": "中标通知书"}) == "中标"
    assert bid_service.service_type_label({"类型": "其他"}) == "其他"
    assert bid_service.service_type_label({}) == ""
