# -*- coding: utf-8 -*-
"""保格式仿写路由测试（规划 mock，应用真实执行，零网络）。"""
import pytest
from fastapi.testclient import TestClient

from app.main import app


@pytest.fixture()
def client():
    with TestClient(app) as c:
        yield c


def _project(client, name="仿写测试项目"):
    return client.post("/api/projects", json={"name": name}).json()["id"]


def _make_mimic_template(path):
    from docx import Document
    doc = Document()
    doc.add_heading("第一章 工作制度", level=1)              # 0
    doc.add_heading("第一条 总则", level=3)                  # 1
    doc.add_paragraph("本公司实行项目经理负责制。")           # 2
    doc.add_paragraph("本制度适用于所有造价咨询项目。")       # 3
    doc.add_heading("第二条 工作程序", level=3)              # 4
    doc.add_paragraph("项目组应按计划开展现场作业。")         # 5
    table = doc.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "制度名称"
    table.rows[0].cells[1].text = "主要内容"
    table.rows[0].cells[2].text = "备注"
    table.rows[1].cells[0].text = "例会制度"
    table.rows[1].cells[1].text = "每周召开"
    table.rows[1].cells[2].text = ""
    doc.save(str(path))


def _register_template(tmp_path):
    from app.db import Database
    path = tmp_path / "模板.docx"
    _make_mimic_template(path)
    db = Database(); db.init_schema()
    return db.create_template("仿写模板", str(path)), path


def _set_key(monkeypatch):
    monkeypatch.setattr("app.services.mimic_service.get_api_key", lambda: "sk-test")
    monkeypatch.setattr("app.services.mimic_service.get_model", lambda: "kimi-k3")


def test_mimic_plan(client, monkeypatch, tmp_path):
    pid = _project(client)
    tid, _ = _register_template(tmp_path)
    _set_key(monkeypatch)
    fake_ops = [
        {"op": "insert", "index": 3, "anchor": "本制度适用于",
         "items": [{"style": "Heading 3", "text": "第三条 专项制度"}]},
        {"op": "set_text", "index": 5, "anchor": "项目组应按计划",
         "text": "项目组应按计划开展现场作业，误差率≤2%。"},
        {"op": "add_row", "table": 0, "match": "例会制度",
         "cells": ["专项制度", "针对本项目", ""]},
    ]
    monkeypatch.setattr("app.services.mimic_service.plan_edits",
                        lambda a, r, p, k, m: fake_ops)
    r = client.post(f"/api/projects/{pid}/mimic/plan", json={"template_id": tid})
    assert r.status_code == 200
    body = r.json()
    assert body["ops"] == fake_ops
    assert "3 处修改" in body["summary"]
    assert body["placeholder_count"] == 0


def test_mimic_apply_and_download(client, monkeypatch, tmp_path):
    pid = _project(client)
    tid, _ = _register_template(tmp_path)
    _set_key(monkeypatch)
    ops = [
        {"op": "insert", "index": 3, "anchor": "本制度适用于",
         "items": [{"style": "Heading 3", "text": "第三条 专项制度"},
                   {"style": "Normal", "text": "针对本项目专项规定。"}]},
        {"op": "set_text", "index": 5, "anchor": "项目组应按计划",
         "text": "项目组应按计划开展现场作业。"},
        {"op": "add_row", "table": 0, "match": "例会制度",
         "cells": ["专项制度", "针对本项目", ""]},
    ]
    r = client.post(f"/api/projects/{pid}/mimic/apply",
                    json={"template_id": tid, "ops": ops})
    assert r.status_code == 200
    body = r.json()
    assert body["output_file"].startswith(f"mimic_{pid}_")
    assert body["report"]["drift"] == 0
    assert body["report"]["invalid_new"] == 0

    d = client.get(body["download_url"])
    assert d.status_code == 200

    from app import config
    from docx import Document
    out = config.FILES_DIR / "mimic_outputs" / body["output_file"]
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    assert any("第三条 专项制度" in t for t in texts)
    assert any("针对本项目专项规定" in t for t in texts)


def test_mimic_apply_anchor_mismatch_422(client, monkeypatch, tmp_path):
    pid = _project(client)
    tid, _ = _register_template(tmp_path)
    _set_key(monkeypatch)
    ops = [{"op": "insert", "index": 3, "anchor": "完全不匹配的开头",
            "items": [{"style": "Heading 3", "text": "第三条"}]}]
    r = client.post(f"/api/projects/{pid}/mimic/apply",
                    json={"template_id": tid, "ops": ops})
    assert r.status_code == 422
    assert "锚点不匹配" in r.json()["detail"]


def test_mimic_plan_missing_project_404(client, tmp_path):
    tid, _ = _register_template(tmp_path)
    r = client.post("/api/projects/999/mimic/plan", json={"template_id": tid})
    assert r.status_code == 404


def test_mimic_plan_missing_template_404(client):
    pid = _project(client)
    r = client.post(f"/api/projects/{pid}/mimic/plan", json={"template_id": 999})
    assert r.status_code == 404


def test_mimic_plan_no_key_502(client, monkeypatch, tmp_path):
    pid = _project(client)
    tid, _ = _register_template(tmp_path)
    monkeypatch.setattr("app.services.mimic_service.get_api_key", lambda: "")
    r = client.post(f"/api/projects/{pid}/mimic/plan", json={"template_id": tid})
    assert r.status_code == 502
