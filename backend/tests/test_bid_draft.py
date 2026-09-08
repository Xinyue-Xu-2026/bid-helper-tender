"""商务标底稿裁切测试：程序化构造招标文件 docx（零网络），
覆盖格式章节定位（样式/文本模式/多命中取最后/无命中）、
底稿裁切（区间保留、表格计数、parse_outline 可解析）、
空区间与标题解析错误。"""
import pytest
from docx import Document

from app.core.bid_draft import (
    BidDraftError,
    cut_draft,
    find_format_chapter,
    list_headings,
    resolve_heading_index,
)
from app.core.template_outline import parse_outline


def _build_tender(path):
    """标准招标文件：第一章(H1) + 正文 + 第三章 投标文件格式(H1)
    + 一、投标函(H2 无样式文本模式) + 1 张表 + 第四章 合同条款(H1)。"""
    doc = Document()
    doc.add_heading("第一章 招标公告", level=1)
    doc.add_paragraph("招标公告正文内容若干。")
    doc.add_paragraph("更多公告正文。")
    doc.add_heading("第三章 投标文件格式", level=1)
    doc.add_paragraph("一、投标函")  # 无样式段落，文本模式兜底判 H2
    doc.add_paragraph("投标函正文内容。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "表格单元格甲"
    doc.add_heading("第四章 合同条款", level=1)
    doc.add_paragraph("合同条款正文内容。")
    doc.save(str(path))


def _all_text(docx_path):
    """产物全文（段落 + 顶层表格单元格）拼接，便于包含性断言。"""
    doc = Document(docx_path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


# ---------- find_format_chapter / list_headings ----------

def test_find_format_chapter_hits_third_chapter(tmp_path):
    path = tmp_path / "tender.docx"
    _build_tender(path)
    headings = list_headings(str(path))
    found = find_format_chapter(headings)
    assert found is not None
    assert found["start"]["title"] == "第三章 投标文件格式"
    assert found["start"]["level"] == 1
    assert found["end"] is not None
    assert found["end"]["title"] == "第四章 合同条款"
    # index 坐标系：body 子元素下标（含正文段/表格占位）
    assert found["start"]["index"] < found["end"]["index"]


def test_find_format_chapter_no_keyword_returns_none(tmp_path):
    doc = Document()
    doc.add_heading("第一章 招标公告", level=1)
    doc.add_heading("第二章 投标人须知", level=1)
    path = tmp_path / "tender.docx"
    doc.save(str(path))
    assert find_format_chapter(list_headings(str(path))) is None


def test_find_format_chapter_takes_last_match(tmp_path):
    """前文摘要出现关键词 + 文末真章节 → 取最后一个命中。"""
    doc = Document()
    doc.add_heading("第一章 招标公告", level=1)
    doc.add_heading("第二章 投标文件格式摘要", level=1)  # 摘要里的关键词命中
    doc.add_paragraph("摘要正文。")
    doc.add_heading("第三章 投标文件格式", level=1)      # 真正的格式章节
    doc.add_paragraph("一、投标函")
    path = tmp_path / "tender.docx"
    doc.save(str(path))
    found = find_format_chapter(list_headings(str(path)))
    assert found is not None
    assert found["start"]["title"] == "第三章 投标文件格式"


def test_find_format_chapter_text_pattern_fallback(tmp_path):
    """全文无 Heading 样式（纯文本"第三章 …"）→ 文本模式仍能命中。"""
    doc = Document()
    doc.add_paragraph("第一章 招标公告")
    doc.add_paragraph("公告正文内容若干。")
    doc.add_paragraph("第三章 投标文件格式")
    doc.add_paragraph("一、投标函")
    doc.add_paragraph("第四章 合同条款")
    path = tmp_path / "tender.docx"
    doc.save(str(path))
    headings = list_headings(str(path))
    found = find_format_chapter(headings)
    assert found is not None
    assert found["start"]["title"] == "第三章 投标文件格式"
    assert found["start"]["level"] == 1
    assert found["end"]["title"] == "第四章 合同条款"


def test_find_format_chapter_end_none_when_last_chapter(tmp_path):
    """格式章节为最后一章 → end 为 None（裁切到文档末尾）。"""
    doc = Document()
    doc.add_heading("第一章 招标公告", level=1)
    doc.add_paragraph("公告正文。")
    doc.add_heading("第三章 投标文件格式", level=1)
    doc.add_paragraph("一、投标函")
    doc.add_paragraph("投标函正文。")
    path = tmp_path / "tender.docx"
    doc.save(str(path))
    found = find_format_chapter(list_headings(str(path)))
    assert found is not None
    assert found["start"]["title"] == "第三章 投标文件格式"
    assert found["end"] is None


# ---------- cut_draft ----------

def test_cut_draft_keeps_format_chapter_only(tmp_path):
    src = tmp_path / "tender.docx"
    _build_tender(src)
    dest = tmp_path / "draft.docx"
    headings = list_headings(str(src))
    found = find_format_chapter(headings)
    result = cut_draft(str(src), str(dest),
                       found["start"]["index"], found["end"]["index"])
    text = _all_text(str(dest))
    # 区间内内容保留
    assert "第三章 投标文件格式" in text
    assert "一、投标函" in text
    assert "投标函正文内容。" in text
    assert "表格单元格甲" in text
    # 区间外内容删除
    assert "第一章 招标公告" not in text
    assert "招标公告正文内容若干。" not in text
    assert "第四章 合同条款" not in text
    assert "合同条款正文内容。" not in text
    # 返回统计
    assert result["tables"] == 1
    assert (1, "第三章 投标文件格式") in result["outline"]
    assert (2, "一、投标函") in result["outline"]
    assert isinstance(result["toc_paragraphs"], int)
    # 产物大纲可解析
    outline = parse_outline(str(dest))
    assert (1, "第三章 投标文件格式") in outline


def test_cut_draft_to_document_end(tmp_path):
    """end_index=-1：格式章节为最后一章时裁切到文档末尾。"""
    doc = Document()
    doc.add_heading("第一章 招标公告", level=1)
    doc.add_paragraph("公告正文。")
    doc.add_heading("第三章 投标文件格式", level=1)
    doc.add_paragraph("一、投标函")
    doc.add_paragraph("投标函正文。")
    src = tmp_path / "tender.docx"
    doc.save(str(src))
    dest = tmp_path / "draft.docx"
    headings = list_headings(str(src))
    found = find_format_chapter(headings)
    assert found["end"] is None
    result = cut_draft(str(src), str(dest), found["start"]["index"])
    text = _all_text(str(dest))
    assert "第三章 投标文件格式" in text
    assert "投标函正文。" in text
    assert "第一章 招标公告" not in text
    assert result["tables"] == 0


def test_cut_draft_empty_range_raises(tmp_path):
    src = tmp_path / "tender.docx"
    _build_tender(src)
    dest = tmp_path / "draft.docx"
    headings = list_headings(str(src))
    start = resolve_heading_index(headings, "第三章 投标文件格式")
    with pytest.raises(BidDraftError, match="裁切区间为空"):
        cut_draft(str(src), str(dest), start, start)


# ---------- resolve_heading_index ----------

def test_resolve_heading_index_found(tmp_path):
    src = tmp_path / "tender.docx"
    _build_tender(src)
    headings = list_headings(str(src))
    idx = resolve_heading_index(headings, " 第三章 投标文件格式 ")
    assert idx == resolve_heading_index(headings, "第三章 投标文件格式")


def test_resolve_heading_index_missing_raises(tmp_path):
    src = tmp_path / "tender.docx"
    _build_tender(src)
    headings = list_headings(str(src))
    with pytest.raises(BidDraftError, match="未找到标题"):
        resolve_heading_index(headings, "第九章 不存在的章节")
