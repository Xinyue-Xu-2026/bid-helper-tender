from docx import Document
from docx.oxml.ns import qn

from app.core.doc_blocks import apply_edits, read_blocks


def _doc_with_table_merge():
    doc = Document()
    doc.add_paragraph("第一段")
    doc.add_paragraph("第二段")
    t = doc.add_table(2, 3)
    t.cell(0, 0).merge(t.cell(0, 1))
    t.cell(0, 0).text = "合并格"
    t.cell(0, 2).text = "右格"
    t.cell(1, 0).text = "A"
    t.cell(1, 1).text = "B"
    t.cell(1, 2).text = "C"
    return doc


def test_read_blocks_order_and_merge_grid(tmp_path):
    p = tmp_path / "d.docx"
    _doc_with_table_merge().save(str(p))
    r = read_blocks(str(p))
    assert r["n_tables"] == 1
    kinds = [b["kind"] for b in r["blocks"]]
    assert kinds == ["paragraph", "paragraph", "table"]
    tbl = r["blocks"][2]
    assert tbl["index"] == 0
    row0 = tbl["rows"][0]
    assert row0[0]["text"] == "合并格" and row0[0]["colspan"] == 2
    assert row0[0]["origin"] is True
    assert row0[1]["text"] == "右格" and row0[1]["colspan"] == 1
    assert len(tbl["rows"][1]) == 3  # A/B/C


def test_apply_edits_text_and_structure_preserved(tmp_path):
    src = tmp_path / "d.docx"
    _doc_with_table_merge().save(str(src))
    out = tmp_path / "o.docx"
    res = apply_edits(
        str(src), str(out),
        paragraphs={0: "改第一段"},
        tables={0: [["新合并", "新右"], ["A2", "B2", "C2"]]})
    assert res["paragraphs"] == 1 and res["cells"] == 5
    doc = Document(str(out))
    assert doc.paragraphs[0].text == "改第一段"
    assert doc.paragraphs[1].text == "第二段"  # 未改
    t = doc.tables[0]
    assert t.cell(0, 0).text == "新合并"
    assert t.cell(0, 2).text == "新右"
    assert t.cell(1, 0).text == "A2"
    assert t.cell(1, 2).text == "C2"
    # 结构保留：合并格仍是 gridSpan=2（row0 的 cells[0] 与 cells[1] 指向同一 tc）
    assert t.rows[0].cells[0]._tc is t.rows[0].cells[1]._tc


def test_apply_edits_clone_separates_tables(tmp_path):
    src = tmp_path / "d.docx"
    doc = Document()
    t = doc.add_table(1, 1)
    t.cell(0, 0).text = "简历表"
    doc.save(str(src))
    out = tmp_path / "o.docx"
    res = apply_edits(str(src), str(out), clones=[{"table_index": 0, "count": 2}])
    assert res["cloned"] == 2
    d = Document(str(out))
    assert len(d.tables) == 3
    # 任意相邻两表之间必有空段（防 Word 合并）
    body = d.element.body
    children = list(body)
    for i in range(1, len(children) - 1):
        if children[i].tag == qn("w:tbl"):
            prev = children[i - 1]
            assert prev.tag != qn("w:tbl"), "相邻表格之间应有空段分隔"
