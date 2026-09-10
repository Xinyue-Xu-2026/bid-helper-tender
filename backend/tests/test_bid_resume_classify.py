"""B1：简历样表优先识别（修复被 person_roster 抢走导致未克隆）。"""
from docx import Document

from app.core.bid_table_classifier import classify_tables


def _save(doc, tmp_path):
    p = tmp_path / "d.docx"
    doc.save(str(p))
    return str(p)


def test_resume_sample_priority_over_person_roster(tmp_path):
    doc = Document()
    t = doc.add_table(6, 2)
    labels = ["姓名", "年龄", "执业资格证书（或上岗证书）名称",
              "拟在本项目任职", "主要工作经历", "学历"]
    for i, lab in enumerate(labels):
        t.cell(i, 0).text = lab
    sugg = classify_tables(_save(doc, tmp_path))
    assert len(sugg) == 1
    s = sugg[0]
    assert s["role"] == "resume_each"
    assert s["mode"] == "per_person"
    assert "name" in s["columns"].values()


def test_person_roster_not_misclassified_as_resume(tmp_path):
    doc = Document()
    t = doc.add_table(2, 5)
    for j, h in enumerate(["序号", "姓名", "职称", "专业工作年限", "执业资格"]):
        t.cell(0, j).text = h
    for j, v in enumerate(["1", "张三", "高工", "10年", "一级注册造价师"]):
        t.cell(1, j).text = v
    sugg = classify_tables(_save(doc, tmp_path))
    assert sugg[0]["role"] == "person_roster"
