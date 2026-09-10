"""简历表「一人一表」整表克隆测试（V1.2 7.4）：选 N 人 → 样表 + N-1 份
XML 级克隆（结构原样：行列数/合并/行高保留），逐人填充；克隆表登记进
verify（table_insertions）不误报；无 mode=per_person 保持旧行为。
另含 7.3 donor 克隆行 vMerge 剔除回归。"""
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm

from app.core.bid_draft_exporter import fill_draft


def _sample_resume_table(doc):
    t = doc.add_table(2, 2)
    t.rows[0].cells[0].text = "姓名"; t.rows[0].cells[1].text = ""
    t.rows[1].cells[0].text = "主要工作经历"; t.rows[1].cells[1].text = ""
    return t


def _persons():
    return [
        {"name": "张三", "is_lead": True, "role": "项目经理",
         "sem": {"name": "张三"}, "perfs_text": "甲项目（2024）", "fields": {}},
        {"name": "李四", "is_lead": False, "role": "造价员",
         "sem": {"name": "李四"}, "perfs_text": "乙项目（2023）", "fields": {}},
    ]


def _bindings(mode="per_person"):
    return {"tables": [{"table_index": 0, "role": "resume_each",
                        "columns": {"姓名": "name", "主要工作经历": "lead_perfs"},
                        "mode": mode, "confirmed": True}],
            "swap_toc": False}


def test_per_person_clone_count_and_structure(tmp_path):
    src = tmp_path / "draft.docx"; dst = tmp_path / "out.docx"
    doc = Document(); _sample_resume_table(doc); doc.save(src)
    data = {"persons": _persons(), "contracts": [],
            "lead_perfs_text": "甲项目（2024）"}
    report = fill_draft(str(src), str(dst), _bindings(), data,
                        project_no="", project_name="", doc_date="")
    out = Document(str(dst))
    assert len(out.tables) == 2                       # 样表 + 1 克隆
    assert out.tables[0].rows[0].cells[1].text == "张三"
    assert out.tables[1].rows[0].cells[1].text == "李四"
    # 逐人 perfs_text：lead_perfs 语义键写本人业绩文本
    assert out.tables[0].rows[1].cells[1].text == "甲项目（2024）"
    assert out.tables[1].rows[1].cells[1].text == "乙项目（2023）"
    assert report["verify"]["ok"] is True             # 克隆已登记，不误报
    assert report["filled"][0]["cloned"] == 2


def test_clone_preserves_row_col_merge_and_height(tmp_path):
    """deepcopy 克隆保留结构：行列数、横向合并（同 _tc）、行高。"""
    src = tmp_path / "draft.docx"; dst = tmp_path / "out.docx"
    doc = Document()
    t = doc.add_table(3, 2)
    t.rows[0].cells[0].text = "姓名"; t.rows[0].cells[1].text = ""
    t.rows[1].cells[0].text = "主要工作经历"; t.rows[1].cells[1].text = ""
    t.rows[2].cells[0].text = "职称"; t.rows[2].cells[1].text = ""
    t.rows[0].cells[0].merge(t.rows[0].cells[1])   # 横向合并首行
    t.rows[1].height = Cm(1.5)
    doc.save(src)
    persons = _persons() + [
        {"name": "王五", "is_lead": False, "role": "组员",
         "sem": {"name": "王五"}, "perfs_text": "", "fields": {}}]
    data = {"persons": persons, "contracts": [], "lead_perfs_text": ""}
    report = fill_draft(str(src), str(dst), _bindings(), data)
    # 行高以底稿自身的落盘值为准（Cm→twips 有取整），克隆须与其完全一致
    draft_height = Document(str(src)).tables[0].rows[1].height
    assert draft_height is not None
    out = Document(str(dst))
    assert len(out.tables) == 3
    for tb in out.tables:
        assert len(tb.rows) == 3 and len(tb.columns) == 2
        # 首行横向合并保留：两格同 _tc
        assert tb.rows[0].cells[0]._tc is tb.rows[0].cells[1]._tc
        # 行高保留（与底稿同值）
        assert tb.rows[1].height == draft_height
    assert report["verify"]["ok"] is True


def test_resume_each_without_mode_keeps_single_lead(tmp_path):
    """无 mode=per_person：resume_each 退化为旧 lead_resume 行为
    （仅负责人单份填充，不克隆）。"""
    src = tmp_path / "draft.docx"; dst = tmp_path / "out.docx"
    doc = Document(); _sample_resume_table(doc); doc.save(src)
    data = {"persons": _persons(), "contracts": [],
            "lead_perfs_text": "甲项目（2024）"}
    report = fill_draft(str(src), str(dst), _bindings(mode=""), data)
    out = Document(str(dst))
    assert len(out.tables) == 1
    assert out.tables[0].rows[0].cells[1].text == "张三"
    assert "cloned" not in report["filled"][0]
    assert report["verify"]["ok"] is True


def test_classifier_resume_each_vs_lead_resume(tmp_path):
    """分类器：键值样表含「主要工作经历」等一人一表标记 → resume_each；
    不含标记的旧样表仍 → lead_resume（向后兼容）。"""
    from app.core.bid_table_classifier import classify_tables
    doc2 = Document()
    t = doc2.add_table(4, 2)
    for r, label in enumerate(["姓名", "性别", "职称", "主要工作经历"]):
        t.rows[r].cells[0].text = label
    path2 = tmp_path / "d2.docx"; doc2.save(path2)
    items2 = classify_tables(str(path2))
    assert items2[0]["role"] == "resume_each"
    # 无标记键值表 → lead_resume（旧行为不回归）
    doc3 = Document()
    t3 = doc3.add_table(4, 2)
    for r, label in enumerate(["姓名", "性别", "职称", "学历"]):
        t3.rows[r].cells[0].text = label
    path3 = tmp_path / "d3.docx"; doc3.save(path3)
    assert classify_tables(str(path3))[0]["role"] == "lead_resume"


def test_fill_strips_vmerge_from_donor(tmp_path):
    """7.3 回归：donor 数据行含 vMerge 元素时，克隆写出的数据行不得携带
    vMerge（否则会延续/开启纵向合并区，破坏结构）。"""
    src = tmp_path / "draft.docx"; dst = tmp_path / "out.docx"
    doc = Document()
    t = doc.add_table(3, 2)
    t.cell(0, 0).text = "姓名"; t.cell(0, 1).text = "职称"
    t.cell(1, 0).text = "旧名"; t.cell(1, 1).text = "旧职称"
    t.cell(2, 0).text = "旧名2"; t.cell(2, 1).text = "旧职称2"
    tc_pr = t.rows[1].cells[0]._tc.get_or_add_tcPr()
    tc_pr.append(tc_pr.makeelement(qn("w:vMerge"), {}))
    doc.save(src)
    bindings = {"tables": [{"table_index": 0, "role": "person_roster",
                            "columns": {"name": 0, "title": 1},
                            "person_scope": "all", "confirmed": True}],
                "swap_toc": False}
    data = {"persons": [
        {"name": "张三", "is_lead": True,
         "sem": {"name": "张三", "title": "高级工程师"}, "fields": {}},
        {"name": "李四", "is_lead": False,
         "sem": {"name": "李四", "title": "工程师"}, "fields": {}}],
        "contracts": []}
    fill_draft(str(src), str(dst), bindings, data)
    out_tbl = Document(str(dst)).tables[0]
    assert not out_tbl._tbl.findall(f".//{qn('w:vMerge')}")
    rows = [[c.text for c in r.cells] for r in out_tbl.rows]
    assert rows[1] == ["张三", "高级工程师"]
    assert rows[2] == ["李四", "工程师"]
