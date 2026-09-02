"""Word 导出测试。"""
import pytest
from docx import Document

from app.core.word_exporter import export_word


SECTIONS = [
    {"title": "第一章 资格审查", "level": 1, "content": "正文一"},
    {"title": "1.1 营业执照", "level": 2, "content": "正文二"},
    {"title": "1.1.1 副本", "level": 3, "content": ""},
    {"title": "超深章", "level": 5, "content": ""},
]


def test_export_word_toc_and_headings(tmp_path):
    dest = tmp_path / "out.docx"
    export_word({"name": "测试项目"}, None, SECTIONS, str(dest))
    assert dest.exists()
    doc = Document(str(dest))
    xml = doc.element.body.xml
    assert "instrText" in xml and "TOC" in xml          # TOC 域存在
    heads = [p for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert heads[0].text == "第一章 资格审查"
    assert heads[0].style.name == "Heading 1"
    assert heads[1].style.name == "Heading 2"
    assert heads[2].style.name == "Heading 3"
    assert heads[3].style.name == "Heading 4"           # 超深章截断到 4 级


def test_export_word_empty_sections(tmp_path):
    with pytest.raises(ValueError, match="章节树为空"):
        export_word({"name": "x"}, None, [], str(tmp_path / "e.docx"))


def test_export_word_with_template(tmp_path):
    tpl = tmp_path / "tpl.docx"
    Document().save(str(tpl))
    dest = tmp_path / "out.docx"
    export_word({"name": "x"}, {"file_path": str(tpl)},
                [{"title": "A", "level": 1, "content": "B"}], str(dest))
    doc = Document(str(dest))
    heads = [p for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert heads[0].text == "A"


def _build_cover_template(path):
    """构造含封面（普通段落+封面表格）+ 旧正文标题/表格的模板。"""
    doc = Document()
    doc.add_paragraph("某某项目投标文件封面")
    doc.add_paragraph("投标单位：测试公司")
    cover_tbl = doc.add_table(rows=1, cols=2)
    cover_tbl.cell(0, 0).text = "封面表A1"
    cover_tbl.cell(0, 1).text = "封面表B1"
    doc.add_heading("第一章 旧模板章节", level=1)
    doc.add_paragraph("旧正文应被清除")
    old_tbl = doc.add_table(rows=1, cols=1)
    old_tbl.cell(0, 0).text = "旧表应被清除"
    doc.save(str(path))


def test_export_word_keeps_cover_and_clears_body(tmp_path):
    tpl = tmp_path / "tpl.docx"
    _build_cover_template(tpl)
    dest = tmp_path / "out.docx"
    export_word({"name": "x"}, {"file_path": str(tpl)},
                [{"title": "第一章 新章节", "level": 1, "content": "新正文"}], str(dest))
    doc = Document(str(dest))
    paras = doc.paragraphs
    # 封面段落保留在文档开头
    assert paras[0].text == "某某项目投标文件封面"
    assert paras[1].text == "投标单位：测试公司"
    # 封面表格保留、旧正文表格被清除
    table_texts = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    assert "封面表A1" in table_texts
    assert "旧表应被清除" not in table_texts
    # 旧标题/旧正文被清除，新章节用 Heading 样式
    all_text = "\n".join(p.text for p in paras)
    assert "旧模板章节" not in all_text and "旧正文应被清除" not in all_text
    heads = [p for p in paras if p.style.name.startswith("Heading")]
    assert heads[0].text == "第一章 新章节"
    assert heads[0].style.name == "Heading 1"
    # TOC 域在封面之后、章节之前
    xml = doc.element.body.xml
    assert "instrText" in xml and "TOC" in xml
    cover_idx = xml.index("投标单位")
    toc_idx = xml.index("instrText")
    head_idx = xml.index("第一章 新章节")
    assert cover_idx < toc_idx < head_idx


def test_export_word_markdown_table_and_bullets(tmp_path):
    content = (
        "说明段落一\n"
        "\n"
        "| 名称 | 数量 |\n"
        "| --- | --- |\n"
        "| 服务器 | 3 台 |\n"
        "| 交换机 | 5 台 |\n"
        "\n"
        "- 要点一\n"
        "- 要点二\n"
        "收尾段落\n"
    )
    dest = tmp_path / "out.docx"
    export_word({"name": "x"}, None,
                [{"title": "第一章", "level": 1, "content": content}], str(dest))
    doc = Document(str(dest))
    # Markdown 表格 → 真表格
    assert len(doc.tables) == 1
    tbl = doc.tables[0]
    assert len(tbl.rows) == 3 and len(tbl.columns) == 2
    assert tbl.cell(0, 0).text == "名称"
    assert tbl.cell(1, 1).text == "3 台"
    assert tbl.cell(2, 0).text == "交换机"
    # 列表行 → List Bullet
    bullets = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert [p.text for p in bullets] == ["要点一", "要点二"]
    # 普通段落按行拆分
    texts = [p.text for p in doc.paragraphs]
    assert "说明段落一" in texts and "收尾段落" in texts
