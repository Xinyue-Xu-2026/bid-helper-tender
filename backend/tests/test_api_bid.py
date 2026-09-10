"""商务标 API 测试（设计文档 §8 测试计划 1-8）：
勾选 PUT→GET 往返 / 覆盖语义 / 校验错误 / xlsx、docx 导出 /
空勾选导出 422 / expiring-detail 明细 / 证书有效期警告（expired、soon）。"""
from datetime import date, timedelta
from io import BytesIO
from urllib.parse import unquote
import json

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.bid_exporter import PERSON_HEADERS
from app.db import Database
from app.deps import get_db
from app.main import app


@pytest.fixture()
def client(db_path):
    def override():
        db = Database(db_path)
        db.init_schema()
        yield db
    app.dependency_overrides[get_db] = override
    yield TestClient(app)
    app.dependency_overrides.clear()


def _make_project(client, name="测试项目", bid_date=""):
    r = client.post("/api/projects", json={"name": name, "bid_date": bid_date})
    assert r.status_code == 200
    return r.json()["id"]


def _make_person(client, name, fields=None):
    r = client.post("/api/assets", json={
        "type": "person", "name": name, "fields": fields or {}})
    assert r.status_code == 200
    return r.json()["id"]


def _make_contract(client, name, fields=None):
    r = client.post("/api/assets", json={
        "type": "contract", "name": name, "fields": fields or {}})
    assert r.status_code == 200
    return r.json()["id"]


def _days(n):
    return (date.today() + timedelta(days=n)).isoformat()


def _bid_url(pid, suffix=""):
    return f"/api/projects/{pid}/bid-assets{suffix}"


# ---------- 1. PUT→GET 往返 ----------

def test_put_get_roundtrip(client):
    pid = _make_project(client)
    far = _days(3650)
    person_id = _make_person(client, "张三", fields={
        "职称": "高级工程师",
        "证书": [{"类型": "一级建造师", "编号": "A001", "有效期至": far}],
    })
    contract_id = _make_contract(client, "某市政工程", fields={
        "类型": "编标", "合同金额": "500万", "年份": "2025"})

    r = client.put(_bid_url(pid), json={
        "persons": [{"asset_id": person_id, "role": "项目经理"}],
        "contracts": [contract_id],
    })
    assert r.status_code == 200
    assert r.json() == {"ok": True, "persons": 1, "contracts": 1}

    data = client.get(_bid_url(pid)).json()
    assert len(data["persons"]) == 1
    p = data["persons"][0]
    assert p["asset_id"] == person_id
    assert p["role"] == "项目经理"
    assert p["name"] == "张三"
    assert p["fields"]["职称"] == "高级工程师"
    assert p["fields"]["证书"][0]["有效期至"] == far
    assert p["cert_warnings"] == []  # 无投标日且证书长期有效 → 无警告
    assert len(data["contracts"]) == 1
    c = data["contracts"][0]
    assert c["asset_id"] == contract_id
    assert c["name"] == "某市政工程"
    assert c["fields"]["合同金额"] == "500万"
    # 旧格式 PUT（无 is_lead/certs、裸 id 合同）→ GET 归一化默认值
    assert p["is_lead"] is False
    assert p["certs"] is None
    assert c["section"] == 1


def test_put_get_roundtrip_with_export_selections(client):
    """新格式 PUT：is_lead/certs/section 持久化并在 GET 还原。"""
    pid = _make_project(client)
    person_a = _make_person(client, "张三", fields={
        "证书": [{"类型": "一级造价师"}, {"类型": "监理工程师"}]})
    person_b = _make_person(client, "李四")
    contract_a = _make_contract(client, "业绩A")
    contract_b = _make_contract(client, "业绩B")

    r = client.put(_bid_url(pid), json={
        "persons": [
            {"asset_id": person_a, "role": "项目经理",
             "is_lead": True, "certs": [1]},
            {"asset_id": person_b, "role": "组员", "certs": None},
        ],
        "contracts": [
            {"asset_id": contract_a, "section": 2},
            {"asset_id": contract_b},  # section 缺省 → 1
        ],
    })
    assert r.status_code == 200
    assert r.json() == {"ok": True, "persons": 2, "contracts": 2}

    data = client.get(_bid_url(pid)).json()
    pa, pb = data["persons"]
    assert pa["asset_id"] == person_a
    assert pa["is_lead"] is True
    assert pa["certs"] == [1]
    assert pb["asset_id"] == person_b
    assert pb["is_lead"] is False
    assert pb["certs"] is None  # null=全部证书
    ca, cb = data["contracts"]
    assert (ca["asset_id"], ca["section"]) == (contract_a, 2)
    assert (cb["asset_id"], cb["section"]) == (contract_b, 1)


def test_get_normalizes_old_format_stored_rows(client, db_path):
    """模拟旧库数据：db 层以旧格式（无 is_lead/certs、裸 id 合同）写入，
    再走原始 SQL 插一行不含新列的记录，GET 均应按默认值归一化。"""
    pid = _make_project(client)
    person_id = _make_person(client, "张三")
    contract_id = _make_contract(client, "旧业绩")

    from app.db import Database
    db = Database(db_path)
    # 旧形状调用（历史 save_bid_assets 的落库方式）
    db.replace_project_assets(pid, [{"asset_id": person_id, "role": "项目经理"}],
                              [contract_id])
    # 更旧：连新列都不写的原始行（依赖列默认值）
    person2 = _make_person(client, "李四")
    with db._connect() as conn:
        conn.execute(
            "INSERT INTO project_assets (project_id, asset_id, role, sort_order) "
            "VALUES (?, ?, ?, ?)", (pid, person2, "组员", 99))

    data = client.get(_bid_url(pid)).json()
    persons = {p["asset_id"]: p for p in data["persons"]}
    assert persons[person_id]["is_lead"] is False
    assert persons[person_id]["certs"] is None
    assert persons[person2]["is_lead"] is False
    assert persons[person2]["certs"] is None
    assert data["contracts"][0]["asset_id"] == contract_id
    assert data["contracts"][0]["section"] == 1


# ---------- 2. PUT 覆盖语义 ----------

def test_put_overwrites_previous_selection(client):
    pid = _make_project(client)
    person_a = _make_person(client, "张三")
    person_b = _make_person(client, "李四")
    contract_id = _make_contract(client, "旧业绩")

    client.put(_bid_url(pid), json={
        "persons": [{"asset_id": person_a, "role": "项目经理"}],
        "contracts": [contract_id],
    })
    r = client.put(_bid_url(pid), json={
        "persons": [{"asset_id": person_b, "role": "技术负责人"}],
        "contracts": [],
    })
    assert r.status_code == 200

    data = client.get(_bid_url(pid)).json()
    assert [p["asset_id"] for p in data["persons"]] == [person_b]
    assert data["persons"][0]["role"] == "技术负责人"
    assert data["contracts"] == []


# ---------- 3. 校验错误 ----------

def test_project_not_found(client):
    assert client.get(_bid_url(99999)).status_code == 404
    r = client.put(_bid_url(99999), json={"persons": [], "contracts": []})
    assert r.status_code == 404


def test_missing_asset_id_422(client):
    pid = _make_project(client)
    r = client.put(_bid_url(pid), json={
        "persons": [{"asset_id": 99999, "role": "项目经理"}],
        "contracts": [],
    })
    assert r.status_code == 422
    assert "资产不存在" in r.json()["detail"]
    r = client.put(_bid_url(pid), json={"persons": [], "contracts": [88888]})
    assert r.status_code == 422


def test_type_mismatch_422(client):
    pid = _make_project(client)
    person_id = _make_person(client, "张三")
    contract_id = _make_contract(client, "某项目")
    # person 槽位放了 contract 资产
    r = client.put(_bid_url(pid), json={
        "persons": [{"asset_id": contract_id, "role": "项目经理"}],
        "contracts": [],
    })
    assert r.status_code == 422
    assert "类型不匹配" in r.json()["detail"]
    # contract 槽位放了 person 资产
    r = client.put(_bid_url(pid), json={
        "persons": [], "contracts": [person_id],
    })
    assert r.status_code == 422
    assert "类型不匹配" in r.json()["detail"]


# ---------- 4/5. 导出 xlsx / docx ----------

def _select_one_each(client, pid, person_id, contract_id):
    r = client.put(_bid_url(pid), json={
        "persons": [{"asset_id": person_id, "role": "项目经理"}],
        "contracts": [contract_id],
    })
    assert r.status_code == 200


def test_export_xlsx(client):
    pid = _make_project(client, name="导出项目")
    far = _days(3650)
    person_id = _make_person(client, "张三", fields={
        "职称": "高级工程师",
        "证书": [{"类型": "一级建造师", "专业": "建筑工程", "有效期至": far},
                 {"类型": "监理工程师", "有效期至": far}],
    })
    contract_id = _make_contract(client, "某市政工程", fields={
        "类型": "编标", "合同金额": "500万", "年份": "2025"})
    _select_one_each(client, pid, person_id, contract_id)

    r = client.get(_bid_url(pid, "/export"), params={"format": "xlsx"})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    wb = load_workbook(BytesIO(r.content))
    ws = wb["人员配备"]
    assert [c.value for c in ws[1]] == PERSON_HEADERS
    row = [c.value for c in ws[2]]
    assert row[0] == 1 and row[1] == "张三" and row[2] == "项目经理"
    assert row[3] == "高级工程师"
    assert f"一级建造师·建筑工程（有效期至{far}）" in row[4]
    assert f"监理工程师（有效期至{far}）" in row[4]
    assert row[5] == "正常"

    ws2 = wb["企业业绩"]
    headers = [c.value for c in ws2[1]]
    assert headers[:2] == ["序号", "项目名称"]
    assert "合同金额" in headers
    row2 = [c.value for c in ws2[2]]
    assert row2[0] == 1 and row2[1] == "某市政工程"
    assert row2[headers.index("合同金额")] == "500万"


def test_export_docx(client):
    pid = _make_project(client, name="导出项目")
    person_id = _make_person(client, "张三", fields={"职称": "工程师"})
    contract_id = _make_contract(client, "某市政工程", fields={
        "类型": "编标", "合同金额": "500万", "年份": "2025"})
    _select_one_each(client, pid, person_id, contract_id)

    r = client.get(_bid_url(pid, "/export"), params={"format": "docx"})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    doc = Document(BytesIO(r.content))
    assert any(p.text == "导出项目 商务标" for p in doc.paragraphs)
    assert len(doc.tables) == 2

    person_table, contract_table = doc.tables
    assert [c.text for c in person_table.rows[0].cells] == PERSON_HEADERS
    person_row = [c.text for c in person_table.rows[1].cells]
    assert person_row[1] == "张三" and person_row[2] == "项目经理"
    assert person_row[3] == "工程师"

    contract_headers = [c.text for c in contract_table.rows[0].cells]
    assert contract_headers[:2] == ["序号", "项目名称"]
    contract_row = [c.text for c in contract_table.rows[1].cells]
    assert contract_row[1] == "某市政工程"
    assert contract_row[contract_headers.index("合同金额")] == "500万"


# ---------- 6. 无勾选导出 → 422 ----------

def test_export_without_selection_422(client):
    pid = _make_project(client)
    for fmt in ("xlsx", "docx"):
        r = client.get(_bid_url(pid, "/export"), params={"format": fmt})
        assert r.status_code == 422
        assert "勾选" in r.json()["detail"]


# ---------- 7. expiring-detail 证书级到期明细 ----------

def test_expiring_detail_person_cert_and_credit(client):
    person_expiry = _days(15)
    credit_expiry = _days(20)
    _make_person(client, "张三", fields={
        "证书": [{"类型": "一级建造师", "编号": "A001",
                  "有效期至": person_expiry}],
    })
    client.post("/api/assets", json={
        "type": "credit", "name": "安全生产许可证",
        "expiry_date": credit_expiry})
    # 超窗与无日期的干扰项不应出现
    _make_person(client, "李四", fields={
        "证书": [{"类型": "造价师", "有效期至": _days(60)}]})
    client.post("/api/assets", json={"type": "credit", "name": "无期限证书"})

    rows = client.get("/api/assets/expiring-detail", params={"days": 30}).json()
    assert len(rows) == 2
    # 按 days_left 升序：人员证书（15 天）在前
    person_row, credit_row = rows
    assert person_row["type"] == "person"
    assert person_row["asset_name"] == "张三"
    assert person_row["cert_name"] == "一级建造师"
    assert person_row["expiry_date"] == person_expiry
    assert person_row["days_left"] == 15
    assert credit_row["type"] == "credit"
    assert credit_row["asset_name"] == "安全生产许可证"
    assert credit_row["days_left"] == 20


# ---------- 8. 证书警告 expired / soon ----------

def test_cert_warning_expired_before_bid_date(client):
    pid = _make_project(client, bid_date=_days(30))
    expiry = _days(10)  # 有效期早于投标日 → expired（最高优先）
    person_id = _make_person(client, "张三", fields={
        "证书": [{"类型": "一级建造师", "编号": "A001",
                  "有效期至": expiry}],
    })
    client.put(_bid_url(pid), json={
        "persons": [{"asset_id": person_id, "role": ""}], "contracts": []})

    data = client.get(_bid_url(pid)).json()
    warnings = data["persons"][0]["cert_warnings"]
    assert len(warnings) == 1
    w = warnings[0]
    assert w["level"] == "expired"
    assert w["cert_name"] == "一级建造师"
    assert w["expiry"] == expiry
    assert "早于投标日" in w["message"]


def test_cert_warning_soon_with_past_bid_date(client):
    # 投标日已过去（不触发 expired），证书 15 天内到期 → soon
    pid = _make_project(client, bid_date=_days(-365))
    expiry = _days(15)
    person_id = _make_person(client, "张三", fields={
        "证书": [{"类型": "一级建造师", "编号": "A001",
                  "有效期至": expiry}],
    })
    client.put(_bid_url(pid), json={
        "persons": [{"asset_id": person_id, "role": ""}], "contracts": []})

    data = client.get(_bid_url(pid)).json()
    warnings = data["persons"][0]["cert_warnings"]
    assert len(warnings) == 1
    w = warnings[0]
    assert w["level"] == "soon"
    assert w["cert_name"] == "一级建造师"
    assert w["days_left"] == 15
    assert "15" in w["message"]


# ---------- 9. 标段字段（V1.2 Task 2） ----------

def test_projects_has_section_columns(client):
    from app.db import Database
    db = Database(); db.init_schema(); db.init_schema()   # 幂等
    with db._connect() as conn:
        cols = {r[1] for r in conn.execute("PRAGMA table_info(projects)")}
    assert {"section_name", "section_no"} <= cols


def _make_section_draft(client, db_path, pid, tmp_path):
    """直接造底稿：含 标段名称：____（标签空白）与 （标段名称）（括号占位）。"""
    doc = Document()
    doc.add_paragraph("标段名称：____")
    doc.add_paragraph("（标段名称）")
    path = tmp_path / "draft.docx"
    doc.save(str(path))
    db = Database(db_path)
    db.create_bid_template(pid, "底稿", str(path))
    return path


def test_export_section_param_fills_label_and_bracket(client, db_path, tmp_path):
    """导出 body 带 section_name/section_no → 底稿标签空白与括号占位被填充；
    section 字段非空时持久化到 projects。"""
    pid = _make_project(client)
    _make_section_draft(client, db_path, pid, tmp_path)
    person_id = _make_person(client, "张三")
    r = client.post(_bid_url(pid, "/export-template"), json={
        "persons": [{"asset_id": person_id, "is_lead": True}],
        "contracts": [],
        "section_name": "一标段", "section_no": "A1"})
    assert r.status_code == 200
    out = Document(BytesIO(r.content))
    texts = [p.text for p in out.paragraphs]
    assert "标段名称：一标段" in texts
    assert "一标段" in texts and "（标段名称）" not in texts
    # 校验端共用同一套规则 → 不误报
    report = json.loads(unquote(r.headers["x-fill-report"]))
    assert report["verify"]["ok"] is True
    # 非空标段持久化到 projects
    row = Database(db_path).get_project(pid)
    assert row["section_name"] == "一标段"
    assert row["section_no"] == "A1"


def test_export_section_empty_skips_fill(client, db_path, tmp_path):
    """section 为空 → 不填充，空白占位保持原样，projects 不写标段。"""
    pid = _make_project(client)
    _make_section_draft(client, db_path, pid, tmp_path)
    person_id = _make_person(client, "张三")
    r = client.post(_bid_url(pid, "/export-template"), json={
        "persons": [{"asset_id": person_id, "is_lead": True}],
        "contracts": []})
    assert r.status_code == 200
    out = Document(BytesIO(r.content))
    texts = [p.text for p in out.paragraphs]
    assert "标段名称：____" in texts
    assert "（标段名称）" in texts
    row = Database(db_path).get_project(pid)
    assert row["section_name"] == "" and row["section_no"] == ""
