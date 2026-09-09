# -*- coding: utf-8 -*-
"""授权页正文填充（P5）：法定代表人身份证明 / 授权委托书的
姓名/身份证号/日期/身份证附图 填充 + verify bound 段落豁免。"""
import tempfile
from pathlib import Path

from docx import Document

from app.core.bid_draft_exporter import fill_draft
from app.core.bid_template_exporter import _para_text

# 1x1 PNG（python-docx 可直接插图；base64 合法最小图）
import base64 as _b64
_PNG = _b64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk"
    "YPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


def _auth_draft(path):
    """构造含授权两页形态的程序化底稿（段落结构与真实招标一致）。"""
    doc = Document()
    lines = [
        "二、投标人代表身份证明",
        "法定代表人（单位负责人）身份证明",
        "投标人名称：",
        "姓名：性别：年龄：职务：",
        "系（投标人名称）的法定代表人（单位负责人）。",
        "特此证明。",
        "投标人：（盖单位公章）",
        "年月日",
        "附：法定代表人（单位负责人）身份证（正反面）电子扫描件。",
        "授权委托书",
        "本人（姓名）系（投标人名称）的法定代表人（单位负责人），现委托（姓名）为我方代理人。"
        "代理人根据授权，以我方名义签署、澄清确认、递交、撤回、修改（工程项目名称）（标段名称）"
        "投标文件、签订合同和处理有关事宜，其法律后果由我方承担。",
        "委托期限：自签署之日至投标有效期期满。",
        "代理人无转委托权。",
        "附：委托双方的身份证（正反面）电子扫描件。",
        "投标人：    （盖单位公章）",
        "法定代表人（单位负责人）：（签名）",
        "身份证号：",
        "委托代理人：（签名）",
        "身份证号：",
        "年月日",
        "其他承诺内容保持不变。",
    ]
    for ln in lines:
        doc.add_paragraph(ln)
    doc.save(path)
    return path


def _legal(name, id_no, img_dir, tag):
    """构造 auth person 字典。"""
    front = img_dir / f"{tag}_front.png"
    back = img_dir / f"{tag}_back.png"
    front.write_bytes(_PNG)
    back.write_bytes(_PNG)
    return {"name": name, "身份证号": id_no,
            "身份证正面扫描件": str(front), "身份证反面扫描件": str(back)}


def _texts(path):
    return [_para_text(p) for p in Document(path).paragraphs if _para_text(p).strip()]


def test_authority_fill_full(tmp_path):
    draft = _auth_draft(str(tmp_path / "draft.docx"))
    dest = str(tmp_path / "out.docx")
    img_dir = tmp_path / "imgs"
    img_dir.mkdir()
    auth = {
        "legal_rep": _legal("张三", "320000199001011234", img_dir, "z"),
        "agent": _legal("李四", "320000199002022345", img_dir, "l"),
        "doc_date": "2026-09-10",
    }
    report = fill_draft(draft, dest, {"tables": [], "swap_toc": False}, {},
                        bidder_name="宏信天德工程顾问有限公司", auth=auth)
    texts = _texts(dest)

    # 括号占位：法人名进 本人（姓名），代理人名进 现委托（姓名）
    body = "".join(texts)
    assert "本人张三系" in body
    assert "现委托李四" in body
    # （投标人名称）→ 投标人名称
    assert "（投标人名称）" not in body
    assert "系宏信天德工程顾问有限公司的法定代表人" in body
    # 身份证号两处
    assert "身份证号：320000199001011234" in body
    assert "身份证号：320000199002022345" in body
    # 年月日 → 日期（两页）
    assert texts.count("2026年9月10日") == 2
    # 附图追加（身份证明页两张 + 委托双方四张 = 6 处图片 run）
    assert report["authority_images"] == 6
    # 图片追加段落含 drawing
    import re as _re
    import zipfile
    with zipfile.ZipFile(dest) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert xml.count("<pic:pic") >= 6
    # verify 通过（bound 段落豁免；图片追加不影响文本比对）
    assert report["verify"]["ok"] is True, report["verify"]["issues"]
    # 未涉及段落未动
    assert "其他承诺内容保持不变。" in body


def test_authority_fill_legal_only(tmp_path):
    draft = _auth_draft(str(tmp_path / "draft.docx"))
    dest = str(tmp_path / "out.docx")
    img_dir = tmp_path / "imgs"
    img_dir.mkdir()
    auth = {"legal_rep": _legal("张三", "320000199001011234", img_dir, "z"),
            "agent": {}, "doc_date": "2026-09-10"}
    report = fill_draft(draft, dest, {"tables": [], "swap_toc": False}, {},
                        bidder_name="宏信天德工程顾问有限公司", auth=auth)
    texts = _texts(dest)
    body = "".join(texts)
    assert "本人张三系" in body
    # 未选代理人 → 现委托（姓名）不替换
    assert "现委托（姓名）" in body
    assert "身份证号：320000199001011234" in body
    # 无代理人身份证号行保留原样（不写入空值/不新增）
    assert "身份证号：320000199002022345" not in body
    assert report["authority_images"] == 4  # 身份证明 2 + 委托书仅法代 2
    assert report["verify"]["ok"] is True


def test_authority_fill_empty_auth_no_change(tmp_path):
    draft = _auth_draft(str(tmp_path / "draft.docx"))
    dest = str(tmp_path / "out.docx")
    report = fill_draft(draft, dest, {"tables": [], "swap_toc": False}, {},
                        bidder_name="", auth={})
    body = "".join(_texts(dest))
    assert "（姓名）" in body and "身份证号：" in body
    assert report["authority_images"] == 0
    assert report["verify"]["ok"] is True
