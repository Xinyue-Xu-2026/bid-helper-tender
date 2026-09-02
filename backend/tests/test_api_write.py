"""编写工作台路由测试（大纲/单节生成 mock，零网络）。"""
import pytest
from fastapi.testclient import TestClient

from app.main import app


@pytest.fixture()
def client():
    with TestClient(app) as c:
        yield c


def _project(client):
    return client.post("/api/projects", json={"name": "编写测试项目"}).json()["id"]


def test_sections_crud(client):
    pid = _project(client)
    root = client.post(f"/api/projects/{pid}/sections",
                       json={"title": "第一章", "level": 1}).json()["id"]
    child = client.post(f"/api/projects/{pid}/sections",
                        json={"parent_id": root, "title": "1.1 子节", "level": 2}).json()["id"]
    tree = client.get(f"/api/projects/{pid}/sections").json()
    assert tree[0]["title"] == "第一章"
    assert tree[0]["children"][0]["title"] == "1.1 子节"
    client.put(f"/api/sections/{child}", json={"content": "正文", "gen_status": "已生成"})
    assert client.delete(f"/api/sections/{root}").json() == {"ok": True}  # 递归删子节
    assert client.get(f"/api/projects/{pid}/sections").json() == []


def test_outline_mock(client, monkeypatch):
    from app.db import Database
    pid = _project(client)
    db = Database(); db.init_schema()
    db.create_requirement(pid, "格式要求", "投标文件应包含资质材料")
    monkeypatch.setattr("app.services.write_service.get_api_key", lambda: "sk-test")
    monkeypatch.setattr("app.services.write_service.get_model", lambda: "kimi-k3")
    monkeypatch.setattr(
        "app.services.write_service.generate_outline",
        lambda reqs, key, model: [
            {"title": "第一章 资格审查", "level": 1, "children": [
                {"title": "1.1 营业执照", "level": 2, "children": []}]}])
    tree = client.post(f"/api/projects/{pid}/outline").json()
    assert tree[0]["title"] == "第一章 资格审查"
    assert tree[0]["children"][0]["title"] == "1.1 营业执照"


def test_generate_section_sse(client, monkeypatch):
    pid = _project(client)
    sid = client.post(f"/api/projects/{pid}/sections",
                      json={"title": "1.1 营业执照", "level": 2}).json()["id"]
    monkeypatch.setattr("app.services.write_service.get_api_key", lambda: "sk-test")
    monkeypatch.setattr("app.services.write_service.get_model", lambda: "kimi-k3")
    monkeypatch.setattr("app.services.write_service.stream_section",
                        lambda prompt, key, model: iter(["我", "方", "承诺"]))
    r = client.get(f"/api/projects/{pid}/sections/{sid}/generate")
    assert r.status_code == 200
    assert "event: chunk" in r.text
    assert "event: done" in r.text
    assert "data: 我方承诺" in r.text
    stored = client.get(f"/api/projects/{pid}/sections").json()[0]
    assert stored["content"] == "我方承诺"
    assert stored["gen_status"] == "已生成"


def test_generate_section_error_strips_newline(client, monkeypatch):
    pid = _project(client)
    sid = client.post(f"/api/projects/{pid}/sections",
                      json={"title": "1.1", "level": 2}).json()["id"]
    monkeypatch.setattr("app.services.write_service.get_api_key", lambda: "sk-test")
    monkeypatch.setattr("app.services.write_service.get_model", lambda: "kimi-k3")

    def boom(prompt, key, model):
        raise Exception("第一行\n第二行")
        yield  # pragma: no cover
    monkeypatch.setattr("app.services.write_service.stream_section", boom)
    r = client.get(f"/api/projects/{pid}/sections/{sid}/generate")
    assert "event: error" in r.text
    assert "第一行 第二行" in r.text
    assert "\n第二行" not in r.text


def test_generate_section_missing_404(client):
    pid = _project(client)
    assert client.get(f"/api/projects/{pid}/sections/999/generate").status_code == 404


def test_generate_section_multiline_sse(client, monkeypatch):
    pid = _project(client)
    sid = client.post(f"/api/projects/{pid}/sections",
                      json={"title": "1.1", "level": 2}).json()["id"]
    monkeypatch.setattr("app.services.write_service.get_api_key", lambda: "sk-test")
    monkeypatch.setattr("app.services.write_service.get_model", lambda: "kimi-k3")
    monkeypatch.setattr("app.services.write_service.stream_section",
                        lambda prompt, key, model: iter(["第一行\n第二行", "第三行"]))
    r = client.get(f"/api/projects/{pid}/sections/{sid}/generate")
    assert r.status_code == 200
    assert "data: 第一行\ndata: 第二行" in r.text
    stored = client.get(f"/api/projects/{pid}/sections").json()[0]
    assert stored["content"] == "第一行\n第二行第三行"
    assert stored["gen_status"] == "已生成"


def test_generate_section_cross_project_404(client):
    pid_a = _project(client)
    pid_b = client.post("/api/projects", json={"name": "另一个项目"}).json()["id"]
    sid = client.post(f"/api/projects/{pid_b}/sections",
                      json={"title": "1.1", "level": 2}).json()["id"]
    assert client.get(f"/api/projects/{pid_a}/sections/{sid}/generate").status_code == 404


# ---------- 模板章节树 from-template ----------

def _make_template_docx(path):
    from docx import Document
    doc = Document()
    doc.add_heading("第一章 项目概述", level=1)
    doc.add_heading("1.1 项目背景", level=2)
    doc.add_paragraph("背" * 1500)
    doc.add_heading("1.2 建设目标", level=2)
    doc.add_paragraph("标" * 500)
    doc.add_heading("第二章 技术方案", level=1)
    doc.add_heading("2.1 总体设计", level=2)
    doc.add_paragraph("设" * 800)
    doc.save(str(path))


def _register_template(tmp_path, name="模板A", profile="", with_headings=True):
    """造一个 docx 模板并登记到 templates 表，返回模板 id。"""
    from app.db import Database
    path = tmp_path / f"{name}.docx"
    if with_headings:
        _make_template_docx(path)
    else:
        from docx import Document
        Document().save(str(path))
    db = Database(); db.init_schema()
    tid = db.create_template(name, str(path))
    if profile:
        db.update_template(tid, style_profile=profile)
    return tid


def test_sections_from_template(client, tmp_path):
    pid = _project(client)
    tid = _register_template(tmp_path)
    r = client.post(f"/api/projects/{pid}/sections/from-template", json={"template_id": tid})
    assert r.status_code == 200
    sections = r.json()["sections"]
    assert [s["title"] for s in sections] == ["第一章 项目概述", "第二章 技术方案"]
    first = sections[0]
    assert first["level"] == 1
    assert [c["title"] for c in first["children"]] == ["1.1 项目背景", "1.2 建设目标"]
    assert first["children"][0]["level"] == 2
    assert first["children"][0]["parent_id"] == first["id"]
    # 再次调用覆盖旧章节（replace_sections）
    r2 = client.post(f"/api/projects/{pid}/sections/from-template", json={"template_id": tid})
    assert [s["title"] for s in r2.json()["sections"]] == ["第一章 项目概述", "第二章 技术方案"]


def test_sections_from_template_no_headings_422(client, tmp_path):
    pid = _project(client)
    tid = _register_template(tmp_path, name="空白模板", with_headings=False)
    r = client.post(f"/api/projects/{pid}/sections/from-template", json={"template_id": tid})
    assert r.status_code == 422
    assert r.json()["detail"] == "模板中未识别到章节标题，请检查模板样式"


def test_sections_from_template_missing_404(client, tmp_path):
    pid = _project(client)
    r = client.post(f"/api/projects/{pid}/sections/from-template", json={"template_id": 999})
    assert r.status_code == 404
    tid = _register_template(tmp_path)
    assert client.post(f"/api/projects/999/sections/from-template",
                       json={"template_id": tid}).status_code == 404


# ---------- 篇幅驱动 + templateId 画像选择 ----------

def test_generate_section_template_id_and_target(client, monkeypatch, tmp_path):
    pid = _project(client)
    # 章节树来自模板A（标题与模板一致，可命中字数统计）
    tid_a = _register_template(tmp_path, name="模板A", profile="画像A：模板A的风格")
    tid_b = _register_template(tmp_path, name="模板B", profile="画像B：模板B的风格")
    sections = client.post(f"/api/projects/{pid}/sections/from-template",
                           json={"template_id": tid_a}).json()["sections"]
    sid = sections[0]["children"][0]["id"]  # 1.1 项目背景

    monkeypatch.setattr("app.services.write_service.get_api_key", lambda: "sk-test")
    monkeypatch.setattr("app.services.write_service.get_model", lambda: "kimi-k3")
    captured = {}

    def fake_stream(prompt, key, model):
        captured["prompt"] = prompt
        return iter(["正文"])
    monkeypatch.setattr("app.services.write_service.stream_section", fake_stream)

    # 指定 templateId=tid_a：画像来自模板A（而非 templates[0] 即后登记的模板B）
    r = client.get(f"/api/projects/{pid}/sections/{sid}/generate?templateId={tid_a}")
    assert r.status_code == 200
    assert "画像A：模板A的风格" in captured["prompt"]
    assert "画像B：模板B的风格" not in captured["prompt"]
    # 篇幅要求：1.1 项目背景 实测 1500 字
    assert "本章不少于 1500 字" in captured["prompt"]

    # 指定 templateId=tid_b：画像来自模板B；模板B结构相同（同名模板docx），字数仍在
    r = client.get(f"/api/projects/{pid}/sections/{sid}/generate?templateId={tid_b}")
    assert "画像B：模板B的风格" in captured["prompt"]
    assert "画像A：模板A的风格" not in captured["prompt"]

    # 不传 templateId：回退 templates[0]（最新登记，即模板B），保持兼容
    r = client.get(f"/api/projects/{pid}/sections/{sid}/generate")
    assert "画像B：模板B的风格" in captured["prompt"]


# ---------- 逐小节生成（progress 事件） ----------

def _make_children_template(path):
    from docx import Document
    doc = Document()
    doc.add_heading("第一章 项目概述", level=1)
    doc.add_heading("1.1 项目背景", level=2)
    doc.add_paragraph("背" * 100)
    doc.add_heading("1.2 建设目标", level=2)
    doc.add_paragraph("标" * 200)
    doc.save(str(path))


def test_generate_section_with_children_progress(client, monkeypatch, tmp_path):
    from app.db import Database
    pid = client.post("/api/projects", json={
        "name": "GZ659商业综合体项目", "client": "扬州城建", "project_type": "造价咨询"}
    ).json()["id"]
    db = Database(); db.init_schema()
    # 自动匹配要求：按子节标题分词命中
    db.create_requirement(pid, "格式要求", "项目背景应结合地块区位详细说明")
    db.create_requirement(pid, "评分项", "建设目标响应程度得分")
    db.create_requirement(pid, "其他", "与本章无关的保证金要求")
    # 模板提供篇幅目标
    tpl = tmp_path / "tpl.docx"
    _make_children_template(tpl)
    tid = db.create_template("模板", str(tpl))
    # 章节树：父章 + 两子节
    root = client.post(f"/api/projects/{pid}/sections",
                       json={"title": "第一章 项目概述", "level": 1}).json()["id"]
    c1 = client.post(f"/api/projects/{pid}/sections",
                     json={"parent_id": root, "title": "1.1 项目背景", "level": 2}).json()["id"]
    c2 = client.post(f"/api/projects/{pid}/sections",
                     json={"parent_id": root, "title": "1.2 建设目标", "level": 2}).json()["id"]

    monkeypatch.setattr("app.services.write_service.get_api_key", lambda: "sk-test")
    monkeypatch.setattr("app.services.write_service.get_model", lambda: "kimi-k3")
    prompts = []

    def fake_stream(prompt, key, model):
        prompts.append(prompt)
        if "当前撰写章节：1.1 项目背景" in prompt:
            return iter(["背景正文"])
        return iter(["目标正文"])
    monkeypatch.setattr("app.services.write_service.stream_section", fake_stream)

    # req_ids 传了也应被忽略（自动匹配）
    r = client.get(f"/api/projects/{pid}/sections/{root}/generate?templateId={tid}&req_ids=999")
    assert r.status_code == 200
    # progress 事件序列
    assert 'event: progress\ndata: {"type": "progress", "current": 1, "total": 2, "title": "1.1 项目背景"}' in r.text
    assert 'event: progress\ndata: {"type": "progress", "current": 2, "total": 2, "title": "1.2 建设目标"}' in r.text
    assert r.text.index("current\": 1") < r.text.index("背景正文") < r.text.index("current\": 2")
    assert "event: done" in r.text
    # 拼接写入本章：子节标题行 + 正文
    tree = client.get(f"/api/projects/{pid}/sections").json()
    assert tree[0]["content"] == "1.1 项目背景\n背景正文\n\n1.2 建设目标\n目标正文"
    assert tree[0]["gen_status"] == "已生成"
    assert tree[0]["children"][0]["content"] == "背景正文"
    assert tree[0]["children"][1]["content"] == "目标正文"
    # prompt 增强：项目信息在开头、自动匹配要求、篇幅目标
    p1, p2 = prompts
    assert p1.startswith("项目信息：\n项目名称：GZ659商业综合体项目\n招标人/客户：扬州城建")
    assert "项目背景应结合地块区位详细说明" in p1
    assert "建设目标响应程度得分" not in p1  # 未命中 1.1 标题词
    assert "与本章无关的保证金要求" not in p1
    assert "本章不少于 100 字" in p1
    assert "建设目标响应程度得分" in p2
    assert "本章不少于 200 字" in p2


def test_generate_section_leaf_unchanged(client, monkeypatch):
    """无子节的章保持原逻辑：无 progress 事件。"""
    pid = _project(client)
    sid = client.post(f"/api/projects/{pid}/sections",
                      json={"title": "1.1 营业执照", "level": 2}).json()["id"]
    monkeypatch.setattr("app.services.write_service.get_api_key", lambda: "sk-test")
    monkeypatch.setattr("app.services.write_service.get_model", lambda: "kimi-k3")
    monkeypatch.setattr("app.services.write_service.stream_section",
                        lambda prompt, key, model: iter(["正文"]))
    r = client.get(f"/api/projects/{pid}/sections/{sid}/generate")
    assert r.status_code == 200
    assert "event: progress" not in r.text
    assert "event: done" in r.text
    stored = client.get(f"/api/projects/{pid}/sections").json()[0]
    assert stored["content"] == "正文"


# ---------- AI 改写章节标题 ----------

def test_adapt_titles_success(client, monkeypatch):
    pid = client.post("/api/projects", json={"name": "GZ659商业综合体项目"}).json()["id"]
    root = client.post(f"/api/projects/{pid}/sections",
                       json={"title": "第一章 工程概况", "level": 1}).json()["id"]
    client.post(f"/api/projects/{pid}/sections",
                json={"parent_id": root, "title": "1.1 工程简介", "level": 2})
    monkeypatch.setattr("app.services.write_service.get_api_key", lambda: "sk-test")
    monkeypatch.setattr("app.services.write_service.get_model", lambda: "kimi-k3")
    monkeypatch.setattr(
        "app.services.write_service.adapt_titles",
        lambda titles, name, summary, key, model: [
            "第一章 GZ659商业综合体项目概况", "1.1 GZ659商业综合体工程简介"])
    r = client.post(f"/api/projects/{pid}/sections/adapt-titles")
    assert r.status_code == 200
    sections = r.json()["sections"]
    assert sections[0]["title"] == "第一章 GZ659商业综合体项目概况"
    assert sections[0]["children"][0]["title"] == "1.1 GZ659商业综合体工程简介"
    # 已落库
    assert client.get(f"/api/projects/{pid}/sections").json()[0]["title"] == \
        "第一章 GZ659商业综合体项目概况"


def test_adapt_titles_failure_502_keeps_titles(client, monkeypatch):
    from app.core.llm_parser import LLMParseError
    pid = _project(client)
    root = client.post(f"/api/projects/{pid}/sections",
                       json={"title": "第一章 工程概况", "level": 1}).json()["id"]
    monkeypatch.setattr("app.services.write_service.get_api_key", lambda: "sk-test")
    monkeypatch.setattr("app.services.write_service.get_model", lambda: "kimi-k3")

    def boom(titles, name, summary, key, model):
        raise LLMParseError("API 调用失败")
    monkeypatch.setattr("app.services.write_service.adapt_titles", boom)
    r = client.post(f"/api/projects/{pid}/sections/adapt-titles")
    assert r.status_code == 502
    assert r.json()["detail"] == "标题改写失败，请重试"
    # 原标题不变
    assert client.get(f"/api/projects/{pid}/sections").json()[0]["title"] == "第一章 工程概况"


def test_adapt_titles_missing_project_404(client):
    assert client.post("/api/projects/999/sections/adapt-titles").status_code == 404


# ---------- 资产字段格式化 ----------

def test_assets_excerpt_performance_list_formatting(db_path):
    from app.db import Database
    from app.services.write_service import WriteService
    db = Database(db_path); db.init_schema()
    aid = db.create_asset("material", "企业业绩", fields={
        "企业名称": "宏信天德",
        "业绩": [
            {"项目名称": "A项目", "年份": "2023", "金额": "500万"},
            {"项目名称": "B项目", "年份": "2024"},
        ],
    })
    svc = WriteService(db_path)
    text = svc._assets_excerpt([aid], [])
    assert "【企业业绩】" in text
    assert "企业名称：宏信天德" in text
    assert "1. 项目名称：A项目，年份：2023，金额：500万" in text
    assert "2. 项目名称：B项目，年份：2024" in text
    assert "'项目名称'" not in text  # 不再是 Python repr


def test_match_requirements_by_title(db_path):
    from app.services.write_service import WriteService
    reqs = [
        {"id": 1, "content": "项目背景应详细说明"},
        {"id": 2, "content": "保证金缴纳方式"},
    ]
    matched = WriteService._match_requirements(reqs, "第一章 项目背景")
    assert [r["id"] for r in matched] == [1]
    # 无有效词（纯编号标题）→ 返回全部
    assert len(WriteService._match_requirements(reqs, "第一章")) == 2
