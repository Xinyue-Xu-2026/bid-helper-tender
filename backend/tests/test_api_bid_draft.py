"""商务标底稿 API 测试（T9）：headings/generate/upload/preview/bindings 回显/
重新生成清空绑定/导出集成分叉（有底稿走 fill_draft 新管线 + X-Fill-Report，
无底稿回退旧模板路径且无该头）/各端点 404。"""
import json
from io import BytesIO
from urllib.parse import unquote

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import config
from app.db import Database
from app.deps import get_db
from app.main import app

PERSON_HEADERS = ["序号", "姓名", "职称", "专业工作年限", "执业资格"]
DOCX_MIME = ("application/vnd.openxmlformats-officedocument"
             ".wordprocessingml.document")


@pytest.fixture()
def client(db_path):
    def override():
        db = Database(db_path)
        db.init_schema()
        yield db
    app.dependency_overrides[get_db] = override
    yield TestClient(app)
    app.dependency_overrides.clear()


def _build_tender(path, with_format_chapter=True, tenderer=""):
    """程序化招标文件：第一章 招标公告 + 第三章 投标文件格式（含人员表）。"""
    doc = Document()
    doc.add_heading("第一章 招标公告", level=1)
    if tenderer:
        doc.add_paragraph(f"招标人：{tenderer}")
    doc.add_paragraph("招标正文……")
    if with_format_chapter:
        doc.add_heading("第三章 投标文件格式", level=1)
        doc.add_heading("一、投标函", level=2)
        table = doc.add_table(rows=2, cols=len(PERSON_HEADERS))
        for j, h in enumerate(PERSON_HEADERS):
            table.cell(0, j).text = h
        for j in range(len(PERSON_HEADERS)):
            table.cell(1, j).text = f"旧{j}"
        doc.add_heading("二、商务条款响应表", level=2)
    else:
        doc.add_heading("第二章 投标人须知", level=1)
        doc.add_paragraph("无格式章节正文")
    doc.save(str(path))
    return path


def _make_project(client, name="底稿项目"):
    return client.post("/api/projects", json={"name": name}).json()["id"]


def _make_person(client, name, fields=None):
    return client.post("/api/assets", json={
        "type": "person", "name": name, "fields": fields or {}}).json()["id"]


def _upload_tender(client, pid, tmp_path, with_format_chapter=True, tenderer=""):
    path = _build_tender(tmp_path / "tender.docx", with_format_chapter, tenderer)
    with open(path, "rb") as f:
        r = client.post(f"/api/projects/{pid}/tender",
                        files={"file": ("tender.docx", f.read(), DOCX_MIME)})
    assert r.status_code == 200
    return r


def _generate(client, pid):
    r = client.post(f"/api/projects/{pid}/bid-draft/generate", json={})
    assert r.status_code == 200
    return r.json()


def _confirm_first_table(client, pid, gen, **over):
    """把 generate 预览中第 0 张表确认绑定（role/columns 沿用分类建议）。"""
    sug = gen["tables"][0]
    binding = {"table_index": sug["table_index"], "role": sug["role"],
               "columns": sug["columns"], "person_scope": "all",
               "perf_scope": "all", "label_kind": "", "person": "",
               "confirmed": True}
    binding.update(over)
    r = client.put(f"/api/projects/{pid}/bid-draft/bindings",
                   json={"tables": [binding], "swap_toc": False})
    assert r.status_code == 200
    return binding


# ---------- 1. headings ----------

def test_headings_requires_tender(client):
    pid = _make_project(client)
    r = client.get(f"/api/projects/{pid}/bid-draft/headings")
    assert r.status_code == 400


def test_headings_after_upload(client, tmp_path):
    pid = _make_project(client)
    _upload_tender(client, pid, tmp_path)
    r = client.get(f"/api/projects/{pid}/bid-draft/headings")
    assert r.status_code == 200
    data = r.json()
    assert any("投标文件格式" in h["title"] for h in data["headings"])
    assert data["suggested"] and "格式" in data["suggested"]["start"]


# ---------- 2. generate ----------

def test_generate_ok(client, tmp_path):
    pid = _make_project(client)
    _upload_tender(client, pid, tmp_path)
    data = _generate(client, pid)
    assert data["source"] == "tender-cut"
    assert "投标文件格式" in data["cut_start"]
    assert data["outline"]
    assert len(data["tables"]) == 1
    assert data["tables"][0]["role"] == "person_roster"
    assert data["tables"][0]["confirmed"] is False
    assert data["tables"][0]["header"] == PERSON_HEADERS


def test_generate_no_format_chapter_422(client, tmp_path):
    pid = _make_project(client)
    _upload_tender(client, pid, tmp_path, with_format_chapter=False)
    r = client.post(f"/api/projects/{pid}/bid-draft/generate", json={})
    assert r.status_code == 422
    assert "格式" in r.json()["detail"]


# ---------- 招标人提取（P3） ----------

def test_generate_extracts_tenderer(client, tmp_path):
    """招标文件封面"招标人：XXX"（前 30 个非空段落内）→ 底稿行与预览回显。"""
    pid = _make_project(client)
    _upload_tender(client, pid, tmp_path, tenderer="高邮市水利建设服务中心")
    data = _generate(client, pid)
    assert data["tenderer"] == "高邮市水利建设服务中心"
    g = client.get(f"/api/projects/{pid}/bid-draft").json()
    assert g["tenderer"] == "高邮市水利建设服务中心"


def test_generate_tenderer_absent_empty(client, tmp_path):
    """招标文件无"招标人："行 → tenderer 为空串。"""
    pid = _make_project(client)
    _upload_tender(client, pid, tmp_path)
    data = _generate(client, pid)
    assert data["tenderer"] == ""


def test_generate_with_explicit_indices(client, tmp_path):
    """索引直达：start_index/end_index 直接裁切，标题仅作展示（Fix C 回归）。"""
    pid = _make_project(client)
    _upload_tender(client, pid, tmp_path)
    headings = client.get(f"/api/projects/{pid}/bid-draft/headings").json()["headings"]
    start = next(h for h in headings if "投标文件格式" in h["title"])
    end = next(h for h in headings if "商务条款" in h["title"])
    # end=-1（文档末尾）
    r = client.post(f"/api/projects/{pid}/bid-draft/generate",
                    json={"start_index": start["index"], "end_index": -1})
    assert r.status_code == 200
    data = r.json()
    assert "投标文件格式" in data["cut_start"]
    assert len(data["tables"]) == 1
    # 显式 end_index：裁到"二、商务条款响应表"之前
    r2 = client.post(f"/api/projects/{pid}/bid-draft/generate",
                     json={"start_index": start["index"],
                           "end_index": end["index"]})
    assert r2.status_code == 200
    data2 = r2.json()
    assert data2["cut_end"] == end["title"]
    outline_titles = [t for _, t in data2["outline"]]
    assert not any("商务条款" in t for t in outline_titles)
    assert len(data2["tables"]) == 1  # 人员表在结束标题之前，保留


def test_generate_with_out_of_range_index_422(client, tmp_path):
    pid = _make_project(client)
    _upload_tender(client, pid, tmp_path)
    r = client.post(f"/api/projects/{pid}/bid-draft/generate",
                    json={"start_index": 99999})
    assert r.status_code == 422
    assert "索引" in r.json()["detail"]


# ---------- 3. upload ----------

def test_upload_rejects_non_docx(client):
    pid = _make_project(client)
    r = client.post(f"/api/projects/{pid}/bid-draft/upload",
                    files={"file": ("draft.pdf", b"%PDF-1.4",
                                    "application/pdf")})
    assert r.status_code == 400


def test_upload_docx_ok(client, tmp_path):
    pid = _make_project(client)
    path = tmp_path / "draft.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=len(PERSON_HEADERS))
    for j, h in enumerate(PERSON_HEADERS):
        table.cell(0, j).text = h
    doc.save(str(path))
    with open(path, "rb") as f:
        r = client.post(f"/api/projects/{pid}/bid-draft/upload",
                        files={"file": ("draft.docx", f.read(), DOCX_MIME)})
    assert r.status_code == 200
    data = r.json()
    assert data["source"] == "upload"
    assert data["cut_start"] == "" and data["cut_end"] == ""
    assert data["tables"][0]["role"] == "person_roster"
    # GET 回显同一份底稿
    g = client.get(f"/api/projects/{pid}/bid-draft")
    assert g.status_code == 200
    assert g.json()["source"] == "upload"


def test_get_draft_none(client):
    pid = _make_project(client)
    r = client.get(f"/api/projects/{pid}/bid-draft")
    assert r.status_code == 200
    assert r.json() == {"draft": None}


# ---------- 4. PUT bindings 校验与回显 ----------

def test_put_bindings_invalid_role_422(client, tmp_path):
    pid = _make_project(client)
    _upload_tender(client, pid, tmp_path)
    _generate(client, pid)
    r = client.put(f"/api/projects/{pid}/bid-draft/bindings", json={
        "tables": [{"table_index": 0, "role": "不存在角色", "confirmed": True}],
        "swap_toc": False})
    assert r.status_code == 422


def test_put_bindings_no_draft_422(client):
    pid = _make_project(client)
    r = client.put(f"/api/projects/{pid}/bid-draft/bindings",
                   json={"tables": [], "swap_toc": False})
    assert r.status_code == 422


def test_put_bindings_invalid_header_rows_422(client, tmp_path):
    """header_rows 合法域 1..4（P1 回归）：0 / 5 → 422。"""
    pid = _make_project(client)
    _upload_tender(client, pid, tmp_path)
    gen = _generate(client, pid)
    sug = gen["tables"][0]
    base = {"table_index": sug["table_index"], "role": sug["role"],
            "columns": sug["columns"], "confirmed": True}
    for bad in (0, 5):
        r = client.put(f"/api/projects/{pid}/bid-draft/bindings",
                       json={"tables": [{**base, "header_rows": bad}],
                             "swap_toc": False})
        assert r.status_code == 422


def test_put_bindings_header_rows_roundtrip(client, tmp_path):
    """header_rows 经 PUT 存库后在 GET 预览中回显（透传）。"""
    pid = _make_project(client)
    _upload_tender(client, pid, tmp_path)
    gen = _generate(client, pid)
    sug = gen["tables"][0]
    assert sug["header_rows"] == 1  # 程序化单行表头 → 1
    r = client.put(f"/api/projects/{pid}/bid-draft/bindings", json={
        "tables": [{"table_index": sug["table_index"], "role": sug["role"],
                    "columns": sug["columns"], "header_rows": 2,
                    "confirmed": True}],
        "swap_toc": False})
    assert r.status_code == 200
    g = client.get(f"/api/projects/{pid}/bid-draft").json()
    assert g["tables"][0]["header_rows"] == 2


def test_put_bindings_header_rows_three_accepted(client, tmp_path):
    """header_rows=3 属合法域 1..4（V1.2 Task 4 边界补充）→ 200 且回显。"""
    pid = _make_project(client)
    _upload_tender(client, pid, tmp_path)
    gen = _generate(client, pid)
    sug = gen["tables"][0]
    r = client.put(f"/api/projects/{pid}/bid-draft/bindings", json={
        "tables": [{"table_index": sug["table_index"], "role": sug["role"],
                    "columns": sug["columns"], "header_rows": 3,
                    "confirmed": True}],
        "swap_toc": False})
    assert r.status_code == 200
    g = client.get(f"/api/projects/{pid}/bid-draft").json()
    assert g["tables"][0]["header_rows"] == 3


def test_put_bindings_mode_validation_and_roundtrip(client, tmp_path):
    """mode（V1.2 一人一表）："per_person"/"" 接受且回显；非法值 → 422。"""
    pid = _make_project(client)
    _upload_tender(client, pid, tmp_path)
    gen = _generate(client, pid)
    sug = gen["tables"][0]
    base = {"table_index": sug["table_index"], "role": sug["role"],
            "columns": sug["columns"], "confirmed": True}
    r = client.put(f"/api/projects/{pid}/bid-draft/bindings",
                   json={"tables": [{**base, "mode": "per_person"}],
                         "swap_toc": False})
    assert r.status_code == 200
    g = client.get(f"/api/projects/{pid}/bid-draft").json()
    assert g["tables"][0]["mode"] == "per_person"
    r = client.put(f"/api/projects/{pid}/bid-draft/bindings",
                   json={"tables": [{**base, "mode": "bogus"}],
                         "swap_toc": False})
    assert r.status_code == 422


# ---------- 占位符预览（V1.2 Task 8a） ----------

def test_placeholders_preview_no_draft_404(client):
    pid = _make_project(client)
    r = client.post(f"/api/projects/{pid}/bid-draft/placeholders", json={})
    assert r.status_code == 404


def test_placeholders_preview_matched(client, tmp_path, db_path):
    """底稿含"招标人：____" → 预览 matched 含 招标人=传入 tenderer。"""
    from app.db import Database
    pid = _make_project(client)
    doc = Document()
    doc.add_paragraph("招标人：____")
    p = tmp_path / "draft.docx"
    doc.save(str(p))
    Database(db_path).create_bid_template(pid, "底稿", str(p))
    r = client.post(f"/api/projects/{pid}/bid-draft/placeholders",
                    json={"tenderer": "某中心"})
    assert r.status_code == 200
    data = r.json()
    assert any(m["label"] == "招标人" and m["value"] == "某中心"
               for m in data["matched"])
    assert "suspicious" in data


def test_put_bindings_roundtrip(client, tmp_path):
    pid = _make_project(client)
    _upload_tender(client, pid, tmp_path)
    gen = _generate(client, pid)
    sug = gen["tables"][0]
    r = client.put(f"/api/projects/{pid}/bid-draft/bindings", json={
        "tables": [{"table_index": sug["table_index"], "role": sug["role"],
                    "columns": sug["columns"], "person_scope": "all",
                    "perf_scope": "all", "label_kind": "", "person": "",
                    "confirmed": True}],
        "swap_toc": True})
    assert r.status_code == 200
    assert r.json() == {"ok": True}
    g = client.get(f"/api/projects/{pid}/bid-draft").json()
    assert g["swap_toc"] is True
    t = g["tables"][0]
    assert t["confirmed"] is True
    assert t["role"] == "person_roster"
    assert t["columns"] == sug["columns"]
    assert t["header"] == PERSON_HEADERS  # 展示字段来自分类建议


# ---------- 5. 重新 generate → 旧绑定清空 ----------

def test_regenerate_clears_confirmed_bindings(client, tmp_path):
    pid = _make_project(client)
    _upload_tender(client, pid, tmp_path)
    gen = _generate(client, pid)
    _confirm_first_table(client, pid, gen)
    g = client.get(f"/api/projects/{pid}/bid-draft").json()
    assert g["tables"][0]["confirmed"] is True
    _generate(client, pid)  # 重新生成
    g2 = client.get(f"/api/projects/{pid}/bid-draft").json()
    assert g2["tables"][0]["confirmed"] is False


# ---------- 6. 导出集成分叉 ----------

def test_export_template_with_draft_uses_fill_pipeline(client, tmp_path):
    pid = _make_project(client)
    _upload_tender(client, pid, tmp_path)
    gen = _generate(client, pid)
    _confirm_first_table(client, pid, gen)
    person_id = _make_person(client, "张三", {"职称": "高级工程师"})
    r = client.post(f"/api/projects/{pid}/bid-assets/export-template",
                    json={"persons": [{"asset_id": person_id,
                                       "is_lead": True}],
                          "contracts": [],
                          "tenderer": "某服务中心",
                          "bidder_name": "宏信天德工程顾问有限公司"})
    assert r.status_code == 200
    header = r.headers.get("x-fill-report")
    assert header, "有底稿导出必须带 X-Fill-Report"
    report = json.loads(unquote(header))
    assert report["verify"]["ok"] is True
    assert report["verify"]["issues"] == []
    assert report["filled"][0]["role"] == "person_roster"
    assert report["filled"][0]["rows"] == 1
    # 产物确实走了新管线：人员表数据行被填为张三
    doc = Document(BytesIO(r.content))
    rows = [[c.text for c in row.cells] for row in doc.tables[0].rows]
    assert rows[1][1] == "张三"
    assert rows[1][2] == "高级工程师"


def test_export_template_without_draft_keeps_legacy(client, tmp_path,
                                                    monkeypatch):
    """无底稿项目：仍走旧模板路径，无 X-Fill-Report 头。"""
    tpl = tmp_path / "商务标模板.docx"
    Document().save(str(tpl))  # 最小旧模板（无表可定位，原样输出）
    monkeypatch.setattr(config, "BID_TEMPLATE_PATH", tpl)
    pid = _make_project(client)
    person_id = _make_person(client, "张三")
    r = client.post(f"/api/projects/{pid}/bid-assets/export-template",
                    json={"persons": [{"asset_id": person_id,
                                       "is_lead": True}],
                          "contracts": []})
    assert r.status_code == 200
    assert "x-fill-report" not in r.headers


# ---------- 7. 项目不存在 → 404 ----------

def test_endpoints_404(client, tmp_path):
    assert client.get(
        "/api/projects/99999/bid-draft/headings").status_code == 404
    assert client.get("/api/projects/99999/bid-draft").status_code == 404
    assert client.post("/api/projects/99999/bid-draft/generate",
                       json={}).status_code == 404
    assert client.put("/api/projects/99999/bid-draft/bindings",
                      json={"tables": []}).status_code == 404
    r = client.post("/api/projects/99999/bid-draft/upload",
                    files={"file": ("draft.docx", b"x", DOCX_MIME)})
    assert r.status_code == 404
