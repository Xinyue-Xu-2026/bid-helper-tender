"""底稿驱动填充导出测试（T7）：程序化构造底稿（封面残留编号/名称/日期 +
toc 段 + 人员表 + 业绩表 + 2列简历表 + 1列图片占位表 + 不绑定承诺函表），
覆盖 person_roster/perf_list/lead_resume/image_slot 填充、横向合并、
stale 替换、swap_toc、报告结构与分类器共享键回归。"""
import base64

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn

from app.core.bid_draft_exporter import fill_draft
from app.core.bid_table_classifier import classify_tables

# 1x1 PNG（python-docx 可插图）
PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk"
    "+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg==")

PROMISE_TEXT = "本单位郑重承诺：所投材料真实有效。"


def _add_table(doc, rows_data):
    table = doc.add_table(rows=len(rows_data), cols=len(rows_data[0]))
    for r, row in enumerate(rows_data):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    return table


def _build_draft(path, with_toc=True):
    """构造底稿：封面残留 + toc 段 + 人员表 + 业绩表 + 简历表 +
    图片占位表 + 不绑定承诺函表。返回路径字符串。"""
    doc = Document()
    doc.add_paragraph("编       号：OLD-2020-001号")
    doc.add_paragraph("项 目 名 称：旧项目名AAA")
    doc.add_paragraph("日      期 ：2020年1月2日")
    if with_toc:
        doc.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
        doc.add_paragraph("一、开标一览表\t1", style="toc 1")
        doc.add_paragraph("二、承诺函\t2", style="toc 1")
    _add_table(doc, [                                   # 表0 人员表
        ["序号", "姓名", "职称", "专业工作年限", "执业资格", "备注"],
        ["9", "旧人", "旧职称", "1", "旧资格", "旧备注"],
    ])
    _add_table(doc, [                                   # 表1 业绩表
        ["序号", "项目名称", "委托单位", "合同签订时间"],
        ["9", "旧项目", "旧单位", "2020-01-01"],
    ])
    _add_table(doc, [                                   # 表2 简历表（2列键值）
        ["姓名", ""], ["职称", ""], ["已完项目", ""],
    ])
    _add_table(doc, [["注册证书扫描件"], [""]])         # 表3 图片占位表
    _add_table(doc, [[PROMISE_TEXT]])                   # 表4 承诺函（不绑定）
    doc.save(str(path))
    return str(path)


def _bindings(**overrides):
    bindings = {
        "tables": [
            {"table_index": 0, "role": "person_roster",
             "columns": {"seq": 0, "name": 1, "title": 2,
                         "work_years": 3, "certs": 4},
             "person_scope": "all", "perf_scope": "", "label_kind": "",
             "person": "", "confirmed": True},
            {"table_index": 1, "role": "perf_list",
             "columns": {"seq": 0, "project_name": 1, "client": 2,
                         "sign_date": 3},
             "person_scope": "", "perf_scope": "all", "label_kind": "",
             "person": "", "confirmed": True},
            {"table_index": 2, "role": "lead_resume",
             "columns": {"姓名": "name", "职称": "title",
                         "已完项目": "lead_perfs"},
             "person_scope": "", "perf_scope": "", "label_kind": "",
             "person": "", "confirmed": True},
            {"table_index": 3, "role": "image_slot", "columns": {},
             "person_scope": "", "perf_scope": "",
             "label_kind": "注册证书", "person": "lead",
             "confirmed": True},
        ],
        "swap_toc": False,
    }
    bindings.update(overrides)
    return bindings


def _person(name, is_lead, title="", certs="", work_years="", scan=""):
    return {
        "name": name, "is_lead": is_lead, "role": "",
        "fields": {"职称": title,
                   "证书": [{"类型": "一级造价师", "扫描件": scan}]},
        "sem": {"seq": "1" if is_lead else "2", "label": "", "name": name,
                "gender": "", "age": "", "education": "", "title": title,
                "work_years": work_years, "certs": certs, "role": ""},
    }


def _contract(name, section, client="某单位", sign="2024-01-01", leader=""):
    return {"name": name, "section": section,
            "fields": {"委托单位": client, "签订日期": sign,
                       "项目负责人": leader},
            "sem": {"seq": "1", "project_name": name, "client": client,
                    "client_contact": "", "sign_date": sign, "amount": "",
                    "project_type": "", "service_type": "", "content": "",
                    "proof_page": ""}}


def _data(scan=""):
    return {
        "persons": [_person("张三", True, title="高级工程师",
                            certs="一级造价师", work_years="15", scan=scan),
                    _person("李四", False, title="工程师",
                            certs="一级建造师", work_years="8")],
        "contracts": [_contract("项目A", 1, leader="张三"),
                      _contract("项目B", 2, leader="王五")],
        "lead_perfs_text": "项目A（2024）",
    }


def _rows(table):
    return [[cell.text for cell in row.cells] for row in table.rows]


# ---------- 1. person_roster 填充 ----------

def test_person_roster_fill(tmp_path):
    draft = _build_draft(tmp_path / "draft.docx")
    out = str(tmp_path / "out.docx")
    report = fill_draft(draft, out, _bindings(), _data())
    rows = _rows(Document(out).tables[0])
    assert rows[0] == ["序号", "姓名", "职称", "专业工作年限", "执业资格", "备注"]
    assert len(rows) == 3  # 表头 + 2 人
    assert rows[1][:5] == ["1", "张三", "高级工程师", "15", "一级造价师"]
    assert rows[2][:5] == ["2", "李四", "工程师", "8", "一级建造师"]
    # columns 未映射的"备注"列被清为 ""
    assert rows[1][5] == "" and rows[2][5] == ""
    filled = {f["role"]: f for f in report["filled"]}
    assert filled["person_roster"]["rows"] == 2
    assert filled["person_roster"]["table_index"] == 0


# ---------- 2. 空 persons → 仅表头 ----------

def test_person_roster_empty_persons(tmp_path):
    data = _data()
    data["persons"] = []
    data["lead_perfs_text"] = ""
    draft = _build_draft(tmp_path / "draft.docx")
    out = str(tmp_path / "out.docx")
    report = fill_draft(draft, out, _bindings(), data)
    rows = _rows(Document(out).tables[0])
    assert len(rows) == 1  # 仅表头
    filled = {f["role"]: f for f in report["filled"]}
    assert filled["person_roster"]["rows"] == 0


# ---------- 3. perf_list perf_scope=section1 ----------

def test_perf_list_section1_scope(tmp_path):
    bindings = _bindings()
    bindings["tables"][1]["perf_scope"] = "section1"
    draft = _build_draft(tmp_path / "draft.docx")
    out = str(tmp_path / "out.docx")
    fill_draft(draft, out, bindings, _data())
    rows = _rows(Document(out).tables[1])
    assert len(rows) == 2  # 表头 + 仅 section==1 的项目A
    assert rows[1][1] == "项目A"


# ---------- 4. lead_resume 填充 / 无负责人不写 ----------

def test_lead_resume_fill(tmp_path):
    draft = _build_draft(tmp_path / "draft.docx")
    out = str(tmp_path / "out.docx")
    report = fill_draft(draft, out, _bindings(), _data())
    rows = _rows(Document(out).tables[2])
    assert rows[0] == ["姓名", "张三"]
    assert rows[1] == ["职称", "高级工程师"]
    assert rows[2] == ["已完项目", "项目A（2024）"]
    filled = {f["role"]: f for f in report["filled"]}
    assert filled["lead_resume"]["rows"] == 3


def test_lead_resume_no_lead_writes_nothing(tmp_path):
    data = _data()
    data["persons"] = [p for p in data["persons"] if not p["is_lead"]]
    draft = _build_draft(tmp_path / "draft.docx")
    out = str(tmp_path / "out.docx")
    report = fill_draft(draft, out, _bindings(), data)
    rows = _rows(Document(out).tables[2])
    assert rows[0] == ["姓名", ""]
    assert rows[2] == ["已完项目", ""]
    filled = {f["role"]: f for f in report["filled"]}
    assert filled["lead_resume"]["rows"] == 0


# ---------- 5. image_slot：插图 / 文件缺失不抛 ----------

def test_image_slot_insert(tmp_path):
    png = tmp_path / "cert.png"
    png.write_bytes(PNG_1X1)
    draft = _build_draft(tmp_path / "draft.docx")
    out = str(tmp_path / "out.docx")
    report = fill_draft(draft, out, _bindings(), _data(scan=str(png)))
    table = Document(out).tables[3]
    assert table._tbl.findall(f".//{qn('w:drawing')}")
    img = report["images"][0]
    assert img == {"table_index": 3, "person": "张三",
                   "label_kind": "注册证书", "ok": True}


def test_image_slot_missing_file_no_throw(tmp_path):
    draft = _build_draft(tmp_path / "draft.docx")
    out = str(tmp_path / "out.docx")
    report = fill_draft(draft, out, _bindings(),
                        _data(scan=str(tmp_path / "不存在.png")))
    assert report["images"][0]["ok"] is False
    assert not Document(out).tables[3]._tbl.findall(f".//{qn('w:drawing')}")


# ---------- 6. 横向合并单元格：同 _tc 只写一次 ----------

def test_merged_cells_written_once(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    for j, h in enumerate(["序号", "姓名", "职称", "备注"]):
        table.cell(0, j).text = h
    for j in range(4):
        table.cell(1, j).text = f"旧{j}"
    table.cell(1, 0).merge(table.cell(1, 1))  # 数据行横向合并
    draft = str(tmp_path / "draft.docx")
    doc.save(draft)
    bindings = {"tables": [
        {"table_index": 0, "role": "person_roster",
         "columns": {"seq": 0, "name": 1, "title": 2},
         "person_scope": "all", "perf_scope": "", "label_kind": "",
         "person": "", "confirmed": True},
    ], "swap_toc": False}
    data = _data()
    data["persons"] = data["persons"][:1]
    fill_draft(draft, str(tmp_path / "out.docx"), bindings, data)
    texts = [c.text for c in Document(str(tmp_path / "out.docx"))
             .tables[0].rows[1].cells]
    # 合并格同一 _tc 只写一次（seq=1 未被二次写入覆盖/叠加）
    assert texts[0] == "1" and texts[1] == "1"
    assert texts[2] == "高级工程师"


# ---------- 7. stale 替换生效 + 未绑定承诺函表零改动 ----------

def test_stale_replacement_and_unbound_table_untouched(tmp_path):
    draft = _build_draft(tmp_path / "draft.docx")
    out = str(tmp_path / "out.docx")
    fill_draft(draft, out, _bindings(), _data(),
               project_no="NEW-2026-123", project_name="新项目XYZ",
               doc_date="2026-09-08")
    doc = Document(out)
    texts = [p.text for p in doc.paragraphs]
    assert "编       号：NEW-2026-123号" in texts
    assert "项 目 名 称：新项目XYZ" in texts
    assert "日      期 ：2026年9月8日" in texts
    assert not any("OLD-2020-001" in t for t in texts)
    assert not any("旧项目名AAA" in t for t in texts)
    # 未绑定承诺函表文本零改动
    assert doc.tables[4].cell(0, 0).text == PROMISE_TEXT


# ---------- 8. swap_toc → toc 段替换为 TOC 域 ----------

def test_swap_toc(tmp_path):
    bindings = _bindings(swap_toc=True)
    draft = _build_draft(tmp_path / "draft.docx")
    out = str(tmp_path / "out.docx")
    report = fill_draft(draft, out, bindings, _data())
    doc = Document(out)
    toc_paras = [p for p in doc.paragraphs
                 if (p.style.name or "").lower().startswith("toc")]
    assert len(toc_paras) == 1  # 连续 toc 段组只留首段
    xml = toc_paras[0]._element.xml
    assert "instrText" in xml and "TOC" in xml
    assert "开标一览表" not in toc_paras[0].text
    assert report["verify"]["ok"] is True


# ---------- 9. report 结构完整 + quote/越界进 skipped + verify 正常 ----------

def test_report_structure_and_quote_skip(tmp_path):
    bindings = _bindings()
    bindings["tables"].append(
        {"table_index": 4, "role": "quote", "columns": {},
         "person_scope": "", "perf_scope": "", "label_kind": "",
         "person": "", "confirmed": True})
    bindings["tables"].append(
        {"table_index": 99, "role": "person_roster", "columns": {},
         "person_scope": "all", "perf_scope": "", "label_kind": "",
         "person": "", "confirmed": True})
    draft = _build_draft(tmp_path / "draft.docx")
    out = str(tmp_path / "out.docx")
    report = fill_draft(draft, out, bindings, _data())
    assert set(report) == {"filled", "images", "skipped", "verify"}
    reasons = {(s.get("table_index"), s.get("reason"))
               for s in report["skipped"]}
    assert (4, "报价表需手工填写") in reasons
    assert (99, "表不存在，可能底稿已重新生成") in reasons
    # quote 表未被填充改动
    assert Document(out).tables[4].cell(0, 0).text == PROMISE_TEXT
    verify = report["verify"]
    assert verify["ok"] is True
    assert verify["issues"] == []
    assert verify["checked_paragraphs"] == 3   # 3 封面（toc 段恒跳过）
    assert verify["checked_tables"] == 1  # quote 表不再豁免校验（fill 不改动它）


def test_verify_counts_unbound_table(tmp_path):
    draft = _build_draft(tmp_path / "draft.docx")
    out = str(tmp_path / "out.docx")
    report = fill_draft(draft, out, _bindings(), _data())
    verify = report["verify"]
    assert verify["ok"] is True
    assert verify["checked_tables"] == 1  # 仅承诺函表未绑定


# ---------- 10. 分类器共享键回归：纯业绩表含"序号" → perf columns 含 seq ----------

def test_classifier_shared_seq_goes_to_perf(tmp_path):
    doc = Document()
    _add_table(doc, [
        ["序号", "项目名称", "委托单位", "合同签订时间"],
        ["1", "某项目", "某单位", "2024-01-01"],
    ])
    path = str(tmp_path / "c.docx")
    doc.save(path)
    item = classify_tables(path)[0]
    assert item["role"] == "perf_list"
    assert item["columns"]["seq"] == 0
    assert item["columns"]["project_name"] == 1
